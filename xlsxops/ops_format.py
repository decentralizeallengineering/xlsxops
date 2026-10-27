# SPDX-License-Identifier: Apache-2.0
"""Charts, conditional formats, data validation and range formatting."""
from __future__ import annotations

import datetime as _dt
import math
import re
import statistics

from openpyxl.chart import BarChart, LineChart, PieChart, Reference, ScatterChart, Series
from openpyxl.drawing.spreadsheet_drawing import AbsoluteAnchor, OneCellAnchor, TwoCellAnchor
from openpyxl.formatting.rule import CellIsRule, ColorScaleRule, DataBarRule, FormulaRule
from openpyxl.styles import Font, PatternFill
from openpyxl.utils.cell import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from . import formula as fx
from ._core import (bolded, S_COLS, S_FORCE, S_HDR, S_OUT, S_PATH, S_SHEET, Book, UserError,
                    ValueReader, autofit_columns, find_table, finish, operation, parse_cell,
                    parse_range, table_warnings)

EMU_PER_PX = 9525
EMU_PER_CM = 360000
_range = range  # several operations take a parameter named "range"


# --------------------------------------------------------------------------
# geometry helpers for chart placement
# --------------------------------------------------------------------------
class _Grid:
    def __init__(self, ws):
        self.ws = ws
        default_w = ws.sheet_format.defaultColWidth or 8.43
        self.default_col = int(default_w * 7 + 5)
        self.default_row = (ws.sheet_format.defaultRowHeight or 15) * 4 / 3
        self.cols = {}
        for dim in ws.column_dimensions.values():
            lo, hi = dim.min or 0, dim.max or 0
            if not lo:
                continue
            px = 0 if dim.hidden else (int(dim.width * 7 + 5) if dim.width else self.default_col)
            for c in range(lo, min(hi, lo + 2000) + 1):
                self.cols[c] = px
        self.rows = {r: (0 if d.hidden else (d.height * 4 / 3 if d.height else self.default_row))
                     for r, d in ws.row_dimensions.items()}

    def col_px(self, c):
        return self.cols.get(c, self.default_col)

    def row_px(self, r):
        return self.rows.get(r, self.default_row)

    def x_of(self, col0):  # 0-based column index -> px
        return sum(self.col_px(c) for c in range(1, col0 + 1))

    def y_of(self, row0):
        return sum(self.row_px(r) for r in range(1, row0 + 1))

    def row_at(self, y):  # px -> 0-based row index containing y
        acc, r = 0.0, 1
        while acc + self.row_px(r) <= y and r < 1_048_576:
            acc += self.row_px(r)
            r += 1
        return r - 1


def _rect(grid, obj):
    a = obj.anchor
    if isinstance(a, str):
        r, c = parse_cell(a)
        x, y = grid.x_of(c - 1), grid.y_of(r - 1)
        if hasattr(obj, "series"):
            w, h = obj.width * EMU_PER_CM / EMU_PER_PX, obj.height * EMU_PER_CM / EMU_PER_PX
        else:
            w, h = obj.width, obj.height
        return (x, y, x + w, y + h)
    if isinstance(a, TwoCellAnchor):
        f, t = a._from, a.to
        return (grid.x_of(f.col) + f.colOff / EMU_PER_PX, grid.y_of(f.row) + f.rowOff / EMU_PER_PX,
                grid.x_of(t.col) + t.colOff / EMU_PER_PX, grid.y_of(t.row) + t.rowOff / EMU_PER_PX)
    if isinstance(a, OneCellAnchor):
        f = a._from
        x, y = grid.x_of(f.col) + f.colOff / EMU_PER_PX, grid.y_of(f.row) + f.rowOff / EMU_PER_PX
        return (x, y, x + a.ext.width / EMU_PER_PX, y + a.ext.height / EMU_PER_PX)
    if isinstance(a, AbsoluteAnchor):
        x, y = a.pos.x / EMU_PER_PX, a.pos.y / EMU_PER_PX
        return (x, y, x + a.ext.width / EMU_PER_PX, y + a.ext.height / EMU_PER_PX)
    return None


def _overlaps(a, b):
    return a[0] < b[2] and b[0] < a[2] and a[1] < b[3] and b[1] < a[3]


def _place(ws, width_cm, height_cm, start_col, start_row):
    """First free spot (column start_col, rows from start_row down) for a chart."""
    grid = _Grid(ws)
    occupied = [r for r in (_rect(grid, o) for o in list(ws._charts) + list(ws._images)) if r]
    w, h = width_cm * EMU_PER_CM / EMU_PER_PX, height_cm * EMU_PER_CM / EMU_PER_PX
    x = grid.x_of(start_col - 1)
    row0 = start_row - 1
    for _ in range(500):
        y = grid.y_of(row0)
        cand = (x, y, x + w, y + h)
        hits = [o for o in occupied if _overlaps(cand, o)]
        if not hits:
            return f"{get_column_letter(start_col)}{row0 + 1}", len(occupied)
        row0 = grid.row_at(max(o[3] for o in hits)) + 2
    raise UserError("could not find free space for the chart")


# --------------------------------------------------------------------------
# chart
# --------------------------------------------------------------------------
@operation("chart", {
    "type": "object", "required": ["path", "type"],
    "properties": {"path": S_PATH, "sheet": S_SHEET,
                   "type": {"enum": ["bar", "line", "pie", "scatter"]},
                   "x": {"type": "string"}, "y": S_COLS, "range": {"type": "string"},
                   "title": {"type": "string"}, "x_title": {"type": "string"},
                   "y_title": {"type": "string"}, "anchor": {"type": "string"},
                   "width_cm": {"type": "number"}, "height_cm": {"type": "number"},
                   "header_row": S_HDR, "out_path": S_OUT, "force": S_FORCE},
})
def chart(path, type, sheet=None, x=None, y=None, range=None, title=None, x_title=None,
          y_title=None, anchor=None, width_cm=15.0, height_cm=7.5, header_row=None,
          out_path=None, force=False):
    """Add a bar/line/pie/scatter chart from table columns (x + y) or a range.

    With 'range', the first row holds series names and the first column the categories.
    The chart is placed right of the data, below any existing charts, never overlapping.
    """
    book = Book(path, force=force)
    ws = book.ws(sheet)
    warnings = []
    if range:
        c1, r1, c2, r2 = parse_range(range)
        if c2 <= c1 or r2 <= r1:
            raise UserError("range needs a header row and at least two columns")
        hdr, first, last, xcol, ycols = r1, r1 + 1, r2, c1, list(_range(c1 + 1, c2 + 1))
        used_last_col = max(c2, max((c for (_, c), cell in ws._cells.items()
                                     if cell.value is not None), default=c2))
    else:
        info = find_table(ws, header_row)
        warnings += table_warnings(info)
        if not y:
            raise UserError("give 'y' (one or more value columns) or a 'range'")
        hdr, first, last = info.header_row, info.first_data, info.last_data
        xcol = info.col(x) if x else None
        ycols = [info.col(c) for c in y]
        used_last_col = info.last_col
    if last < first:
        raise UserError("no data rows to chart")
    if type == "pie" and len(ycols) != 1:
        raise UserError("a pie chart takes exactly one value column")
    if type == "scatter" and xcol is None:
        raise UserError("a scatter chart needs an 'x' column of numbers")
    reader = ValueReader(book)
    for c in ycols:
        vals = [reader.get(ws, r, c) for r in _range(first, last + 1)]
        nums = sum(1 for v in vals if isinstance(v, (int, float)) and not isinstance(v, bool))
        if nums == 0:
            raise UserError(f"column {get_column_letter(c)} has no numbers to plot")
        if nums < sum(1 for v in vals if v is not None):
            warnings.append(f"column {get_column_letter(c)} has non-numeric values; they plot as gaps")
    if type == "bar":
        ch = BarChart()
        ch.type = "col"
    elif type == "line":
        ch = LineChart()
    elif type == "pie":
        ch = PieChart()
    else:
        ch = ScatterChart()
        ch.style = 13
    ch.width, ch.height = width_cm, height_cm
    if type == "scatter":
        xref = Reference(ws, min_col=xcol, min_row=first, max_row=last)
        for c in ycols:
            s = Series(Reference(ws, min_col=c, min_row=hdr, max_row=last), xref, title_from_data=True)
            s.marker.symbol = "circle"
            s.graphicalProperties.line.noFill = True
            ch.series.append(s)
    else:
        for c in ycols:
            ch.add_data(Reference(ws, min_col=c, min_row=hdr, max_row=last), titles_from_data=True)
        if xcol is not None:
            ch.set_categories(Reference(ws, min_col=xcol, min_row=first, max_row=last))
    if title:
        ch.title = title
    if type == "pie":
        if x_title or y_title:
            warnings.append("pie charts have no axes; axis titles ignored")
    else:
        if x_title:
            ch.x_axis.title = x_title
        if y_title:
            ch.y_axis.title = y_title
        # openpyxl >= 3.1 hides axes unless told otherwise.
        ch.x_axis.delete = False
        ch.y_axis.delete = False
    if anchor:
        r, c = parse_cell(anchor)
        grid = _Grid(ws)
        x0, y0 = grid.x_of(c - 1), grid.y_of(r - 1)
        cand = (x0, y0, x0 + width_cm * EMU_PER_CM / EMU_PER_PX, y0 + height_cm * EMU_PER_CM / EMU_PER_PX)
        if any(o and _overlaps(cand, o) for o in (_rect(grid, x) for x in list(ws._charts) + list(ws._images))):
            warnings.append(f"chart at {anchor} overlaps an existing chart or image")
        pos, existing = anchor.upper(), len(ws._charts)
    else:
        pos, existing = _place(ws, width_cm, height_cm, used_last_col + 2, hdr)
    ws.add_chart(ch, pos)
    details = {"sheet": ws.title, "type": type, "anchor": pos, "series": len(ycols),
               "data_rows": last - first + 1, "existing_charts_and_images": existing,
               "size_cm": [width_cm, height_cm]}
    return finish(book, out_path, details, warnings + reader.warnings())



# --------------------------------------------------------------------------
# conditional_format
# --------------------------------------------------------------------------
_HEX = re.compile(r"^#?([0-9A-Fa-f]{6})$")
_CELLIS = {">": "greaterThan", ">=": "greaterThanOrEqual", "<": "lessThan",
           "<=": "lessThanOrEqual", "=": "equal", "!=": "notEqual", "between": "between",
           "not_between": "notBetween"}


def _color(c, default):
    if c is None:
        return default
    m = _HEX.match(str(c))
    if not m:
        raise UserError(f"color {c!r} must be hex like 'FF0000'")
    return m.group(1).upper()


def _target(ws, rng, column, header_row):
    if rng:
        c1, r1, c2, r2 = parse_range(rng)
        return c1, r1, c2, r2, None
    if not column:
        raise UserError("give 'range' (e.g. 'C2:C50') or 'column' (header name)")
    info = find_table(ws, header_row)
    c = info.col(column)
    if info.n_rows == 0:
        raise UserError("the table has no data rows")
    return c, info.first_data, c, info.last_data, info


def _crit_literal(v):
    if isinstance(v, bool):
        return "TRUE" if v else "FALSE"
    if isinstance(v, (int, float)):
        return fx.general_text(v)
    s = str(v)
    if s.startswith("="):
        return s[1:]
    x = fx._text_to_number(s)
    if x is not None and not re.match(r"\d{4}-", s.strip()):
        return fx.general_text(x)
    return '"' + s.replace('"', '""') + '"'


@operation("conditional_format", {
    "type": "object", "required": ["path", "rule"],
    "properties": {"path": S_PATH, "sheet": S_SHEET, "range": {"type": "string"},
                   "column": {"type": "string"},
                   "rule": {"enum": ["color_scale", "data_bar", "cell_is", "outliers"]},
                   "operator": {"enum": list(_CELLIS)}, "values": {"type": "array", "maxItems": 2},
                   "colors": {"type": "array", "items": {"type": "string"}, "maxItems": 3},
                   "z": {"type": "number", "exclusiveMinimum": 0}, "header_row": S_HDR,
                   "out_path": S_OUT, "force": S_FORCE},
})
def conditional_format(path, rule, sheet=None, range=None, column=None, operator=None,
                       values=None, colors=None, z=3.0, header_row=None, out_path=None,
                       force=False):
    """Add a color scale, data bar, cell-is rule, or z-score outlier highlight."""
    book = Book(path, force=force)
    ws = book.ws(sheet)
    c1, r1, c2, r2, info = _target(ws, range, column, header_row)
    sq = f"{get_column_letter(c1)}{r1}:{get_column_letter(c2)}{r2}"
    colors = colors or []
    warnings, details = [], {"sheet": ws.title, "range": sq, "rule": rule}
    if rule == "color_scale":
        if len(colors) == 2:
            cf = ColorScaleRule(start_type="min", start_color=_color(colors[0], "F8696B"),
                                end_type="max", end_color=_color(colors[1], "63BE7B"))
        else:
            cf = ColorScaleRule(start_type="min", start_color=_color(colors[0] if colors else None, "F8696B"),
                                mid_type="percentile", mid_value=50,
                                mid_color=_color(colors[1] if len(colors) > 1 else None, "FFEB84"),
                                end_type="max", end_color=_color(colors[2] if len(colors) > 2 else None, "63BE7B"))
    elif rule == "data_bar":
        cf = DataBarRule(start_type="min", end_type="max",
                         color=_color(colors[0] if colors else None, "638EC6"), showValue=True)
    elif rule == "cell_is":
        if not operator or not values:
            raise UserError("cell_is needs 'operator' and 'values'")
        need = 2 if operator in ("between", "not_between") else 1
        if len(values) != need:
            raise UserError(f"operator {operator!r} needs {need} value(s)")
        fill = _color(colors[0] if colors else None, "FFC7CE")
        font = _color(colors[1] if len(colors) > 1 else None, "9C0006")
        cf = CellIsRule(operator=_CELLIS[operator], formula=[_crit_literal(v) for v in values],
                        fill=PatternFill(start_color=fill, end_color=fill, fill_type="solid"),
                        font=Font(color=font))
    else:
        tl = f"{get_column_letter(c1)}{r1}"
        ab = f"${get_column_letter(c1)}${r1}:${get_column_letter(c2)}${r2}"
        formula = (f"AND(ISNUMBER({tl}),ABS({tl}-AVERAGE({ab}))>{fx.general_text(z)}*STDEV({ab}))")
        fill = _color(colors[0] if colors else None, "FFC7CE")
        font = _color(colors[1] if len(colors) > 1 else None, "9C0006")
        cf = FormulaRule(formula=[formula],
                         fill=PatternFill(start_color=fill, end_color=fill, fill_type="solid"),
                         font=Font(color=font))
        reader = ValueReader(book)
        cells = [(r, c, reader.get(ws, r, c)) for r in _range(r1, r2 + 1) for c in _range(c1, c2 + 1)]
        nums = [(r, c, v) for r, c, v in cells if isinstance(v, (int, float)) and not isinstance(v, bool)]
        flagged = []
        if len(nums) >= 2:
            mean = statistics.fmean(v for _, _, v in nums)
            sd = statistics.stdev(v for _, _, v in nums)
            if sd > 0:
                flagged = [{"cell": f"{get_column_letter(c)}{r}", "value": v,
                            "z": round((v - mean) / sd, 3)}
                           for r, c, v in nums if abs(v - mean) > z * sd]
            details.update(mean=fx.clean_number(mean), stdev=fx.clean_number(sd))
        else:
            warnings.append("fewer than two numbers in the range; nothing can be an outlier")
        details.update(z=z, formula="=" + formula, outliers=flagged[:100],
                       outlier_count=len(flagged))
        if len(nums) < 8:
            warnings.append(f"only {len(nums)} numbers; z-scores are unreliable on small samples "
                            f"(|z| cannot exceed {round((len(nums) - 1) / math.sqrt(len(nums)), 2)})"
                            if nums else "no numbers in range")
    ws.conditional_formatting.add(sq, cf)
    details["rules_on_sheet"] = sum(len(x.rules) for x in ws.conditional_formatting)
    return finish(book, out_path, details, warnings)


# --------------------------------------------------------------------------
# data_validation
# --------------------------------------------------------------------------
def _iso_date(s, what):
    try:
        return _dt.date.fromisoformat(str(s)[:10])
    except ValueError:
        raise UserError(f"{what} must be an ISO date like 2024-01-31, got {s!r}")


@operation("data_validation", {
    "type": "object", "required": ["path", "type"],
    "properties": {"path": S_PATH, "sheet": S_SHEET, "range": {"type": "string"},
                   "column": {"type": "string"},
                   "type": {"enum": ["list", "whole", "decimal", "date"]},
                   "values": {"type": "array"}, "source": {"type": "string"},
                   "min": {"type": ["number", "string"]}, "max": {"type": ["number", "string"]},
                   "allow_blank": {"type": "boolean"}, "error": {"type": "string"},
                   "prompt": {"type": "string"}, "header_row": S_HDR,
                   "out_path": S_OUT, "force": S_FORCE},
})
def data_validation(path, type, sheet=None, range=None, column=None, values=None, source=None,
                    min=None, max=None, allow_blank=True, error=None, prompt=None,
                    header_row=None, out_path=None, force=False):
    """Restrict input: a list (inline values or a source range), whole/decimal numbers, dates.

    Existing cells that would fail the rule are reported.
    """
    book = Book(path, force=force)
    ws = book.ws(sheet)
    c1, r1, c2, r2, _ = _target(ws, range, column, header_row)
    sq = f"{get_column_letter(c1)}{r1}:{get_column_letter(c2)}{r2}"
    warnings = []
    kw = {"allow_blank": allow_blank, "showErrorMessage": True, "errorStyle": "stop"}
    if type == "list":
        if bool(values) == bool(source):
            raise UserError("list validation needs exactly one of 'values' or 'source'")
        if values:
            items = [fx.general_text(v) for v in values]
            if any("," in i or '"' in i for i in items):
                raise UserError("inline list items cannot contain commas or quotes; put them in "
                                "cells and pass 'source' (e.g. Lists!$A$1:$A$10)")
            f1 = '"' + ",".join(items) + '"'
            if len(f1) > 257:
                raise UserError("inline list is longer than Excel's 255-character limit; "
                                "use 'source' with a range")
            allowed = {i.casefold() for i in items}
        else:
            f1 = source[1:] if source.startswith("=") else source
            try:
                toks = fx.references(f1)
            except (fx.FormulaSyntaxError, fx.Unsupported):
                toks = []
            if len(toks) != 1:
                raise UserError(f"source {source!r} must be one range like Lists!$A$1:$A$10")
            reader = ValueReader(book)
            try:
                rng = reader.ev.eval(fx.parse(f1), ws.title)
                allowed = {fx.general_text(v).casefold() for v in rng.flat() if v is not None}
            except Exception:
                allowed = None
                warnings.append("could not read the source range to check existing values")
        dv = DataValidation(type="list", formula1=f1, **kw)

        def ok(v):
            return allowed is None or fx.general_text(v).casefold() in allowed
        details_rule = {"values": values, "source": source}
    else:
        if min is None and max is None:
            raise UserError(f"{type} validation needs 'min' and/or 'max'")
        if type == "date":
            lo = _iso_date(min, "min") if min is not None else None
            hi = _iso_date(max, "max") if max is not None else None
            fmt = lambda d: f"DATE({d.year},{d.month},{d.day})"  # noqa: E731
            from openpyxl.utils.datetime import to_excel
            lo_n = to_excel(lo) if lo else None
            hi_n = to_excel(hi) if hi else None
        else:
            try:
                lo = float(min) if min is not None else None
                hi = float(max) if max is not None else None
            except (TypeError, ValueError):
                raise UserError("min/max must be numbers")
            if type == "whole":
                lo = math.ceil(lo) if lo is not None else None
                hi = math.floor(hi) if hi is not None else None
            fmt = fx.general_text
            lo_n, hi_n = lo, hi
        if lo is not None and hi is not None:
            if lo_n > hi_n:
                raise UserError("min is greater than max")
            dv = DataValidation(type=type, operator="between", formula1=fmt(lo), formula2=fmt(hi), **kw)
        elif lo is not None:
            dv = DataValidation(type=type, operator="greaterThanOrEqual", formula1=fmt(lo), **kw)
        else:
            dv = DataValidation(type=type, operator="lessThanOrEqual", formula1=fmt(hi), **kw)
        reader = ValueReader(book)

        def ok(v):
            if not isinstance(v, (int, float)) or isinstance(v, bool):
                return False
            if type == "whole" and float(v) != int(v):
                return False
            return (lo_n is None or v >= lo_n) and (hi_n is None or v <= hi_n)
        details_rule = {"min": min, "max": max}
    if error:
        dv.error, dv.errorTitle = error, "Invalid entry"
    if prompt:
        dv.prompt, dv.showInputMessage = prompt, True
    overlap = [str(d.sqref) for d in ws.data_validations.dataValidation
               if any(cr.min_col <= c2 and cr.max_col >= c1 and cr.min_row <= r2 and cr.max_row >= r1
                      for cr in d.sqref.ranges)]
    if overlap:
        warnings.append(f"overlaps existing validation on {overlap}; Excel applies only one")
    ws.add_data_validation(dv)
    dv.add(sq)
    reader = ValueReader(book)
    bad = []
    for r in _range(r1, r2 + 1):
        for c in _range(c1, c2 + 1):
            v = reader.get(ws, r, c)
            if v is None or (v == "" and allow_blank):
                continue
            if not ok(v):
                bad.append({"cell": f"{get_column_letter(c)}{r}", "value": v})
    if bad:
        warnings.append(f"{len(bad)} existing cell(s) do not satisfy the new rule")
    details = {"sheet": ws.title, "range": sq, "type": type, **details_rule,
               "invalid_existing": bad[:50], "invalid_count": len(bad)}
    return finish(book, out_path, details, warnings)


# --------------------------------------------------------------------------
# format_range
# --------------------------------------------------------------------------
PRESETS = {"integer": "#,##0", "decimal": "#,##0.00", "currency": '"$"#,##0.00',
           "percent": "0.0%", "date": "yyyy-mm-dd", "datetime": "yyyy-mm-dd hh:mm",
           "text": "@", "general": "General"}


@operation("format_range", {
    "type": "object", "required": ["path"],
    "properties": {"path": S_PATH, "sheet": S_SHEET, "range": {"type": "string"},
                   "columns": S_COLS, "number_format": {"type": "string"},
                   "bold_header": {"type": "boolean"}, "freeze": {"type": "string"},
                   "autofit": {"type": "boolean"}, "header_row": S_HDR,
                   "out_path": S_OUT, "force": S_FORCE},
})
def format_range(path, sheet=None, range=None, columns=None, number_format=None,
                 bold_header=False, freeze=None, autofit=False, header_row=None, out_path=None,
                 force=False):
    """Number format (presets: integer, decimal, currency, percent, date, datetime, text),
    bold header, freeze panes ('header', a cell like 'B2', or 'none'), autofit widths.
    """
    if not any([number_format, bold_header, freeze, autofit]):
        raise UserError("nothing to do: give number_format, bold_header, freeze or autofit")
    book = Book(path, force=force)
    ws = book.ws(sheet)
    info = None
    if not range or bold_header or freeze == "header" or columns:
        try:
            info = find_table(ws, header_row)
        except UserError:
            if columns or bold_header or freeze == "header" or not range:
                raise
    blocks = []
    if range:
        blocks.append(parse_range(range))
    if columns:
        for col in columns:
            c = info.col(col)
            blocks.append((c, info.first_data, c, max(info.last_data, info.first_data)))
    if not blocks and info is not None:
        blocks.append((info.first_col, info.first_data, info.last_col, max(info.last_data, info.first_data)))
    details = {"sheet": ws.title, "ranges": [f"{get_column_letter(a)}{b}:{get_column_letter(c)}{d}"
                                             for a, b, c, d in blocks]}
    warnings = []
    if number_format:
        fmt = PRESETS.get(number_format.lower(), number_format)
        total = sum((c2 - c1 + 1) * (r2 - r1 + 1) for c1, r1, c2, r2 in blocks)
        if total > 500_000:
            raise UserError("range too large to format cell by cell (> 500k cells)")
        n = 0
        for c1, r1, c2, r2 in blocks:
            for r in _range(r1, r2 + 1):
                for c in _range(c1, c2 + 1):
                    ws.cell(row=r, column=c).number_format = fmt
                    n += 1
        details["number_format"] = fmt
        details["cells_formatted"] = n
    if bold_header:
        hr = info.header_row
        for c in _range(info.first_col, info.last_col + 1):
            cell = ws._cells.get((hr, c))
            if cell is not None and cell.value is not None:
                cell.font = bolded(cell.font)
        details["bold_header_row"] = hr
    if freeze:
        if freeze.lower() == "none":
            ws.freeze_panes = None
        elif freeze.lower() == "header":
            if info is None:
                raise UserError("freeze='header' needs a header row")
            ws.freeze_panes = f"A{info.header_row + 1}"
        else:
            r, c = parse_cell(freeze)
            ws.freeze_panes = f"{get_column_letter(c)}{r}"
        details["freeze_panes"] = ws.freeze_panes
    if autofit:
        cols = set()
        rows = None
        for c1, r1, c2, r2 in blocks:
            cols.update(_range(c1, c2 + 1))
        if info is not None and not range:
            rows = set(_range(info.header_row, max(info.last_data, info.total_row or 0) + 1))
        details["column_widths"] = autofit_columns(ws, cols, rows)
    return finish(book, out_path, details, warnings)
