# SPDX-License-Identifier: Apache-2.0
"""The pure-Python evaluator and verify_formulas."""
from __future__ import annotations

import datetime as dt

import openpyxl
import pytest

import xlsxops as X
from xlsxops import formula as fx


def evaluate(formula, setup=None):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "S"
    other = wb.create_sheet("Other Sheet")
    rows = [["Name", "Qty", "Price", "When"],
            ["apple", 3, 1.5, dt.date(2024, 1, 2)],
            ["Banana", 5, 0.25, dt.date(2024, 2, 3)],
            ["cherry", None, 4, dt.date(2024, 3, 4)],
            ["apple", 2, "n/a", None]]
    for r in rows:
        ws.append(r)
    other["A1"] = 10
    other["A2"] = 20
    if setup:
        setup(wb)
    ws["Z1"] = formula
    report, values, _ = fx.verify(wb)
    key = ("S", 1, 26)
    if key in values:
        return values[key]
    return report


CASES = [
    ("=1+2*3-4/2", 5), ("=-2^2", 4), ("=2^3^2", 64), ("=10%", 0.1), ('="a"&"b"&1', "ab1"),
    ("=1/3*3", 1), ("=0.1+0.2", 0.3), ('=1="1"', False), ('="ABC"="abc"', True), ("=2<\"a\"", True),
    ("=SUM(B2:B5)", 10), ("=SUM(B2:C5)", 10 + 5.75), ("=SUM(1,\"2\",TRUE)", 4),
    ("=AVERAGE(B2:B5)", 10 / 3), ("=COUNT(A2:D5)", 9), ("=COUNTA(A2:D5)", 14),
    ('=COUNTIF(A2:A5,"apple")', 2), ('=COUNTIF(B2:B5,">2")', 2), ('=COUNTIF(A2:A5,"b*")', 1),
    ('=COUNTIF(B2:B5,"")', 1), ('=COUNTIF(A2:A5,"<>apple")', 2),
    ('=COUNTIFS(A2:A5,"apple",B2:B5,">2")', 1), ('=SUMIF(A2:A5,"apple",B2:B5)', 5),
    ('=SUMIF(B2:B5,">=3")', 8), ('=SUMIFS(B2:B5,A2:A5,"?pple",B2:B5,"<3")', 2),
    ('=AVERAGEIF(A2:A5,"apple",B2:B5)', 2.5), ("=MIN(B2:B5)", 2), ("=MAX(B2:C5)", 5),
    ('=IF(B2>2,"big","small")', "big"), ("=IF(FALSE,1)", False), ("=IFERROR(1/0,-1)", -1),
    ('=IFERROR(VLOOKUP("zzz",A2:B5,2,FALSE),"none")', "none"),
    ("=ROUND(2.5,0)", 3), ("=ROUND(-2.5,0)", -3), ("=ROUND(1234.567,-2)", 1200), ("=ROUND(1.005,2)", 1.01),
    ('=VLOOKUP("banana",A2:C5,3,FALSE)', 0.25), ("=VLOOKUP(4,B2:C3,2)", 1.5),
    ('=XLOOKUP("cherry",A2:A5,C2:C5)', 4), ('=XLOOKUP("kiwi",A2:A5,C2:C5,"?")', "?"),
    ('=XLOOKUP("apple",A2:A5,B2:B5,,0,-1)', 2), ('=_xlfn.XLOOKUP(4,B2:B3,A2:A3,,1)', "Banana"),
    ("=INDEX(A2:C5,2,3)", 0.25), ("=INDEX(B2:B5,3)", 0), ('=MATCH("cherry",A2:A5,0)', 3),
    ("=MATCH(4,B2:B3,1)", 1), ('=INDEX(C2:C5,MATCH("banana",A2:A5,0))', 0.25),
    ('=TEXT(1234.5,"#,##0.00")', "1,234.50"), ('=TEXT(0.256,"0.0%")', "25.6%"),
    ('=TEXT(D2,"yyyy-mm-dd")', "2024-01-02"), ('=TEXT(D3,"mmm d, yyyy")', "Feb 3, 2024"),
    ('=TEXT(-3,"$#,##0.00")', "-$3.00"), ('=TEXT(7,"000")', "007"),
    ("='Other Sheet'!A1+'Other Sheet'!A2", 30), ("=SUM('Other Sheet'!A:A)", 30),
    ("=D3-D2", 32), ("=AND(B2>1,OR(B3<0,TRUE))", True), ('=LEN("abc")&UPPER("x")', "3X"),
]


@pytest.mark.parametrize("formula,expected", CASES, ids=[c[0] for c in CASES])
def test_evaluator(formula, expected):
    got = evaluate(formula)
    if isinstance(expected, float):
        assert got == pytest.approx(expected)
    else:
        assert got == expected and type(got) is type(expected) or got == expected


ERRORS = [
    ("=1/0", "#DIV/0!"), ('="a"+1', "#VALUE!"), ("=SUM(C2:C5)+C5", "#VALUE!"),
    ('=VLOOKUP("zzz",A2:B5,2,FALSE)', "#N/A"), ("=VLOOKUP(1,B2:C5,5)", "#REF!"),
    ("=AVERAGE(A2:A5)", "#DIV/0!"), ("=INDEX(A2:B3,5,1)", "#REF!"), ("=NoSuchName+1", "#NAME?"),
]


@pytest.mark.parametrize("formula,code", ERRORS, ids=[e[0] for e in ERRORS])
def test_evaluator_errors(formula, code):
    assert evaluate(formula) == fx.XLError(code)


def test_defined_names():
    def setup(wb):
        from openpyxl.workbook.defined_name import DefinedName
        wb.defined_names["Rate"] = DefinedName("Rate", attr_text="'Other Sheet'!$A$2")
    assert evaluate("=Rate*2", setup) == 40


def test_future_prefix_roundtrip():
    assert fx.add_future_prefixes("=XLOOKUP(A1,B:B,C:C)+IFNA(1,2)") == \
        "=_xlfn.XLOOKUP(A1,B:B,C:C)+_xlfn.IFNA(1,2)"
    assert fx.add_future_prefixes('=SUM(A1)&"XLOOKUP("') == '=SUM(A1)&"XLOOKUP("'


def test_verify_formulas_broken_reference(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Data"
    ws["A1"], ws["A2"], ws["A3"] = 1, 2, 3
    ws["B1"] = "=SUM(A1:A3)"            # fine
    ws["B2"] = "=SUM(Missing!A1:A3)"    # sheet does not exist -> #REF!
    ws["B3"] = "=A1+#REF!"              # deleted reference -> #REF!
    ws["B4"] = "=B2*2"                  # inherits the error (not the origin)
    ws["B5"] = '=A1&"x"+1'              # type error
    ws["C1"] = "=C2+1"                  # circular
    ws["C2"] = "=C1+1"
    ws["C3"] = "=C1*10"                 # depends on the cycle
    ws["D1"] = "=NPV(0.1,A1:A3)"        # not supported by the evaluator
    ws["D2"] = "=D1+1"                  # depends on an unsupported formula
    path = tmp_path / "broken.xlsx"
    wb.save(path)
    r = X.verify_formulas(str(path))
    assert r["ok"] and r["changed"] is False
    d = r["details"]
    errs = {e["cell"]: e for e in d["errors"]}
    assert errs["Data!B2"]["error"] == "#REF!" and errs["Data!B2"]["kind"] == "ref"
    assert errs["Data!B3"]["error"] == "#REF!" and errs["Data!B3"]["origin"] is True
    assert errs["Data!B4"]["error"] == "#REF!" and errs["Data!B4"]["origin"] is False
    assert errs["Data!B5"]["kind"] == "type"
    assert errs["Data!C3"]["kind"] == "circular_dependency"
    assert {c["cell"] for c in d["circular"]} == {"Data!C1", "Data!C2"}
    unsup = {u["cell"]: u["reason"] for u in d["unsupported"]}
    assert "NPV" in unsup["Data!D1"] and "Data!D2" in unsup
    assert "Data!B1" not in errs and d["evaluated"] >= 5
    assert any("errors" in w for w in r["warnings"])
    # write_cache stores computed values so viewers without recalculation show numbers
    r = X.verify_formulas(str(path), write_cache=True)
    assert r["changed"] is True
    cached = openpyxl.load_workbook(path, data_only=True)["Data"]
    assert cached["B1"].value == 6 and cached["B2"].value == "#REF!"


def test_unsupported_formula_keeps_previous_cached_value(tmp_path):
    path = tmp_path / "npv.xlsx"
    wb = openpyxl.Workbook()
    wb.active["A1"] = 100
    wb.active["B1"] = "=NPV(0.1,A1)"
    wb.save(path)
    # Simulate Excel having calculated it: inject a cached value, then edit another cell.
    from xlsxops._core import _inject_cached_values
    tmp = tmp_path / "npv2.xlsx"
    _inject_cached_values(path, tmp, {("Sheet", 1, 2): 90.9090909})
    r = X.write_cells(str(tmp), cells={"C1": 1})
    assert r["ok"] and r["details"]["formula_check"]["kept_previous_cached_value"] == 1
    assert openpyxl.load_workbook(tmp, data_only=True).active["B1"].value == pytest.approx(90.9090909)
