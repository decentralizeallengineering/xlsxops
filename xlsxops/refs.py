# SPDX-License-Identifier: Apache-2.0
"""Rewrite cell references across a workbook after rows move.

Used when an operation removes rows from a table (dedupe): formulas, defined
names, chart series, conditional formats, data validations, Excel tables and
autofilters that ended at the old last row are shrunk to the new last row.
"""
from __future__ import annotations

from copy import copy

from openpyxl.formatting.formatting import ConditionalFormattingList
from openpyxl.utils.cell import get_column_letter, range_boundaries
from openpyxl.worksheet.cell_range import MultiCellRange

from . import formula as fx


def _rewrite(text, fn, host):
    if not text:
        return text
    try:
        return fx.rewrite_refs(text, lambda t: fn(t, host))
    except (fx.FormulaSyntaxError, fx.Unsupported):
        return text


def _rewrite_sqref(sqref: str, fn, host) -> str:
    parts = []
    for part in str(sqref).split():
        parts.append(_rewrite(part, fn, host))
    return " ".join(parts)


def _chart_refs(chart):
    for s in getattr(chart, "series", []) or []:
        for attr in ("val", "cat", "xVal", "yVal", "bubbleSize"):
            obj = getattr(s, attr, None)
            if obj is None:
                continue
            for sub in ("numRef", "strRef", "multiLvlStrRef"):
                ref = getattr(obj, sub, None)
                if ref is not None and getattr(ref, "f", None):
                    yield ref
        tx = getattr(s, "tx", None)
        if tx is not None and tx.strRef is not None and tx.strRef.f:
            yield tx.strRef


def apply(wb, fn) -> int:
    """Apply ``fn(RefTok, host_sheet_title) -> RefTok|None`` everywhere. Returns edits."""
    edits = 0
    for ws in wb.worksheets:
        host = ws.title
        for cell in list(ws._cells.values()):
            if cell.data_type == "f" and isinstance(cell.value, str):
                new = _rewrite(cell.value, fn, host)
                if new != cell.value:
                    cell._value = new
                    edits += 1
        for chart in ws._charts:
            for ref in _chart_refs(chart):
                new = _rewrite(ref.f, fn, host)
                if new != ref.f:
                    ref.f = new
                    edits += 1
        cfl = ws.conditional_formatting
        rebuilt = ConditionalFormattingList()
        touched = False
        for cf in cfl:
            sq = _rewrite_sqref(str(cf.sqref), fn, host)
            touched |= sq != str(cf.sqref)
            for rule in cf.rules:
                if rule.formula:
                    nf = [_rewrite(f, fn, host) for f in rule.formula]
                    touched |= nf != list(rule.formula)
                    rule.formula = nf
                rebuilt.add(sq, rule)
        if touched:
            ws.conditional_formatting = rebuilt
            edits += 1
        for dv in ws.data_validations.dataValidation:
            sq = _rewrite_sqref(str(dv.sqref), fn, host)
            if sq != str(dv.sqref):
                dv.sqref = MultiCellRange(sq)
                edits += 1
            for attr in ("formula1", "formula2"):
                f = getattr(dv, attr)
                if f:
                    nf = _rewrite(f, fn, host)
                    if nf != f:
                        setattr(dv, attr, nf)
                        edits += 1
        for t in ws.tables.values():
            nr = _rewrite(t.ref, fn, host)
            if nr != t.ref:
                t.ref = nr
                if t.autoFilter is not None and t.autoFilter.ref:
                    t.autoFilter.ref = nr
                edits += 1
        if ws.auto_filter.ref:
            nr = _rewrite(ws.auto_filter.ref, fn, host)
            if nr != ws.auto_filter.ref:
                ws.auto_filter.ref = nr
                edits += 1
    for name, dn in list(wb.defined_names.items()):
        if dn.attr_text:
            nt = _rewrite(dn.attr_text, fn, None)
            if nt != dn.attr_text:
                dn.attr_text = nt
                edits += 1
    for ws in wb.worksheets:
        for name, dn in list(getattr(ws, "defined_names", {}).items()):
            if dn.attr_text:
                nt = _rewrite(dn.attr_text, fn, ws.title)
                if nt != dn.attr_text:
                    dn.attr_text = nt
                    edits += 1
    return edits


def shrink_block(wb, ws, first, old_last, new_last, c1, c2, old_total=None, new_total=None):
    """References that end at ``old_last`` now end at ``new_last``; the total row moves.

    Returns (edits, single_cell_refs_into_block).
    """
    target = ws.title.casefold()
    stray = [0]

    def fn(tok, host):
        if tok.external or tok.kind not in ("cell", "area"):
            return None
        sheet = tok.sheet if tok.sheet is not None else host
        if sheet is None or sheet.casefold() != target:
            return None
        if tok.c2 < c1 or tok.c1 > c2:
            return None
        t = copy(tok)
        changed = False
        if tok.kind == "area" and tok.r2 == old_last and tok.r1 <= first and tok.r1 < tok.r2:
            t.r2 = new_last
            changed = True
        if old_total is not None:
            if t.r1 == old_total:
                t.r1, changed = new_total, True
            if t.r2 == old_total:
                t.r2, changed = new_total, True
        if (not changed and tok.kind == "cell" and first <= tok.r1 <= old_last
                and (host is None or host.casefold() != target)):
            stray[0] += 1
        return t if changed else None

    edits = apply(wb, fn)
    return edits, stray[0]


def range_text(c1, r1, c2, r2):
    return f"{get_column_letter(c1)}{r1}:{get_column_letter(c2)}{r2}"


__all__ = ["apply", "shrink_block", "range_text", "range_boundaries"]
