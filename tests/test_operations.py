# SPDX-License-Identifier: Apache-2.0
"""One test per operation (plus a few edge cases that matter for agents)."""
from __future__ import annotations

import datetime as dt
import json

import openpyxl
import pytest

import xlsxops as X


def ok(report):
    json.dumps(report)  # always JSON-serializable
    assert set(report) == {"ok", "changed", "warnings", "details"}
    assert report["ok"], report["details"]
    return report


def cached(path, sheet, ref):
    return openpyxl.load_workbook(path, data_only=True)[sheet][ref].value


def load(path):
    return openpyxl.load_workbook(path)


def test_read_summary(sales):
    r = ok(X.read_summary(sales))
    assert r["changed"] is False
    s = {x["name"]: x for x in r["details"]["sheets"]}
    assert s["Sales"]["headers"][:4] == ["Region", "Rep", "Product", "Qty"]
    assert s["Sales"]["row_count"] == 10
    assert s["Sales"]["used_range"] == "A1:G11"
    assert s["Sales"]["formulas_count"] == 10
    assert s["Sales"]["charts"] == 1
    assert len(s["Sales"]["sample_rows"]) == 5
    # formula cells show formula + computed value even with no cached value in the file
    assert s["Sales"]["sample_rows"][0][6] == {"formula": "=D2*E2", "value": 25}
    assert {"name": "TaxRate", "value": "Notes!$B$1"} in r["details"]["defined_names"]


def test_write_cells(sales):
    r = ok(X.write_cells(sales, sheet="Notes", cells={
        "D1": 5, "D2": "=D1*2", "D3": "2024-05-01", "D4": "'=not a formula", "D5": "=XLOOKUP(1,{1},{7})"}))
    assert r["details"]["formula_values"]["D2"] == 10
    wb = load(sales)
    n = wb["Notes"]
    assert n["D3"].value == dt.datetime(2024, 5, 1) and n["D3"].number_format == "yyyy-mm-dd"
    assert n["D4"].value == "=not a formula" and n["D4"].data_type == "s"
    assert n["D5"].value.startswith("=_xlfn.XLOOKUP(")  # stored the way Excel expects
    assert cached(sales, "Notes", "D2") == 10
    ok(X.write_cells(sales, sheet="Notes", start="F1", values=[[1, 2], [3, "=F1+G1"]]))
    assert cached(sales, "Notes", "G2") == 3


def test_write_cells_rejects_bad_formula_without_writing(sales):
    before = load(sales)["Notes"]["D1"].value
    r = X.write_cells(sales, sheet="Notes", cells={"D1": "=SUM(A1:A3"})
    assert r["ok"] is False and r["details"]["error_type"] == "formula_syntax"
    assert load(sales)["Notes"]["D1"].value == before


def test_add_sheet(sales, tmp_path):
    ok(X.add_sheet(sales, name="Budget", index=1, headers=["Item", "Amount"]))
    wb = load(sales)
    assert wb.sheetnames[1] == "Budget"
    assert wb["Budget"]["A1"].font.b and wb["Budget"].freeze_panes == "A2"
    assert X.add_sheet(sales, name="budget")["details"]["error_type"] == "sheet_exists"
    assert X.add_sheet(sales, name="bad/name")["ok"] is False
    fresh = str(tmp_path / "new.xlsx")
    r = ok(X.add_sheet(fresh, name="First"))
    assert r["details"]["created_workbook"] and load(fresh).sheetnames == ["First"]


def test_add_column(sales):
    r = ok(X.add_column(sales, sheet="Sales", header="Size", fill='=IF(D2>5,"big","small")'))
    assert r["details"]["column"] == "H"
    ws = load(sales)["Sales"]
    assert ws["H1"].value == "Size" and ws["H1"].font.b  # header style copied from neighbour
    assert ws["H1"].fill.fgColor.rgb == "FF1F4E78"
    assert ws["H3"].value == '=IF(D3>5,"big","small")'
    assert cached(sales, "Sales", "H2") == "big" and cached(sales, "Sales", "H3") == "small"
    ok(X.add_column(sales, sheet="Sales", header="Note", values=["a", "b"]))
    assert load(sales)["Sales"]["I3"].value == "b"
    r = X.add_column(sales, sheet="Sales", header="X", column="C")
    assert r["ok"] is False and "inside the table" in r["details"]["error"]


def test_calculated_column(sales):
    r = ok(X.calculated_column(sales, sheet="Sales", header="With tax",
                               formula="=[@Total]*(1+TaxRate)", number_format="0.00"))
    assert r["details"]["first_formula"] == "=G2*(1+TaxRate)"
    assert r["details"]["sample_values"][0] == pytest.approx(25 * 1.0825)
    ws = load(sales)["Sales"]
    assert ws["H11"].value == "=G11*(1+TaxRate)" and ws["H11"].number_format == "0.00"
    # relative-reference form, re-run on an existing header refills that column
    r = ok(X.calculated_column(sales, sheet="Sales", header="With tax", formula="=D2+1"))
    assert load(sales)["Sales"]["H5"].value == "=D5+1"
    assert cached(sales, "Sales", "H5") == 2


def test_set_formula(sales):
    r = ok(X.set_formula(sales, sheet="Notes", cell="D1:D3", formula="=ROUND(B1*100,1)"))
    assert [x["formula"] for x in r["details"]["results"]] == [
        "=ROUND(B1*100,1)", "=ROUND(B2*100,1)", "=ROUND(B3*100,1)"]
    assert r["details"]["results"][0]["value"] == 8.3
    assert r["details"]["results"][1]["value"] == 8500  # B2 is SUM of quantities = 85
    assert cached(sales, "Notes", "D1") == 8.3


def test_total_row(sales):
    r = ok(X.total_row(sales, sheet="Sales", columns=["Qty", "Total"], funcs={"Price": "AVERAGE"}))
    assert r["details"]["row"] == 12
    assert r["details"]["values"] == {"Qty": 85, "Total": 345.75, "Price": 5.025}
    ws = load(sales)["Sales"]
    assert ws["A12"].value == "Total" and ws["D12"].font.b
    assert ws["G12"].number_format == '"$"#,##0.00'
    # running again replaces rather than stacking another total row
    r = ok(X.total_row(sales, sheet="Sales", columns=["Qty"], func="COUNT"))
    assert r["details"]["replaced_existing"] and r["details"]["row"] == 12
    ws = load(sales)["Sales"]
    assert ws["D12"].value == "=COUNT(D2:D11)" and ws["G12"].value is None and ws["A13"].value is None


def test_sort(sales):
    ok(X.total_row(sales, sheet="Sales", columns=["Qty"]))
    r = ok(X.sort(sales, sheet="Sales", keys=[{"column": "Region"}, {"column": "Qty", "descending": True}]))
    ws = load(sales)["Sales"]
    regions = [ws.cell(row=i, column=1).value for i in range(2, 12)]
    assert regions == sorted(regions, key=str.casefold)
    east = [(ws.cell(row=i, column=4).value) for i in range(2, 12) if ws.cell(row=i, column=1).value == "East"]
    assert east == sorted(east, reverse=True)
    # formulas moved with their rows and still point at their own row
    for i in range(2, 12):
        assert ws.cell(row=i, column=7).value == f"=D{i}*E{i}"
    assert ws["D12"].value == "=SUM(D2:D11)"  # total row was not sorted into the data
    assert ws["E2"].number_format == '"$"#,##0.00'  # styles travel with cells
    assert r["details"]["preview"][0]["Region"] == "East"


def test_filter_rows(sales):
    r = ok(X.filter_rows(sales, sheet="Sales", new_sheet="Big West",
                         conditions=[{"column": "Qty", "op": ">=", "value": 5},
                                     {"column": "Region", "op": "=", "value": "west"}]))
    assert r["details"]["matched"] == 2
    ws = load(sales)["Big West"]
    assert [c.value for c in ws[1]][:3] == ["Region", "Rep", "Product"]
    assert ws["G2"].value == 25 and ws["G2"].number_format == '"$"#,##0.00'  # formula -> value
    assert ws["F2"].is_date
    r = ok(X.filter_rows(sales, sheet="Sales", new_sheet="Feb", match="all",
                         conditions=[{"column": "Date", "op": ">=", "value": "2024-02-01"},
                                     {"column": "Date", "op": "<", "value": "2024-03-01"}]))
    assert r["details"]["matched"] == 3


def test_dedupe(sales):
    ok(X.total_row(sales, sheet="Sales", columns=["Qty"]))
    r = ok(X.dedupe(sales, sheet="Sales", columns=["Region", "Rep", "Product", "Qty"]))
    d = r["details"]
    assert d["removed_count"] == 2 and d["rows_after"] == 8
    assert [x["row"] for x in d["removed"]] == [6, 9]
    assert d["removed"][1]["duplicate_of_row"] == 5  # 'north'/'widget' matched case-insensitively
    wb = load(sales)
    ws = wb["Sales"]
    assert ws["A10"].value == "Total" and ws["D10"].value == "=SUM(D2:D9)"
    assert ws["A11"].value is None
    assert wb["Notes"]["B2"].value == "=SUM(Sales!D2:D9)"  # references shrink with the table
    ser = ws._charts[0].series[0]
    assert ser.val.numRef.f.endswith("$D$2:$D$9")
    assert str(ws.data_validations.dataValidation[0].sqref) == "A2:A9"
    assert cached(sales, "Sales", "D10") == 81
    r = ok(X.dedupe(sales, sheet="Sales"))
    assert r["changed"] is False and r["details"]["removed_count"] == 0


def test_split_column(sales):
    ok(X.write_cells(sales, sheet="Products", cells={"D1": "Code", "D2": "HW-001-x", "D3": "EL-002-y",
                                                     "D4": "EL-010-z", "D5": "HW-7"}))
    r = ok(X.split_column(sales, sheet="Products", column="Code", delimiter="-",
                          headers=["Dept", "Num", "Suffix"]))
    assert r["details"]["new_columns"] == {"Dept": "E", "Num": "F", "Suffix": "G"}
    ws = load(sales)["Products"]
    assert ws["F2"].value == "001"  # leading zeros kept as text
    assert ws["F5"].value == 7 and ws["G5"].value is None
    r = ok(X.split_column(sales, sheet="Products", column="Code", regex=r"^([A-Z]+)-(\d+)",
                          headers=["Dept2", "N2"]))
    assert load(sales)["Products"]["I3"].value == "002"


@pytest.mark.parametrize("style", ["xlookup", "vlookup", "index_match"])
def test_lookup_formulas(sales, style):
    r = ok(X.lookup(sales, sheet="Sales", key="Product", source_sheet="Products",
                    columns=["Category", "Cost"], style=style, not_found="none"))
    d = r["details"]
    assert d["matched"] == 9 and d["unmatched_keys"] == ["Doohickey"]
    ws = load(sales)["Sales"]
    assert ws["H1"].value == "Category" and ws["H2"].data_type == "f"
    assert cached(sales, "Sales", "H2") == "Hardware"
    assert cached(sales, "Sales", "I3") == 6
    assert cached(sales, "Sales", "H7") == "none"  # Doohickey
    assert cached(sales, "Sales", "H9") == "Hardware"  # 'widget' matches 'Widget' like Excel


def test_lookup_values(sales):
    ok(X.write_cells(sales, sheet="Products", cells={"D1": "Code", "D2": 1001}))
    r = ok(X.lookup(sales, sheet="Sales", key="Product", source_sheet="Products",
                    columns=["Category"], mode="value", not_found="#N/A"))
    ws = load(sales)["Sales"]
    assert ws["H2"].value == "Hardware" and ws["H2"].data_type != "f"
    assert ws["H7"].value == "#N/A" and ws["H7"].data_type == "e"
    assert r["details"]["unmatched"] == 1


def test_pivot_summary(sales):
    r = ok(X.pivot_summary(sales, sheet="Sales", group_by=["Region"], new_sheet="By region",
                           values=[{"column": "Qty", "agg": "sum"}, {"column": "Price", "agg": "avg"},
                                   {"column": "Rep", "agg": "count"}], formulas=True))
    d = r["details"]
    assert d["groups"] == 4  # 'north' and 'North' group together, like SUMIFS
    assert d["headers"] == ["Region", "Sum of Qty", "Average of Price", "Count of Rep"]
    assert d["formula_version_matches_values"] is True
    wb = load(sales)
    v = wb["By region"]
    assert [v.cell(row=i, column=1).value for i in range(2, 6)] == ["East", "North", "South", "West"]
    assert v["B2"].value == 46 and v["B6"].value == 85 and v["A6"].value == "Total"
    f = wb[d["formula_sheet"]]
    assert f["B2"].value.startswith("=SUMIFS(Sales!$D$2:$D$11")
    assert cached(sales, d["formula_sheet"], "B2") == 46
    assert cached(sales, d["formula_sheet"], "D6") == 10


def test_chart(sales):
    r = ok(X.chart(sales, sheet="Sales", type="line", x="Rep", y=["Qty", "Total"], title="Qty",
                   x_title="Rep", y_title="Units"))
    first = r["details"]["anchor"]
    r = ok(X.chart(sales, sheet="Sales", type="scatter", x="Qty", y=["Total"]))
    second = r["details"]["anchor"]
    ok(X.chart(sales, sheet="Sales", type="pie", x="Rep", y=["Qty"]))
    ok(X.write_cells(sales, sheet="Products", start="E1", values=[["Month", "Units"], ["Jan", 3], ["Feb", 5]]))
    r = ok(X.chart(sales, sheet="Products", type="bar", range="E1:F3", title="Units"))
    assert r["details"]["anchor"] == "H1"  # right of everything on the sheet
    bad = X.chart(sales, sheet="Products", type="bar", range="A1:B5")
    assert bad["ok"] is False and "no numbers" in bad["details"]["error"]
    wb = load(sales)
    ws = wb["Sales"]
    assert len(ws._charts) == 4 and len(wb["Products"]._charts) == 1
    assert first != second
    # no two charts on the sheet overlap
    from xlsxops.ops_format import _Grid, _overlaps, _rect
    g = _Grid(ws)
    rects = [_rect(g, c) for c in ws._charts]
    for i in range(len(rects)):
        for j in range(i + 1, len(rects)):
            assert not _overlaps(rects[i], rects[j])
    line = ws._charts[1]
    assert line.x_axis.title is not None and line.x_axis.delete is False
    bad = X.chart(sales, sheet="Sales", type="pie", x="Rep", y=["Qty", "Total"])
    assert bad["ok"] is False


def test_conditional_format(sales):
    ok(X.conditional_format(sales, sheet="Sales", column="Total", rule="color_scale"))
    ok(X.conditional_format(sales, sheet="Sales", range="E2:E11", rule="cell_is", operator=">",
                            values=[5], colors=["FFEB9C", "9C5700"]))
    r = ok(X.conditional_format(sales, sheet="Sales", column="Qty", rule="outliers", z=2))
    assert [o["cell"] for o in r["details"]["outliers"]] == ["D10"]
    ws = load(sales)["Sales"]
    rules = {str(cf.sqref): [x.type for x in cf.rules] for cf in ws.conditional_formatting}
    assert rules["G2:G11"] == ["colorScale"] and rules["E2:E11"] == ["cellIs"]
    assert "expression" in rules["D2:D11"] and "dataBar" in rules["D2:D11"]
    bad = X.conditional_format(sales, sheet="Sales", column="Qty", rule="cell_is", operator="between",
                               values=[1])
    assert bad["ok"] is False


def test_data_validation(sales):
    r = ok(X.data_validation(sales, sheet="Sales", column="Qty", type="whole", min=0, max=20,
                             error="0-20 only"))
    assert r["details"]["invalid_existing"] == [{"cell": "D10", "value": 40}]
    ok(X.data_validation(sales, sheet="Products", range="B2:B10", type="list",
                         values=["Hardware", "Electronics"]))
    r = ok(X.data_validation(sales, sheet="Sales", column="Date", type="date",
                             min="2024-01-01", max="2024-03-31"))
    assert r["details"]["invalid_count"] == 2
    r = ok(X.data_validation(sales, sheet="Sales", range="C2:C11", type="list",
                             source="Products!$A$2:$A$5"))
    assert r["details"]["invalid_count"] == 1  # Doohickey
    wb = load(sales)
    dvs = {str(d.sqref): d for d in wb["Sales"].data_validations.dataValidation}
    assert dvs["D2:D11"].type == "whole" and dvs["D2:D11"].formula2 == "20"
    assert dvs["F2:F11"].formula1 == "DATE(2024,1,1)"
    assert wb["Products"].data_validations.dataValidation[0].formula1 == '"Hardware,Electronics"'
    bad = X.data_validation(sales, sheet="Sales", column="Rep", type="list", values=["a,b"])
    assert bad["ok"] is False


def test_format_range(sales):
    r = ok(X.format_range(sales, sheet="Sales", columns=["Qty"], number_format="decimal",
                          bold_header=True, freeze="header", autofit=True))
    ws = load(sales)["Sales"]
    assert ws["D2"].number_format == "#,##0.00" and ws["D11"].number_format == "#,##0.00"
    assert ws.freeze_panes == "A2"
    assert ws["D1"].font.b and ws["D1"].fill.fgColor.rgb == "FF1F4E78"  # fill untouched
    assert r["details"]["column_widths"]["D"] >= 6
    ok(X.format_range(sales, sheet="Products", range="C2:C5", number_format="currency",
                      freeze="B2", autofit=True))
    p = load(sales)["Products"]
    assert p["C2"].number_format == '"$"#,##0.00' and p.freeze_panes == "B2"


def test_from_csv(tmp_path):
    csv_path = tmp_path / "in.csv"
    csv_path.write_text(
        "id,zip,amount,when,ratio,flag,note,card\n"
        "1,02134,\"1,234.50\",2024-01-31,12%,TRUE,=cmd|' /C calc'!A0,4111111111111111\n"
        "2,90210,10,2024-02-29,5.5%,false,hello,4111111111111112\n"
        "3,00501,-7.25,2024-12-01,0%,TRUE,,4111111111111113\n", encoding="utf-8")
    out = str(tmp_path / "out.xlsx")
    r = ok(X.from_csv(out, str(csv_path)))
    t = r["details"]["column_types"]
    assert t["id"] == "number" and t["amount"] == "number" and t["when"] == "date"
    assert t["zip"].startswith("text") and t["card"].startswith("text") and t["flag"] == "bool"
    ws = load(out)["in"]
    assert ws["A2"].value == 1 and ws["B2"].value == "02134" and ws["B2"].number_format == "@"
    assert ws["C2"].value == 1234.5 and ws["C2"].number_format == "#,##0.00"
    assert ws["D3"].value == dt.datetime(2024, 2, 29) and ws["D3"].is_date
    assert ws["E2"].value == pytest.approx(0.12) and ws["E2"].number_format == "0.0%"
    assert ws["F3"].value is False
    assert ws["G2"].data_type == "s" and ws["G2"].value.startswith("=")  # never a live formula
    assert ws["H2"].value == "4111111111111111"
    # a second CSV into the same workbook becomes another sheet; day-first dates
    csv2 = tmp_path / "eu.csv"
    csv2.write_text("d;v\n31/01/2024;1,5\n01/02/2024;2\n", encoding="utf-8")
    r = ok(X.from_csv(out, str(csv2), sheet="EU"))
    eu = load(out)["EU"]
    assert eu["A2"].value == dt.datetime(2024, 1, 31) and eu["A3"].value == dt.datetime(2024, 2, 1)


def test_multisheet_report(sales):
    r = ok(X.multisheet_report(sales, sheet="Sales", group_by="Region", prefix="R-"))
    sheets = {x["group"]: x for x in r["details"]["sheets"]}
    assert sheets["East"]["rows"] == 3 and sheets["North"]["rows"] == 2
    wb = load(sales)
    assert "R-East" in wb.sheetnames and "R-Index" in wb.sheetnames
    east = wb["R-East"]
    assert east["A1"].value == "Region" and east["A1"].font.b and east.freeze_panes == "A2"
    assert east["G2"].value == 30  # computed value of the formula
    idx = wb["R-Index"]
    assert idx["B2"].hyperlink.location == "'R-East'!A1"
    assert X.multisheet_report(sales, sheet="Sales", group_by="Rep", max_groups=2)["ok"] is False


def test_verify_formulas(sales):
    r = ok(X.verify_formulas(sales))
    d = r["details"]
    assert d["formulas"] == 12 and d["error_count"] == 0 and d["unsupported_count"] == 0
    assert r["changed"] is False


def test_dedupe_keeps_comments_and_hyperlinks(sales):
    from openpyxl.comments import Comment
    wb = load(sales)
    ws = wb["Sales"]
    ws["B2"].comment = Comment("top rep", "me")      # row that does not move
    ws["B11"].comment = Comment("last row", "me")    # row that moves up
    ws["C11"].hyperlink = "https://example.com/widget"
    wb.save(sales)
    ok(X.dedupe(sales, sheet="Sales", columns=["Region", "Rep", "Product", "Qty"]))
    ws = load(sales)["Sales"]
    assert ws["B2"].comment.text == "top rep"
    assert ws["B9"].comment.text == "last row" and ws["C9"].hyperlink.target.endswith("/widget")
    assert ws["B11"].comment is None


def test_running_total_is_not_mistaken_for_a_total_row(sales):
    ok(X.calculated_column(sales, sheet="Sales", header="Running", formula="=SUM(D$2:D2)"))
    r = ok(X.read_summary(sales, sheet="Sales"))
    assert r["details"]["sheets"][0]["row_count"] == 10
    assert r["details"]["sheets"][0]["total_row"] is None
    assert cached(sales, "Sales", "H11") == 85
