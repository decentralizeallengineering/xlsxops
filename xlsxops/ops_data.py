# SPDX-License-Identifier: Apache-2.0
"""Table operations: sort, filter, dedupe, split, lookup, pivot, reports, CSV import."""
from __future__ import annotations

import csv
import datetime as _dt
import io
import re
from pathlib import Path

from openpyxl.styles import Font
from openpyxl.utils.cell import get_column_letter
from openpyxl.worksheet.hyperlink import Hyperlink

from . import formula as fx
from . import refs
from ._core import (bolded, S_COLS, S_FORCE, S_HDR, S_OUT, S_PATH, S_SHEET, Book, UserError,
                    ValueReader, append_column, autofit_columns, check_new_sheet_name,
                    copy_style, find_table, finish, norm_key, operation, out_value, put,
                    safe_sheet_name, snapshot_row, table_warnings, write_row)


# --------------------------------------------------------------------------
# helpers
# --------------------------------------------------------------------------
def _rank(v):
    """Excel sort order: numbers < text < logicals < errors; blanks always last."""
    if isinstance(v, bool):
        return (2, int(v))
    if isinstance(v, (int, float)):
        return (0, float(v))
    if isinstance(v, fx.XLError):
        return (3, v.code)
    return (1, str(v).casefold())


def _is_blank(v):
    return v is None or v == ""


def _new_sheet(book, name, after_ws):
    check_new_sheet_name(book.wb, name)
    idx = book.wb.worksheets.index(after_ws) + 1
    return book.wb.create_sheet(title=name, index=idx)


def _copy_out(src_ws, info, reader, dst, rows, cols=None, freeze=True):
    """Copy header + rows (values; formulas become their computed values) with styles."""
    cols = cols or list(range(info.first_col, info.last_col + 1))
    for j, c in enumerate(cols, 1):
        src = src_ws._cells.get((info.header_row, c))
        cell = dst.cell(row=1, column=j)
        copy_style(src, cell)
        put(cell, info.headers.get(c), "text")
        L = get_column_letter(c)
        if L in src_ws.column_dimensions and src_ws.column_dimensions[L].width:
            dst.column_dimensions[get_column_letter(j)].width = src_ws.column_dimensions[L].width
    for i, r in enumerate(rows, 2):
        for j, c in enumerate(cols, 1):
            src = src_ws._cells.get((r, c))
            if src is None:
                continue
            v = reader.get(src_ws, r, c, raw=True)
            cell = dst.cell(row=i, column=j)
            copy_style(src, cell)
            put(cell, out_value(v), "text" if isinstance(v, str) else "value")
    if freeze:
        dst.freeze_panes = "A2"
    return len(rows)


def _row_dict(reader, ws, info, r, limit=12):
    out = {}
    for c in list(range(info.first_col, info.last_col + 1))[:limit]:
        out[info.name(c)] = reader.get(ws, r, c, raw=True)
    return out


# --------------------------------------------------------------------------
# sort
# --------------------------------------------------------------------------
@operation("sort", {
    "type": "object", "required": ["path", "keys"],
    "properties": {"path": S_PATH, "sheet": S_SHEET, "header_row": S_HDR,
                   "keys": {"type": "array", "minItems": 1, "items": {
                       "type": "object", "required": ["column"],
                       "properties": {"column": {"type": "string"},
                                      "descending": {"type": "boolean"}}}},
                   "out_path": S_OUT, "force": S_FORCE},
})
def sort(path, keys, sheet=None, header_row=None, out_path=None, force=False):
    """Sort the table's data rows in place by one or more columns (header stays put)."""
    book = Book(path, force=force)
    ws = book.ws(sheet)
    info = find_table(ws, header_row)
    warnings = table_warnings(info)
    rows = list(info.data_rows)
    if len(rows) < 2:
        return False, warnings + ["fewer than two data rows; nothing to sort"], {"rows": len(rows)}
    c1, c2 = info.first_col, info.last_col
    for m in ws.merged_cells.ranges:
        if m.min_row <= rows[-1] and m.max_row >= rows[0] and m.min_col <= c2 and m.max_col >= c1:
            raise UserError(f"merged cells {m} overlap the rows to sort; unmerge them first",
                            "merged_cell")
    reader = ValueReader(book)
    key_cols = [(info.col(k["column"]), bool(k.get("descending"))) for k in keys]
    vals = {r: {c: reader.get(ws, r, c) for c, _ in key_cols} for r in rows}
    order = list(rows)
    for c, desc in reversed(key_cols):
        filled = [r for r in order if not _is_blank(vals[r][c])]
        blank = [r for r in order if _is_blank(vals[r][c])]
        filled.sort(key=lambda r: _rank(vals[r][c]), reverse=desc)
        order = filled + blank
    if order == rows:
        return False, warnings + reader.warnings(), {"sheet": ws.title, "rows": len(rows),
                                                      "note": "already in order"}
    cross = _cross_row_formulas(ws, info)
    if cross:
        warnings.append(f"{len(cross)} formula(s) in the sorted rows refer to other rows "
                        f"(e.g. {cross[0]}); like Excel, their references shift with the row")
    snaps = {r: snapshot_row(ws, r, c1, c2) for r in rows}
    for dst, src in zip(rows, order):
        write_row(ws, dst, snaps[src], c1, src)
    details = {"sheet": ws.title, "rows": len(rows),
               "range": refs.range_text(c1, rows[0], c2, rows[-1]),
               "keys": [{"column": info.name(c), "descending": d} for c, d in key_cols],
               "new_order_source_rows": order[:50]}
    changed, warns, details = finish(book, out_path, details, warnings + reader.warnings())
    reader2 = ValueReader(book)
    details["preview"] = [_row_dict(reader2, ws, info, r) for r in rows[:5]]
    return changed, warns, details


def _cross_row_formulas(ws, info):
    out = []
    lo, hi = info.first_data, info.last_data
    for r in info.data_rows:
        for c in range(info.first_col, info.last_col + 1):
            cell = ws._cells.get((r, c))
            if cell is None or cell.data_type != "f" or not isinstance(cell.value, str):
                continue
            try:
                toks = fx.references(cell.value)
            except (fx.FormulaSyntaxError, fx.Unsupported):
                continue
            for t in toks:
                if t.sheet is not None and t.sheet.casefold() != ws.title.casefold():
                    continue
                if t.kind in ("cell", "area") and not t.abs_r1 and \
                        (t.r1 != r or t.r2 != r) and lo <= t.r1 <= hi:
                    out.append(f"{get_column_letter(c)}{r}")
                    break
    return out


# --------------------------------------------------------------------------
# filter_rows
# --------------------------------------------------------------------------
_FILTER_OPS = ["=", "!=", ">", ">=", "<", "<=", "contains", "not_contains", "starts_with",
               "ends_with", "in", "not_in", "blank", "not_blank", "regex"]


def _cmp_values(cell_v, want):
    """Return (a, b) comparable, numbers if both numeric/date-like else casefolded text."""
    a = cell_v
    b = want
    if isinstance(a, bool) or isinstance(b, bool):
        return (str(a).casefold(), str(b).casefold())
    if isinstance(a, (int, float)):
        if isinstance(b, (int, float)):
            return a, b
        if isinstance(b, str):
            x = fx._text_to_number(b)
            if x is not None:
                return a, x
    if isinstance(a, fx.XLError):
        a = a.code
    return (fx.general_text(a).casefold(), fx.general_text(b).casefold())


def _make_pred(op, want):
    if op == "blank":
        return _is_blank
    if op == "not_blank":
        return lambda v: not _is_blank(v)
    if op in ("in", "not_in"):
        if not isinstance(want, list):
            want = [want]
        def isin(v):
            return any(_eq(v, w) for w in want)
        return isin if op == "in" else (lambda v: not isin(v))
    if op == "regex":
        try:
            rx = re.compile(str(want))
        except re.error as e:
            raise UserError(f"invalid regex {want!r}: {e}")
        return lambda v: v is not None and rx.search(fx.general_text(v)) is not None
    if want is None and op not in ("=", "!="):
        raise UserError(f"operator {op!r} needs a 'value'")
    if op == "=":
        return lambda v: _eq(v, want)
    if op == "!=":
        return lambda v: not _eq(v, want)
    if op in (">", ">=", "<", "<="):
        def cmp(v):
            if _is_blank(v):
                return False
            a, b = _cmp_values(v, want)
            if type(a) is not type(b) and not (isinstance(a, (int, float)) and isinstance(b, (int, float))):
                return False
            return {">": a > b, ">=": a >= b, "<": a < b, "<=": a <= b}[op]
        return cmp
    w = fx.general_text(want).casefold()
    text = {"contains": lambda s: w in s, "not_contains": lambda s: w not in s,
            "starts_with": lambda s: s.startswith(w), "ends_with": lambda s: s.endswith(w)}[op]
    return lambda v: text(fx.general_text(v).casefold() if not isinstance(v, fx.XLError) else v.code.casefold())


def _eq(v, want):
    if want is None or want == "":
        return _is_blank(v)
    if _is_blank(v):
        return False
    a, b = _cmp_values(v, want)
    return a == b


@operation("filter_rows", {
    "type": "object", "required": ["path", "conditions", "new_sheet"],
    "properties": {"path": S_PATH, "sheet": S_SHEET, "new_sheet": {"type": "string"},
                   "conditions": {"type": "array", "items": {
                       "required": ["column", "op"],
                       "properties": {"column": {"type": "string"}, "op": {"enum": _FILTER_OPS},
                                      "value": {}}}},
                   "match": {"enum": ["all", "any"]}, "header_row": S_HDR,
                   "out_path": S_OUT, "force": S_FORCE},
})
def filter_rows(path, conditions, new_sheet, sheet=None, match="all", header_row=None,
                out_path=None, force=False):
    """Copy rows that match the conditions (with the header) to a new sheet, as values."""
    book = Book(path, force=force)
    ws = book.ws(sheet)
    info = find_table(ws, header_row)
    warnings = table_warnings(info)
    reader = ValueReader(book)
    if not conditions:
        raise UserError("give at least one condition")
    preds = []
    for cond in conditions:
        if cond["op"] not in ("blank", "not_blank") and "value" not in cond:
            raise UserError(f"condition on {cond['column']!r} with op {cond['op']!r} needs 'value'")
        want = cond.get("value")
        c = info.col(cond["column"])
        if isinstance(want, str) and cond["op"] in (">", ">=", "<", "<=", "=", "!="):
            first = next((ws._cells.get((r, c)) for r in info.data_rows
                          if ws._cells.get((r, c)) is not None), None)
            if first is not None and first.is_date:
                x = fx._text_to_number(want)
                want = x if x is not None else want
        preds.append((c, _make_pred(cond["op"], want)))
    combine = all if match == "all" else any
    hits = [r for r in info.data_rows
            if combine(p(reader.get(ws, r, c)) for c, p in preds)]
    dst = _new_sheet(book, new_sheet, ws)
    _copy_out(ws, info, reader, dst, hits)
    if any(ws._cells.get((r, c)) is not None and ws._cells[(r, c)].data_type == "f"
           for r in hits for c in range(info.first_col, info.last_col + 1)):
        warnings.append("formulas were copied as their computed values")
    details = {"source_sheet": ws.title, "new_sheet": dst.title, "matched": len(hits),
               "of_rows": info.n_rows, "source_rows": hits[:100]}
    return finish(book, out_path, details, warnings + reader.warnings())


# --------------------------------------------------------------------------
# dedupe
# --------------------------------------------------------------------------
@operation("dedupe", {
    "type": "object", "required": ["path"],
    "properties": {"path": S_PATH, "sheet": S_SHEET, "columns": S_COLS,
                   "keep": {"enum": ["first", "last"]}, "ignore_case": {"type": "boolean"},
                   "trim": {"type": "boolean"}, "header_row": S_HDR,
                   "out_path": S_OUT, "force": S_FORCE},
})
def dedupe(path, sheet=None, columns=None, keep="first", ignore_case=True, trim=True,
           header_row=None, out_path=None, force=False):
    """Remove duplicate data rows (by key columns) in place; report the removed rows."""
    book = Book(path, force=force)
    ws = book.ws(sheet)
    info = find_table(ws, header_row)
    warnings = table_warnings(info)
    reader = ValueReader(book)
    rows = list(info.data_rows)
    c1, c2 = info.first_col, info.last_col
    key_cols = [info.col(c) for c in columns] if columns else list(range(c1, c2 + 1))
    seen, removed = {}, []
    scan = rows if keep == "first" else list(reversed(rows))
    for r in scan:
        key = tuple(norm_key(reader.get(ws, r, c), ignore_case, trim) for c in key_cols)
        if key in seen:
            removed.append((r, seen[key]))
        else:
            seen[key] = r
    details = {"sheet": ws.title, "key_columns": [info.name(c) for c in key_cols],
               "rows_before": len(rows), "removed_count": len(removed)}
    if not removed:
        details["rows_after"] = len(rows)
        details["removed"] = []
        return False, warnings + reader.warnings(), details
    for m in ws.merged_cells.ranges:
        if m.min_row <= rows[-1] and m.max_row >= rows[0] and m.min_col <= c2 and m.max_col >= c1:
            raise UserError(f"merged cells {m} overlap the table rows; unmerge them first",
                            "merged_cell")
    removed.sort()
    details["removed"] = [{"row": r, "duplicate_of_row": k, "values": _row_dict(reader, ws, info, r)}
                          for r, k in removed[:200]]
    gone = {r for r, _ in removed}
    keep_rows = [r for r in rows if r not in gone]
    first, old_last = info.first_data, info.last_data
    new_last = first + len(keep_rows) - 1
    snaps = {r: snapshot_row(ws, r, c1, c2) for r in rows}
    total_snap = snapshot_row(ws, info.total_row, c1, c2) if info.total_row else None
    old_total = info.total_row
    for i, src in enumerate(keep_rows):
        write_row(ws, first + i, snaps[src], c1, src)
    clear_to = old_total or old_last
    for r in range(new_last + 1, clear_to + 1):
        for c in range(c1, c2 + 1):
            ws._cells.pop((r, c), None)
        if r in ws.row_dimensions:
            ws.row_dimensions[r].height = None
    new_total = None
    if total_snap is not None:
        new_total = new_last + (old_total - old_last)
        write_row(ws, new_total, total_snap, c1, new_total)
    edits, stray = refs.shrink_block(book.wb, ws, first, old_last, new_last, c1, c2,
                                     old_total, new_total)
    if stray:
        warnings.append(f"{stray} formula(s) on other sheets point at single cells inside the "
                        "table; after removing rows they point at whatever row moved there")
    details.update(rows_after=len(keep_rows), references_updated=edits)
    return finish(book, out_path, details, warnings + reader.warnings())


# --------------------------------------------------------------------------
# split_column
# --------------------------------------------------------------------------
_NUM_RX = re.compile(r"^(?P<neg>-)?(?P<cur>[$€£])?(?P<num>\d{1,3}(?:,\d{3})+(?:\.\d+)?|\d+(?:\.\d+)?|\.\d+)(?P<pct>%)?$")


def infer_scalar(text: str):
    """Text -> number where unambiguous; keeps leading zeros and long IDs as text."""
    t = text.strip()
    if not t:
        return None, "blank"
    neg = False
    if t.startswith("(") and t.endswith(")"):
        neg, t = True, t[1:-1]
    m = _NUM_RX.match(t.replace("+", "", 1) if t.startswith("+") else t)
    if not m:
        return text, "text"
    digits = re.sub(r"\D", "", m.group("num"))
    if re.match(r"0\d", m.group("num")) or len(digits.lstrip("0")) > 15:
        return text, "text"
    x = float(m.group("num").replace(",", ""))
    if m.group("pct"):
        x /= 100
    if neg or m.group("neg"):
        x = -x
    if x.is_integer() and not m.group("pct") and "." not in m.group("num"):
        x = int(x)
    return x, "number"


@operation("split_column", {
    "type": "object", "required": ["path", "column"],
    "properties": {"path": S_PATH, "sheet": S_SHEET, "column": {"type": "string"},
                   "delimiter": {"type": "string", "minLength": 1}, "regex": {"type": "string"},
                   "headers": S_COLS, "max_parts": {"type": "integer", "minimum": 1, "maximum": 50},
                   "infer_types": {"type": "boolean"}, "header_row": S_HDR,
                   "out_path": S_OUT, "force": S_FORCE},
})
def split_column(path, column, sheet=None, delimiter=None, regex=None, headers=None,
                 max_parts=None, infer_types=True, header_row=None, out_path=None, force=False):
    """Split a column into new columns (appended after the table) by delimiter or regex.

    A regex with capture groups yields one column per group; otherwise it splits.
    """
    if bool(delimiter) == bool(regex):
        raise UserError("give exactly one of 'delimiter' or 'regex'")
    rx = None
    if regex:
        try:
            rx = re.compile(regex)
        except re.error as e:
            raise UserError(f"invalid regex: {e}")
    book = Book(path, force=force)
    ws = book.ws(sheet)
    info = find_table(ws, header_row)
    warnings = table_warnings(info)
    reader = ValueReader(book)
    src = info.col(column)
    parts_by_row, no_match = [], 0
    for r in info.data_rows:
        v = reader.get(ws, r, src, raw=True)
        if v is None:
            parts_by_row.append([])
            continue
        s = v.isoformat() if isinstance(v, (_dt.date, _dt.datetime)) else fx.general_text(v)
        if rx is not None and rx.groups:
            m = rx.search(s)
            if m is None:
                no_match += 1
                parts = []
            else:
                parts = [g if g is not None else "" for g in m.groups()]
        elif rx is not None:
            parts = rx.split(s, maxsplit=(max_parts - 1) if max_parts else 0)
        else:
            parts = s.split(delimiter, (max_parts - 1) if max_parts else -1)
        parts_by_row.append([p.strip() for p in parts])
    n = max((len(p) for p in parts_by_row), default=0)
    if max_parts:
        n = min(n, max_parts)
    if n == 0:
        raise UserError("nothing to split: no values matched")
    if n > 50:
        raise UserError(f"splitting would create {n} columns (max 50); set max_parts")
    base = info.name(src)
    names = list(headers or [])
    if len(names) < n:
        names += [f"{base} {i}" for i in range(len(names) + 1, n + 1)]
    elif len(names) > n:
        warnings.append(f"{len(names) - n} extra header(s) ignored; values split into {n} parts")
        names = names[:n]
    new_cols = []
    for i, name in enumerate(names):
        items = []
        for parts in parts_by_row:
            p = parts[i] if i < len(parts) else None
            if p is None or p == "":
                items.append((None, "blank"))
            elif infer_types:
                items.append(infer_scalar(p))
            else:
                items.append((p, "text"))
        new_cols.append(get_column_letter(append_column(ws, info, name, items)))
    if no_match:
        warnings.append(f"regex did not match {no_match} row(s); their new cells are blank")
    details = {"sheet": ws.title, "source_column": base, "new_columns": dict(zip(names, new_cols)),
               "rows": info.n_rows, "preview": [p[:n] for p in parts_by_row[:5]]}
    return finish(book, out_path, details, warnings + reader.warnings())


# --------------------------------------------------------------------------
# lookup
# --------------------------------------------------------------------------
def _literal(v):
    if isinstance(v, bool):
        return "TRUE" if v else "FALSE"
    if isinstance(v, (int, float)):
        return fx.general_text(v)
    return '"' + str(v).replace('"', '""') + '"'


@operation("lookup", {
    "type": "object",
    "required": ["path", "key", "source_sheet", "columns"],
    "properties": {"path": S_PATH, "sheet": S_SHEET, "key": {"type": "string"},
                   "source_sheet": {"type": "string"}, "source_key": {"type": "string"},
                   "columns": S_COLS, "mode": {"enum": ["formula", "value"]},
                   "style": {"enum": ["xlookup", "vlookup", "index_match"]},
                   "not_found": {}, "header_row": S_HDR,
                   "source_header_row": S_HDR, "out_path": S_OUT, "force": S_FORCE},
})
def lookup(path, key, source_sheet, columns, sheet=None, source_key=None, mode="formula",
           style="xlookup", not_found="", header_row=None, source_header_row=None,
           out_path=None, force=False):
    """Join columns from another sheet by a key (exact match), as formulas or values.

    not_found="#N/A" leaves Excel's #N/A for keys that are missing.
    """
    book = Book(path, force=force)
    ws = book.ws(sheet)
    src_ws = book.ws(source_sheet)
    info = find_table(ws, header_row)
    sinfo = find_table(src_ws, source_header_row)
    warnings = table_warnings(info) + table_warnings(sinfo)
    reader = ValueReader(book)
    kcol = info.col(key)
    skcol = sinfo.col(source_key or info.name(kcol))
    rcols = [sinfo.col(c) for c in columns]
    if sinfo.n_rows == 0:
        raise UserError(f"source sheet {src_ws.title!r} has no data rows")
    # Python-side matching (Excel exact match: case-insensitive text, type-strict).
    index, loose, dups = {}, {}, 0
    for r in sinfo.data_rows:
        v = reader.get(src_ws, r, skcol)
        if _is_blank(v):
            continue
        k = norm_key(v, True, False)
        if k in index:
            dups += 1
            continue
        index[k] = r
        loose.setdefault(fx.general_text(v).strip().casefold(), r)
    matches, unmatched, type_hint = [], [], 0
    for r in info.data_rows:
        v = reader.get(ws, r, kcol)
        hit = None if _is_blank(v) else index.get(norm_key(v, True, False))
        if hit is None and not _is_blank(v):
            unmatched.append(v)
            if fx.general_text(v).strip().casefold() in loose:
                type_hint += 1
        matches.append(hit)
    if dups:
        warnings.append(f"{dups} duplicate key(s) in {src_ws.title!r}; the first match is used")
    if type_hint:
        warnings.append(f"{type_hint} key(s) only match if numbers and text are treated alike "
                        "(e.g. 1001 vs '1001'); convert one side so the types agree")
    na = isinstance(not_found, str) and not_found.strip().upper() == "#N/A"
    first, last = sinfo.first_data, sinfo.last_data
    sp = fx.quote_sheet(src_ws.title) + "!"
    KL = get_column_letter(kcol)

    def rng(c):
        L = get_column_letter(c)
        return f"{sp}${L}${first}:${L}${last}"

    added, blank_hits = {}, 0
    used_names = {str(h).strip().casefold() for h in info.headers.values() if h is not None}
    for rc in rcols:
        name = sinfo.name(rc)
        if name.strip().casefold() in used_names:
            name = f"{name} ({src_ws.title})"
        used_names.add(name.casefold())
        fmt_cell = src_ws._cells.get((first, rc))
        nf = fmt_cell.number_format if fmt_cell is not None else None
        items = []
        if mode == "value":
            for hit in matches:
                if hit is None:
                    items.append((fx.NA.code if na else (not_found if not_found != "" else None),
                                  "value"))
                else:
                    v = reader.get(src_ws, hit, rc, raw=True)
                    items.append((out_value(v), "text" if isinstance(v, str) else "value"))
        else:
            use = style
            if use == "vlookup" and rc < skcol:
                use = "index_match"
                warnings.append(f"VLOOKUP cannot return {name!r} (left of the key); "
                                "used INDEX/MATCH instead")
            for r, hit in zip(info.data_rows, matches):
                kc = f"${KL}{r}"
                if use == "xlookup":
                    f = f"XLOOKUP({kc},{rng(skcol)},{rng(rc)}" + ("" if na else f",{_literal(not_found)}") + ")"
                else:
                    if use == "vlookup":
                        tbl = f"{sp}${get_column_letter(skcol)}${first}:${get_column_letter(rc)}${last}"
                        core = f"VLOOKUP({kc},{tbl},{rc - skcol + 1},FALSE)"
                    else:
                        core = f"INDEX({rng(rc)},MATCH({kc},{rng(skcol)},0))"
                    f = core if na else f"IFNA({core},{_literal(not_found)})"
                items.append((fx.add_future_prefixes("=" + f), "formula"))
                if hit is not None and _is_blank(reader.get(src_ws, hit, rc)):
                    blank_hits += 1
        col = append_column(ws, info, name, items, nf if nf and nf != "General" else None)
        added[name] = get_column_letter(col)
    if blank_hits:
        warnings.append(f"{blank_hits} matched source cell(s) are blank; Excel lookup formulas "
                        "show 0 for those")
    details = {"sheet": ws.title, "source_sheet": src_ws.title, "mode": mode,
               "style": style if mode == "formula" else None, "new_columns": added,
               "rows": info.n_rows, "matched": sum(1 for m in matches if m is not None),
               "unmatched": len(unmatched), "unmatched_keys": unmatched[:20]}
    return finish(book, out_path, details, warnings + reader.warnings())


# --------------------------------------------------------------------------
# pivot_summary
# --------------------------------------------------------------------------
_AGGS = ["sum", "count", "avg", "min", "max"]


@operation("pivot_summary", {
    "type": "object", "required": ["path", "group_by", "values", "new_sheet"],

    "properties": {"path": S_PATH, "sheet": S_SHEET, "group_by": S_COLS,
                   "values": {"type": "array", "items": {
                       "required": ["column", "agg"],
                       "properties": {"column": {"type": "string"}, "agg": {"enum": _AGGS}}}},
                   "new_sheet": {"type": "string"}, "formulas": {"type": "boolean"},
                   "total": {"type": "boolean"}, "header_row": S_HDR,
                   "out_path": S_OUT, "force": S_FORCE},
})
def pivot_summary(path, group_by, values, new_sheet, sheet=None, formulas=False, total=True,
                  header_row=None, out_path=None, force=False):
    """Group-by summary (sum/count/avg/min/max) into a new sheet as values.

    formulas=true also writes a live SUMIFS/COUNTIFS version on '<new_sheet> (formulas)'.
    """
    book = Book(path, force=force)
    ws = book.ws(sheet)
    info = find_table(ws, header_row)
    warnings = table_warnings(info)
    reader = ValueReader(book)
    if not values:
        raise UserError("give at least one entry in 'values', e.g. {\"column\": \"Amount\", \"agg\": \"sum\"}")
    gcols = [info.col(c) for c in group_by]
    specs = [(info.col(v["column"]), v["agg"]) for v in values]
    groups: dict = {}
    display: dict = {}
    for r in info.data_rows:
        raw = [reader.get(ws, r, c, raw=True) for c in gcols]
        k = tuple(norm_key(reader.get(ws, r, c), True, False) for c in gcols)
        groups.setdefault(k, []).append(r)
        display.setdefault(k, raw)
    if len(groups) > 10000:
        raise UserError(f"{len(groups)} groups; too many for a summary sheet")

    def sort_key(k):
        return tuple((1,) if p == ("blank",) else (0, _rank(p[1] if len(p) > 1 else None)) for p in k)

    keys = sorted(groups, key=sort_key)

    def agg(rows, c, how):
        vals = [reader.get(ws, r, c) for r in rows]
        if how == "count":
            return sum(1 for v in vals if not _is_blank(v))
        nums = [v for v in vals if isinstance(v, (int, float)) and not isinstance(v, bool)]
        if how == "sum":
            return fx.clean_number(sum(nums)) if nums else 0
        if how == "avg":
            return fx.clean_number(sum(nums) / len(nums)) if nums else None
        if how == "min":
            return min(nums) if nums else 0
        return max(nums) if nums else 0

    labels = {"sum": "Sum", "count": "Count", "avg": "Average", "min": "Min", "max": "Max"}
    out_headers = [info.name(c) for c in gcols] + [f"{labels[a]} of {info.name(c)}" for c, a in specs]
    dst = _new_sheet(book, new_sheet, ws)
    all_rows = list(info.data_rows)

    def fmt_of(c):
        cell = ws._cells.get((info.first_data, c))
        return cell.number_format if cell is not None else "General"

    def header(sheet_obj):
        for j, h in enumerate(out_headers, 1):
            cell = sheet_obj.cell(row=1, column=j)
            copy_style(ws._cells.get((info.header_row, gcols[0])), cell)
            put(cell, h, "text")
            cell.font = bolded(cell.font)
        sheet_obj.freeze_panes = "A2"

    header(dst)
    preview = []
    for i, k in enumerate(keys, 2):
        rows = groups[k]
        for j, c in enumerate(gcols, 1):
            v = display[k][j - 1]
            cell = dst.cell(row=i, column=j)
            if _is_blank(v):
                put(cell, "(blank)", "text")
            else:
                put(cell, out_value(v), "text" if isinstance(v, str) else "value")
                cell.number_format = fmt_of(c)
        row_vals = []
        for j, (c, a) in enumerate(specs, len(gcols) + 1):
            v = agg(rows, c, a)
            cell = dst.cell(row=i, column=j)
            cell.value = v
            if a != "count":
                cell.number_format = fmt_of(c)
            row_vals.append(v)
        if len(preview) < 10:
            preview.append([fx.general_text(display[k][0]) if display[k][0] is not None else "(blank)"] + row_vals)
    total_row_idx = None
    if total and keys:
        total_row_idx = len(keys) + 2
        cell = dst.cell(row=total_row_idx, column=1)
        put(cell, "Total", "text")
        cell.font = Font(bold=True)
        for j, (c, a) in enumerate(specs, len(gcols) + 1):
            cell = dst.cell(row=total_row_idx, column=j)
            cell.value = agg(all_rows, c, a)
            cell.font = Font(bold=True)
            if a != "count":
                cell.number_format = fmt_of(c)
    autofit_columns(dst, set(range(1, len(out_headers) + 1)))
    details = {"source_sheet": ws.title, "new_sheet": dst.title, "groups": len(keys),
               "headers": out_headers, "preview": preview}
    if formulas:
        fname = safe_sheet_name(book.wb, f"{new_sheet} (formulas)")
        fws = _new_sheet(book, fname, dst)
        header(fws)
        first, last = info.first_data, info.last_data
        sp = fx.quote_sheet(ws.title) + "!"

        def rng(c):
            L = get_column_letter(c)
            return f"{sp}${L}${first}:${L}${last}"

        for i, k in enumerate(keys, 2):
            crit = []
            for j, c in enumerate(gcols, 1):
                v = display[k][j - 1]
                cell = fws.cell(row=i, column=j)
                if _is_blank(v):
                    put(cell, "(blank)", "text")
                    crit.append(f'{rng(c)},""')
                    continue
                put(cell, out_value(v), "text" if isinstance(v, str) else "value")
                cell.number_format = fmt_of(c)
                if isinstance(v, str) and any(ch in v for ch in "*?~"):
                    esc = re.sub(r"([*?~])", r"~\1", v)
                    crit.append(f"{rng(c)},{_literal('=' + esc)}")
                else:
                    crit.append(f'{rng(c)},"="&${get_column_letter(j)}{i}')
            cs = ",".join(crit)
            for j, (c, a) in enumerate(specs, len(gcols) + 1):
                if a == "sum":
                    f = f"=SUMIFS({rng(c)},{cs})"
                elif a == "count":
                    f = f'=COUNTIFS({cs},{rng(c)},"<>")'
                elif a == "avg":
                    f = f'=IFERROR(AVERAGEIFS({rng(c)},{cs}),"")'
                else:
                    f = f"={a.upper()}IFS({rng(c)},{cs})"
                cell = fws.cell(row=i, column=j)
                put(cell, fx.add_future_prefixes(f), "formula")
                if a != "count":
                    cell.number_format = fmt_of(c)
        if total_row_idx:
            cell = fws.cell(row=total_row_idx, column=1)
            put(cell, "Total", "text")
            cell.font = Font(bold=True)
            for j, (c, a) in enumerate(specs, len(gcols) + 1):
                L = get_column_letter(j)
                f = {"sum": f"=SUM({L}2:{L}{total_row_idx - 1})",
                     "count": f"=SUM({L}2:{L}{total_row_idx - 1})",
                     "avg": f'=IFERROR(AVERAGE({rng(c)}),"")',
                     "min": f"=MIN({rng(c)})", "max": f"=MAX({rng(c)})"}[a]
                cell = fws.cell(row=total_row_idx, column=j)
                put(cell, f, "formula")
                cell.font = Font(bold=True)
                if a != "count":
                    cell.number_format = fmt_of(c)
        autofit_columns(fws, set(range(1, len(out_headers) + 1)))
        details["formula_sheet"] = fws.title
    changed, warns, details = finish(book, out_path, details, warnings + reader.warnings())
    if formulas:
        mism = _compare_sheets(book, dst, details["formula_sheet"], len(keys) + (1 if total else 0),
                               len(gcols), len(out_headers))
        details["formula_version_matches_values"] = not mism
        if mism:
            warns.append(f"formula version differs from values at {mism[:5]}")
    return changed, warns, details


def _compare_sheets(book, dst, fname, nrows, first_col, ncols):
    vals = book.last_values
    out = []
    for i in range(2, nrows + 2):
        for j in range(first_col + 1, ncols + 1):
            a = dst._cells.get((i, j))
            a = a.value if a is not None else None
            b = vals.get((fname, i, j), "?")
            if b == "":
                b = None
            if isinstance(a, float) or isinstance(b, float):
                if a is not None and b not in (None, "?") and not isinstance(b, fx.XLError) \
                        and abs(float(a) - float(b)) <= 1e-9 * max(1, abs(float(a))):
                    continue
            if a != b:
                out.append(f"{get_column_letter(j)}{i}")
    return out


# --------------------------------------------------------------------------
# multisheet_report
# --------------------------------------------------------------------------
@operation("multisheet_report", {
    "type": "object", "required": ["path", "group_by"],
    "properties": {"path": S_PATH, "sheet": S_SHEET, "group_by": {"type": "string"},
                   "prefix": {"type": "string", "maxLength": 20}, "index": {"type": "boolean"},
                   "max_groups": {"type": "integer", "minimum": 1, "maximum": 250},
                   "header_row": S_HDR, "out_path": S_OUT, "force": S_FORCE},
})
def multisheet_report(path, group_by, sheet=None, prefix="", index=True, max_groups=50,
                      header_row=None, out_path=None, force=False):
    """One new sheet per distinct value of a column (rows copied as values), plus an index."""
    book = Book(path, force=force)
    ws = book.ws(sheet)
    info = find_table(ws, header_row)
    warnings = table_warnings(info)
    reader = ValueReader(book)
    gcol = info.col(group_by)
    groups, display = {}, {}
    for r in info.data_rows:
        k = norm_key(reader.get(ws, r, gcol), True, True)
        groups.setdefault(k, []).append(r)
        display.setdefault(k, reader.get(ws, r, gcol, raw=True))
    if len(groups) > max_groups:
        raise UserError(f"{len(groups)} distinct values in {info.name(gcol)!r} (max_groups="
                        f"{max_groups}); raise max_groups or group by another column")
    keys = sorted(groups, key=lambda k: (1,) if k == ("blank",) else (0, _rank(k[1])))
    created, taken = [], set()
    anchor = ws
    for k in keys:
        v = display[k]
        label = "(blank)" if _is_blank(v) else (v.date().isoformat() if isinstance(v, _dt.datetime)
                                                else fx.general_text(v))
        name = safe_sheet_name(book.wb, f"{prefix}{label}", taken)
        taken.add(name.casefold())
        dst = book.wb.create_sheet(title=name, index=book.wb.worksheets.index(anchor) + 1)
        anchor = dst
        _copy_out(ws, info, reader, dst, groups[k])
        last_col = get_column_letter(info.last_col - info.first_col + 1)
        if groups[k]:
            dst.auto_filter.ref = f"A1:{last_col}{len(groups[k]) + 1}"
        created.append({"group": label, "sheet": name, "rows": len(groups[k])})
    index_name = None
    if index:
        index_name = safe_sheet_name(book.wb, f"{prefix}Index" if prefix else "Report Index", taken)
        iws = book.wb.create_sheet(title=index_name, index=book.wb.worksheets.index(ws) + 1)
        for j, h in enumerate([info.name(gcol), "Sheet", "Rows"], 1):
            c = iws.cell(row=1, column=j, value=h)
            c.font = Font(bold=True)
        for i, item in enumerate(created, 2):
            put(iws.cell(row=i, column=1), item["group"], "text")
            link = iws.cell(row=i, column=2, value=item["sheet"])
            link.hyperlink = Hyperlink(ref=link.coordinate,
                                       location=f"{fx.quote_sheet(item['sheet'])}!A1")
            link.style = "Hyperlink"
            iws.cell(row=i, column=3, value=item["rows"])
        iws.freeze_panes = "A2"
        autofit_columns(iws, {1, 2, 3})
    details = {"source_sheet": ws.title, "group_by": info.name(gcol), "sheets": created,
               "index_sheet": index_name}
    return finish(book, out_path, details, warnings + reader.warnings())


# --------------------------------------------------------------------------
# from_csv
# --------------------------------------------------------------------------
_DATE_FORMATS = [
    ("%Y-%m-%d", None), ("%Y/%m/%d", None), ("%m/%d/%Y", "m"), ("%d/%m/%Y", "d"),
    ("%m-%d-%Y", "m"), ("%d-%m-%Y", "d"), ("%d.%m.%Y", "d"), ("%m/%d/%y", "m"), ("%d/%m/%y", "d"),
    ("%d-%b-%Y", None), ("%d %b %Y", None), ("%b %d, %Y", None), ("%B %d, %Y", None),
    ("%Y-%m-%dT%H:%M:%S", None), ("%Y-%m-%d %H:%M:%S", None), ("%Y-%m-%d %H:%M", None),
    ("%Y-%m-%dT%H:%M", None), ("%m/%d/%Y %H:%M", "m"), ("%m/%d/%Y %H:%M:%S", "m"),
    ("%d/%m/%Y %H:%M", "d"), ("%d/%m/%Y %H:%M:%S", "d"),
]
_DATEISH = re.compile(r"^\s*(\d{1,4}[-/.]\d{1,2}[-/.]\d{2,4}|\d{1,2}[ -][A-Za-z]{3}|[A-Za-z]{3,9} \d{1,2},)")


def _parse_all(values, fmt):
    out = []
    for v in values:
        try:
            out.append(_dt.datetime.strptime(v.strip(), fmt))
        except ValueError:
            return None
    return out


def _infer_column(values, dayfirst):
    """Decide a column's type from all its non-empty strings."""
    filled = [v for v in values if v.strip()]
    if not filled:
        return "text", None, None
    if all(v.strip().lower() in ("true", "false") for v in filled):
        return "bool", None, None
    if all(_DATEISH.match(v) for v in filled):
        cands = [(f, order) for f, order in _DATE_FORMATS if _parse_all(filled, f) is not None]
        if cands:
            orders = {o for _, o in cands if o}
            note = None
            if orders == {"m", "d"}:
                want = "d" if dayfirst else "m"
                if dayfirst is None:
                    note = "ambiguous day/month order; read as month-first (set dayfirst)"
                cands = [c for c in cands if c[1] in (want, None)]
            fmt = cands[0][0]
            has_time = "%H" in fmt
            return "date", fmt, ("yyyy-mm-dd hh:mm:ss" if has_time else "yyyy-mm-dd", note)
    results = [infer_scalar(v) for v in filled]
    if any(k == "text" and re.fullmatch(r"\s*[-+]?0\d+(\.\d+)?\s*|\s*\d{16,}\s*", v)
           for v, (x, k) in zip(filled, results)):
        return "text_id", None, None
    nums = sum(1 for _, k in results if k == "number")
    if nums == len(filled):
        return "number", None, _number_format(filled)
    if nums:
        return "mixed", None, _number_format([v for v, (_, k) in zip(filled, results) if k == "number"])
    return "text", None, None


def _number_format(samples):
    s = [x.strip() for x in samples]
    if all(x.endswith("%") for x in s):
        dec = max((len(x.split(".")[1]) - 1 if "." in x else 0) for x in s)
        return "0" + ("." + "0" * dec if dec else "") + "%"
    if all(re.match(r"^\(?-?[$]", x) for x in s):
        return '"$"#,##0.00'
    if all(re.match(r"^\(?-?[€]", x) for x in s):
        return '"€"#,##0.00'
    if all(re.match(r"^\(?-?[£]", x) for x in s):
        return '"£"#,##0.00'
    if any("," in x for x in s):
        return "#,##0.00" if any("." in x for x in s) else "#,##0"
    return None


@operation("from_csv", {
    "type": "object", "required": ["path", "csv_path"],
    "properties": {"path": S_PATH, "csv_path": {"type": "string"}, "sheet": S_SHEET,
                   "delimiter": {"type": "string", "minLength": 1, "maxLength": 1},
                   "encoding": {"type": "string"}, "header": {"type": "boolean"},
                   "dayfirst": {"type": "boolean"}, "infer_types": {"type": "boolean"},
                   "force": S_FORCE},
})
def from_csv(path, csv_path, sheet=None, delimiter=None, encoding=None, header=True,
             dayfirst=None, infer_types=True, force=False):
    """Import a CSV into a new sheet (new workbook if path does not exist) with type inference.

    Numbers, dates, percents and currency become typed cells; values with leading zeros
    or more than 15 digits stay text; text starting with '=' is never a formula.
    """
    src = Path(csv_path)
    if not src.is_file():
        raise UserError(f"CSV file not found: {src}", "file_not_found")
    raw = src.read_bytes()
    warnings = []
    text = None
    for enc in ([encoding] if encoding else ["utf-8-sig", "cp1252", "latin-1"]):
        try:
            text = raw.decode(enc)
            if enc != (encoding or "utf-8-sig"):
                warnings.append(f"CSV is not UTF-8; decoded as {enc}")
            break
        except (UnicodeDecodeError, LookupError) as e:
            if encoding:
                raise UserError(f"cannot decode CSV as {encoding}: {e}")
    if delimiter is None:
        sample = text[:65536]
        try:
            delimiter = csv.Sniffer().sniff(sample, delimiters=",;\t|").delimiter
        except csv.Error:
            delimiter = ","
    rows = list(csv.reader(io.StringIO(text, newline=""), delimiter=delimiter))
    rows = [r for r in rows if any(x.strip() for x in r)] if rows else rows
    if not rows:
        raise UserError("the CSV is empty")
    width = max(len(r) for r in rows)
    ragged = sum(1 for r in rows if len(r) != width)
    if ragged:
        warnings.append(f"{ragged} row(s) had fewer fields than {width}; padded with blanks")
    rows = [r + [""] * (width - len(r)) for r in rows]
    if header:
        heads, body = rows[0], rows[1:]
        heads = [h.strip() or f"Column{j + 1}" for j, h in enumerate(heads)]
        seen = {}
        for j, h in enumerate(heads):
            if h.casefold() in seen:
                seen[h.casefold()] += 1
                heads[j] = f"{h}_{seen[h.casefold()]}"
            else:
                seen[h.casefold()] = 1
    else:
        heads, body = None, rows
    book = Book(path, create=True, force=force)
    name = sheet or safe_sheet_name(book.wb, src.stem)
    if book.created:
        ws = book.wb.active
        ws.title = check_new_sheet_name(None, name)
    else:
        check_new_sheet_name(book.wb, name)
        ws = book.wb.create_sheet(title=name)
    r0 = 1
    if heads:
        for j, h in enumerate(heads, 1):
            cell = ws.cell(row=1, column=j)
            put(cell, h, "text")
            cell.font = Font(bold=True)
        r0 = 2
    col_types = {}
    for j in range(width):
        col = [r[j] for r in body]
        label = heads[j] if heads else get_column_letter(j + 1)
        kind, fmt, extra = _infer_column(col, dayfirst) if infer_types else ("text", None, None)
        col_types[label] = {"text_id": "text (kept as text: leading zeros / long digits)"}.get(kind, kind)
        number_format = None
        if kind == "date":
            number_format, note = extra
            if note:
                warnings.append(f"column {label!r}: {note}")
        elif kind in ("number", "mixed"):
            number_format = extra
            if kind == "mixed":
                bad = [v for v in col if v.strip() and infer_scalar(v)[1] != "number"][:3]
                warnings.append(f"column {label!r} mixes numbers and text (e.g. {bad}); "
                                "numbers were converted, text kept")
        for i, v in enumerate(col):
            if not v.strip():
                continue
            cell = ws.cell(row=r0 + i, column=j + 1)
            if kind == "date":
                d = _dt.datetime.strptime(v.strip(), fmt)
                put(cell, d if "%H" in fmt else d.date(), "value")
                cell.number_format = number_format
            elif kind == "bool":
                cell.value = v.strip().lower() == "true"
            elif kind in ("number", "mixed"):
                x, k = infer_scalar(v)
                put(cell, x, "text" if k == "text" else "value")
                if k == "number" and number_format:
                    cell.number_format = number_format
            else:
                put(cell, v, "text")
                if kind == "text_id":
                    cell.number_format = "@"
    if heads:
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = f"A1:{get_column_letter(width)}{max(len(body) + 1, 2)}"
    autofit_columns(ws, set(range(1, width + 1)))
    details = {"sheet": ws.title, "rows": len(body), "columns": width, "delimiter": delimiter,
               "column_types": col_types, "created_workbook": book.created}
    return finish(book, None, details, warnings)
