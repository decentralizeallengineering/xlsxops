# SPDX-License-Identifier: Apache-2.0
"""Shared plumbing: reports, argument validation, safe load/save, table model."""
from __future__ import annotations

import contextlib
import datetime as _dt
import functools
import inspect
import json
import math
import os
import re
import tempfile
import threading
import warnings
import zipfile
from copy import copy
from dataclasses import dataclass, field
from pathlib import Path
from xml.etree import ElementTree as ET

import openpyxl
from jsonschema import Draft202012Validator
from openpyxl.drawing.image import Image as _OpenpyxlImage
from openpyxl.formula.translate import Translator
from openpyxl.utils.cell import (column_index_from_string, get_column_letter,
                                 range_boundaries)
from openpyxl.utils.exceptions import InvalidFileException
from openpyxl.worksheet.worksheet import Worksheet

from . import formula as fx

SCHEMAS: dict[str, dict] = {}
OPERATIONS: dict[str, callable] = {}
SCHEMA_LIMIT = 600
SUFFIXES = (".xlsx", ".xlsm")


class UserError(Exception):
    """A problem with the caller's input or file; reported, never raised."""

    def __init__(self, message: str, code: str = "invalid_request", **extra):
        super().__init__(message)
        self.code = code
        self.extra = extra


# --------------------------------------------------------------------------
# Reports
# --------------------------------------------------------------------------
def jsonable(v):
    """Convert a value into something json.dumps accepts."""
    if v is None or isinstance(v, (bool, int, str)):
        return v
    if isinstance(v, float):
        return v if math.isfinite(v) else None
    if isinstance(v, dict):
        return {str(k): jsonable(x) for k, x in v.items()}
    if isinstance(v, (list, tuple, set)):
        return [jsonable(x) for x in v]
    if isinstance(v, (_dt.datetime, _dt.date, _dt.time)):
        return v.isoformat()
    if isinstance(v, _dt.timedelta):
        return v.total_seconds()
    if isinstance(v, fx.XLError):
        return v.code
    return str(v)


def _report(ok, changed, warnings, details):
    return {"ok": bool(ok), "changed": bool(changed),
            "warnings": [str(w) for w in warnings], "details": jsonable(details)}


def _fail(message, code, extra=None):
    details = {"error": message, "error_type": code}
    if extra:
        details.update(extra)
    return _report(False, False, [], details)


def _schema_errors(validator, params):
    msgs = []
    for e in sorted(validator.iter_errors(params), key=lambda e: list(e.path)):
        where = "/".join(str(p) for p in e.path) or "arguments"
        msgs.append(f"{where}: {e.message}")
    return msgs


def operation(name: str, schema: dict):
    """Register a public operation: validate args, catch errors, shape report.

    The wrapped function returns ``(changed, warnings, details)``.
    """
    size = len(json.dumps(schema))
    if size > SCHEMA_LIMIT:
        raise AssertionError(f"schema for {name} is {size} chars (> {SCHEMA_LIMIT})")
    Draft202012Validator.check_schema(schema)
    validator = Draft202012Validator(schema)

    def deco(fn):
        sig = inspect.signature(fn)

        @functools.wraps(fn)
        def wrapper(*args, **kwargs):
            try:
                bound = sig.bind(*args, **kwargs)
            except TypeError as e:
                return _fail(f"bad arguments: {e}", "invalid_arguments")
            params = {k: v for k, v in bound.arguments.items() if v is not None}
            problems = _schema_errors(validator, params)
            if problems:
                return _fail("; ".join(problems[:5]), "invalid_arguments")
            try:
                changed, warnings, details = fn(**params)
                return _report(True, changed, warnings, details)
            except UserError as e:
                return _fail(str(e), e.code, e.extra)
            except FileNotFoundError as e:
                return _fail(f"file not found: {e.filename or e}", "file_not_found")
            except PermissionError as e:
                return _fail(f"permission denied (is the file open in Excel?): "
                             f"{e.filename or e}", "permission_denied")
            except (zipfile.BadZipFile, InvalidFileException) as e:
                return _fail(f"not a valid .xlsx workbook: {e}", "invalid_file")
            except Exception as e:  # never let a tool call crash the agent loop
                return _fail(f"{type(e).__name__}: {e}", "internal_error")

        wrapper.schema = schema
        SCHEMAS[name] = schema
        OPERATIONS[name] = wrapper
        return wrapper
    return deco


# --------------------------------------------------------------------------
# Preflight: features openpyxl cannot carry through a load/save
# --------------------------------------------------------------------------
_PIL_AVAILABLE = bool(getattr(__import__("openpyxl.drawing.image", fromlist=["PILImage"]),
                              "PILImage", False))


def preflight(path: Path) -> tuple[list[str], list[str]]:
    """Return (blocking, notices): content that would be lost on save."""
    blocking, notices = [], []
    try:
        z = zipfile.ZipFile(path)
    except (zipfile.BadZipFile, OSError):
        return blocking, notices
    with z:
        names = z.namelist()

        def has(prefix):
            return any(n.startswith(prefix) for n in names)

        def add(cond, msg, lst):
            if cond and msg not in lst:
                lst.append(msg)

        add(has("xl/slicers/") or has("xl/slicerCaches/"), "slicers", blocking)
        add(has("xl/timelines/") or has("xl/timelineCaches/"), "timelines", blocking)
        add(has("xl/ctrlProps/"), "form controls", blocking)
        add(any(n.startswith("xl/charts/chartEx") for n in names),
            "new-style charts (waterfall/treemap/etc.)", blocking)
        add(has("xl/model/"), "Power Pivot data model", blocking)
        add("xl/connections.xml" in names or has("xl/queryTables/"),
            "external data connections / queries", blocking)
        add(has("xl/threadedComments/"), "threaded comments (plain notes are kept)", blocking)
        add(has("xl/activeX/"), "ActiveX controls", blocking)
        add(has("xl/charts/style") or has("xl/charts/colors"),
            "chart style/color parts (charts keep data, titles and formatting)", notices)
        add(has("customXml/"), "custom XML parts", notices)
        for n in names:
            if n.startswith("xl/drawings/") and n.endswith(".xml") and "/_rels/" not in n:
                data = z.read(n)
                add(b"<xdr:sp>" in data or b"<xdr:sp " in data or b"<xdr:cxnSp" in data
                    or b"<xdr:grpSp" in data, "shapes / text boxes", blocking)
            elif n.startswith("xl/media/") and n.lower().endswith(".wmf"):
                add(True, "WMF images", blocking)
            elif n.startswith("xl/worksheets/") and n.endswith(".xml") and "/_rels/" not in n:
                data = z.read(n)
                if b"<extLst" in data:
                    add(b"sparklineGroup" in data, "sparklines", blocking)
                    add(b"x14:conditionalFormatting" in data,
                        "Excel 2010+ conditional formats (e.g. icon sets/data bar extensions)",
                        blocking)
                    add(b"x14:dataValidation" in data,
                        "data validations that refer to other sheets (x14 format)", blocking)
    return blocking, notices


class _RawImage(_OpenpyxlImage):
    """An embedded picture kept as raw bytes, so images survive without Pillow."""

    def __init__(self, fp):
        data = fp.read()
        self.ref = None
        self._raw = data
        self.width, self.height = 100, 100
        if data.startswith(b"\x89PNG"):
            self.format = "png"
            self.width = int.from_bytes(data[16:20], "big")
            self.height = int.from_bytes(data[20:24], "big")
        elif data.startswith(b"\xff\xd8"):
            self.format = "jpeg"
        elif data[:4] in (b"GIF8",):
            self.format = "gif"
            self.width = int.from_bytes(data[6:8], "little")
            self.height = int.from_bytes(data[8:10], "little")
        elif data.startswith(b"BM"):
            self.format = "bmp"
        elif data[40:44] == b" EMF":
            self.format = "emf"
        elif data.startswith(b"\xd7\xcd\xc6\x9a"):
            self.format = "wmf"
        else:
            self.format = "png"

    def _data(self):
        return self._raw


_LOAD_LOCK = threading.Lock()


@contextlib.contextmanager
def _keep_images_without_pil():
    if _PIL_AVAILABLE:
        yield
        return
    import openpyxl.reader.drawings as rd
    with _LOAD_LOCK:
        old = rd.PILImage, rd.Image
        rd.PILImage, rd.Image = True, _RawImage
        try:
            yield
        finally:
            rd.PILImage, rd.Image = old


# --------------------------------------------------------------------------
# Book: load, inspect, save with cached formula values
# --------------------------------------------------------------------------
def _check_path(path, must_exist=True) -> Path:
    p = Path(path)
    if p.suffix.lower() not in SUFFIXES:
        raise UserError(f"{p.name}: only .xlsx/.xlsm workbooks are supported", "invalid_file")
    if must_exist and not p.exists():
        raise UserError(f"file not found: {p}", "file_not_found")
    if must_exist and p.is_dir():
        raise UserError(f"{p} is a directory", "invalid_file")
    return p


class Book:
    """An opened workbook plus what is needed to save it faithfully."""

    def __init__(self, path, *, create=False, force=False, for_write=True):
        self.path = _check_path(path, must_exist=not create)
        self.warnings: list[str] = []
        self.created = not self.path.exists()
        self._cache = None
        self._snapshot = {}
        if self.created:
            self.wb = openpyxl.Workbook()
            return
        blocking, notices = preflight(self.path)
        if for_write and blocking and not force:
            raise UserError(
                "this workbook contains content openpyxl cannot preserve: "
                + ", ".join(blocking) + ". Saving would remove it. Re-run with force=true to "
                "accept the loss (use out_path to keep the original untouched).",
                "would_lose_content", lossy_features=blocking)
        if for_write:
            for b in blocking:
                self.warnings.append(f"removed on save: {b}")
            for n in notices:
                self.warnings.append(f"not preserved: {n}")
        self.lossy, self.notices = blocking, notices
        with _keep_images_without_pil(), warnings.catch_warnings(record=True) as caught:
            warnings.simplefilter("always")
            self.wb = openpyxl.load_workbook(self.path, keep_vba=self.path.suffix.lower() == ".xlsm",
                                             keep_links=True, rich_text=True)
        if for_write:
            for msg in dict.fromkeys(str(w.message) for w in caught):
                self.warnings.append(f"openpyxl: {msg}")
        for ws in self.wb.worksheets:
            for (r, c), cell in ws._cells.items():
                if cell.data_type == "f":
                    self._snapshot[(ws.title, r, c)] = _formula_sig(cell.value)

    # -- sheets -------------------------------------------------------------
    def ws(self, name=None) -> Worksheet:
        return get_ws(self.wb, name)

    # -- values ---------------------------------------------------------------
    def cached_values(self):
        """Cached (last-calculated) values of formula cells, from the file on disk."""
        if self._cache is None:
            self._cache = {}
            if not self.created and self._snapshot:
                with _keep_images_without_pil():
                    dwb = openpyxl.load_workbook(self.path, data_only=True)
                for title, r, c in self._snapshot:
                    try:
                        v = dwb[title]._cells.get((r, c))
                    except KeyError:
                        continue
                    if v is not None and v.value is not None:
                        self._cache[(title, r, c)] = v.value
                dwb.close()
        return self._cache

    def evaluator(self):
        fallback = {}
        if self._snapshot:
            for k, v in self.cached_values().items():
                fallback[k] = _to_eval_value(self.wb, v)
        return fx.Evaluator(self.wb, fallback=fallback)

    # -- saving -------------------------------------------------------------
    def save(self, out_path=None) -> dict:
        dest = _check_path(out_path, must_exist=False) if out_path else self.path
        if not dest.parent.exists():
            raise UserError(f"folder does not exist: {dest.parent}", "file_not_found")
        self.wb.calculation.fullCalcOnLoad = True
        report, values, unsupported = fx.verify(self.wb)
        kept = 0
        if unsupported and self._snapshot:
            cache = self.cached_values()
            for key in unsupported:
                cell = self.wb[key[0]]._cells.get((key[1], key[2]))
                if cell is not None and self._snapshot.get(key) == _formula_sig(cell.value) \
                        and key in cache:
                    values[key] = cache[key]
                    kept += 1
        fd, tmp1 = tempfile.mkstemp(suffix=dest.suffix, prefix=".xlsxops-", dir=dest.parent)
        os.close(fd)
        fd, tmp2 = tempfile.mkstemp(suffix=dest.suffix, prefix=".xlsxops-", dir=dest.parent)
        os.close(fd)
        try:
            self.wb.save(tmp1)
            _inject_cached_values(tmp1, tmp2, values)
            os.replace(tmp2, dest)
        finally:
            for t in (tmp1, tmp2):
                with contextlib.suppress(OSError):
                    os.remove(t)
        summary = {k: report[k] for k in ("formulas", "evaluated", "error_count",
                                          "circular_count", "unsupported_count")}
        summary["kept_previous_cached_value"] = kept
        for e in report["errors"][:5]:
            self.warnings.append(f"formula error at {e['cell']}: {e['error']} ({e['formula']})")
        if report["error_count"] > 5:
            self.warnings.append(f"{report['error_count'] - 5} more formula errors; "
                                 "run verify_formulas for the full list")
        for e in report["circular"][:5]:
            self.warnings.append(f"circular reference at {e['cell']}")
        if report["unsupported_count"]:
            self.warnings.append(
                f"{report['unsupported_count']} formula(s) not evaluated by xlsxops "
                f"({kept} kept their previous cached value); Excel recalculates them on open")
        self.saved_to = dest
        self.last_values = values
        return summary


def _formula_sig(v):
    """Comparable identity of a formula cell's content (plain or array formula)."""
    if isinstance(v, str):
        return v
    return (type(v).__name__, str(getattr(v, "ref", "")), str(getattr(v, "text", v)))


def _to_eval_value(wb, v):
    if isinstance(v, (_dt.datetime, _dt.date, _dt.time)):
        from openpyxl.utils.datetime import to_excel
        return to_excel(v, epoch=wb.epoch)
    if isinstance(v, str) and v in fx.ERROR_CODES:
        return fx.XLError(v)
    return v


_ILLEGAL_XML = re.compile("[\x00-\x08\x0b\x0c\x0e-\x1f]")
_FCELL = re.compile(r'<c r="([A-Z]{1,3}[0-9]+)"([^>]*)>(<f(?:\s[^>]*)?>[^<]*</f>)(?:<v\s*/>|<v></v>)</c>')


def _xml_escape(s):
    s = _ILLEGAL_XML.sub("", s)
    return s.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


def _sheet_parts(z: zipfile.ZipFile) -> dict[str, str]:
    ns = {"m": "http://schemas.openxmlformats.org/spreadsheetml/2006/main",
          "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
          "p": "http://schemas.openxmlformats.org/package/2006/relationships"}
    wbx = ET.fromstring(z.read("xl/workbook.xml"))
    rels = ET.fromstring(z.read("xl/_rels/workbook.xml.rels"))
    targets = {}
    for rel in rels.findall("p:Relationship", ns):
        t = rel.get("Target")
        t = t.lstrip("/") if t.startswith("/") else "xl/" + t
        targets[rel.get("Id")] = t
    out = {}
    for sh in wbx.findall("m:sheets/m:sheet", ns):
        rid = sh.get(f"{{{ns['r']}}}id")
        if rid in targets:
            out[sh.get("name")] = targets[rid]
    return out


def _inject_cached_values(src, dst, values):
    """Copy the package, filling empty <v/> of formula cells with computed values."""
    by_sheet: dict[str, dict[str, object]] = {}
    for (title, r, c), v in values.items():
        by_sheet.setdefault(title, {})[f"{get_column_letter(c)}{r}"] = v
    with zipfile.ZipFile(src) as zin:
        parts = {part: by_sheet[name] for name, part in _sheet_parts(zin).items() if name in by_sheet}
        with zipfile.ZipFile(dst, "w", zipfile.ZIP_DEFLATED) as zout:
            for item in zin.infolist():
                data = zin.read(item.filename)
                vals = parts.get(item.filename)
                if vals:
                    data = _patch_sheet(data.decode("utf-8"), vals).encode("utf-8")
                zout.writestr(item, data)


def _patch_sheet(xml: str, vals: dict) -> str:
    def sub(m):
        ref, attrs, f = m.group(1), m.group(2), m.group(3)
        if ref not in vals:
            return m.group(0)
        v = vals[ref]
        attrs = re.sub(r'\s+t="[^"]*"', "", attrs)
        if isinstance(v, fx.XLError):
            t, text = "e", v.code
        elif isinstance(v, bool):
            t, text = "b", "1" if v else "0"
        elif isinstance(v, (int, float)):
            if isinstance(v, float) and not math.isfinite(v):
                return m.group(0)
            t, text = None, repr(v) if isinstance(v, float) else str(v)
        elif isinstance(v, (_dt.datetime, _dt.date)):
            return m.group(0)
        elif v is None:
            return m.group(0)
        else:
            t, text = "str", _xml_escape(str(v))
        tattr = f' t="{t}"' if t else ""
        return f'<c r="{ref}"{attrs}{tattr}>{f}<v>{text}</v></c>'
    return _FCELL.sub(sub, xml)


# --------------------------------------------------------------------------
# Sheet / range helpers
# --------------------------------------------------------------------------
def get_ws(wb, name=None) -> Worksheet:
    if name is None:
        ws = wb.active
        if isinstance(ws, Worksheet):
            return ws
        if wb.worksheets:
            return wb.worksheets[0]
        raise UserError("workbook has no worksheets")
    for ws in wb.worksheets:
        if ws.title == name:
            return ws
    for ws in wb.worksheets:
        if ws.title.casefold() == str(name).casefold():
            return ws
    raise UserError(f"sheet {name!r} not found; sheets: {wb.sheetnames}", "sheet_not_found")


_BAD_SHEET_CHARS = re.compile(r"[\[\]:*?/\\]")


def check_new_sheet_name(wb, name: str) -> str:
    if not name or not name.strip():
        raise UserError("sheet name is empty")
    if len(name) > 31:
        raise UserError(f"sheet name {name!r} is longer than 31 characters")
    if _BAD_SHEET_CHARS.search(name) or name.startswith("'") or name.endswith("'"):
        raise UserError(f"sheet name {name!r} contains a character Excel forbids ([]:*?/\\ or edge ')")
    if name.casefold() == "history":
        raise UserError("'History' is a reserved sheet name in Excel")
    if wb is not None and any(s.casefold() == name.casefold() for s in wb.sheetnames):
        raise UserError(f"a sheet named {name!r} already exists", "sheet_exists")
    return name


def safe_sheet_name(wb, raw: str, taken: set | None = None) -> str:
    """Sanitise arbitrary text into a unique valid sheet name."""
    base = _BAD_SHEET_CHARS.sub("_", str(raw)).strip().strip("'") or "Sheet"
    if base.casefold() == "history":
        base = "History_"
    base = base[:31]
    taken = {s.casefold() for s in wb.sheetnames} | (taken or set())
    name, i = base, 2
    while name.casefold() in taken:
        suffix = f" ({i})"
        name = base[:31 - len(suffix)] + suffix
        i += 1
    return name


def parse_range(ref: str):
    """'A1:C10' -> (min_col, min_row, max_col, max_row); single cells allowed."""
    try:
        ref = ref.replace("$", "").strip()
        if "!" in ref:
            raise UserError(f"range {ref!r}: give the sheet separately via 'sheet'")
        c1, r1, c2, r2 = range_boundaries(ref)
    except (ValueError, TypeError):
        raise UserError(f"invalid range {ref!r}; expected something like 'B2:D20'")
    if None in (c1, r1, c2, r2):
        raise UserError(f"range {ref!r} must have explicit rows and columns, e.g. 'B2:B20'")
    return c1, r1, c2, r2


def parse_cell(ref: str):
    c1, r1, c2, r2 = parse_range(ref)
    if (c1, r1) != (c2, r2):
        raise UserError(f"{ref!r} is a range; a single cell like 'B2' is required")
    return r1, c1


def cell_get(ws, r, c):
    """Read a cell without creating it."""
    return ws._cells.get((r, c))


def cell_value(ws, r, c):
    cell = ws._cells.get((r, c))
    return None if cell is None else cell.value


def is_formula(cell) -> bool:
    return cell is not None and cell.data_type == "f"


def to_number_like(v):
    """Comparable form: dates -> Excel serials; others unchanged."""
    if isinstance(v, (_dt.datetime, _dt.date)):
        from openpyxl.utils.datetime import to_excel
        return to_excel(v)
    return v


# --------------------------------------------------------------------------
# Header-row tables
# --------------------------------------------------------------------------
_TOTAL_RX = re.compile(
    r"^=\s*(?:SUM|AVERAGE|COUNT|COUNTA|MIN|MAX|SUBTOTAL)\(\s*(?:\d+\s*,\s*)?\$?([A-Z]{1,3})\$?(\d+):\$?\1\$?(\d+)\s*\)$",
    re.IGNORECASE)


@dataclass
class TableInfo:
    ws: Worksheet
    header_row: int
    first_col: int
    last_col: int
    headers: dict = field(default_factory=dict)  # col -> header text
    last_data: int = 0
    total_row: int | None = None
    trailing_rows: int = 0

    @property
    def first_data(self):
        return self.header_row + 1

    @property
    def data_rows(self):
        return range(self.first_data, self.last_data + 1)

    @property
    def n_rows(self):
        return max(0, self.last_data - self.header_row)

    def header_list(self):
        return [self.headers.get(c) for c in range(self.first_col, self.last_col + 1)]

    def col(self, spec) -> int:
        """Resolve a header name (case-insensitive) or column letter to an index."""
        if isinstance(spec, int):
            return spec
        s = str(spec).strip()
        for c, h in self.headers.items():
            if h is not None and str(h).strip() == s:
                return c
        for c, h in self.headers.items():
            if h is not None and str(h).strip().casefold() == s.casefold():
                return c
        if re.fullmatch(r"[A-Za-z]{1,3}", s):
            idx = column_index_from_string(s.upper())
            if idx <= fx.MAX_COL:
                return idx
        raise UserError(f"column {spec!r} not found; headers: "
                        f"{[h for h in self.header_list() if h is not None]}", "column_not_found")

    def name(self, c) -> str:
        h = self.headers.get(c)
        return str(h) if h is not None else get_column_letter(c)


def find_table(ws, header_row=None) -> TableInfo:
    """Locate the header row and the contiguous data block below it."""
    cells = ws._cells
    if header_row is None:
        rows = sorted({r for (r, c), cell in cells.items() if cell.value is not None})
        if not rows:
            raise UserError(f"sheet {ws.title!r} is empty", "empty_sheet")
        header_row = rows[0]
    hdr = {c: cell.value for (r, c), cell in cells.items() if r == header_row and cell.value is not None}
    if not hdr:
        raise UserError(f"row {header_row} of {ws.title!r} has no headers", "no_header")
    first_col, last_col = min(hdr), max(hdr)
    filled = {}
    for (r, c), cell in cells.items():
        if r > header_row and first_col <= c <= last_col and cell.value is not None:
            filled[r] = filled.get(r, 0) + 1
    last = header_row
    while (last + 1) in filled:
        last += 1
    info = TableInfo(ws, header_row, first_col, last_col, hdr, last)
    info.trailing_rows = sum(1 for r in filled if r > last + 1)
    # A total row written under the table (SUM(B2:B9) in row 10) is not data.
    if last > header_row + 1:
        for c in range(first_col, last_col + 1):
            cell = cells.get((last, c))
            if is_formula(cell) and isinstance(cell.value, str):
                m = _TOTAL_RX.match(cell.value.replace(" ", ""))
                if m and int(m.group(2)) == header_row + 1 and last - 2 <= int(m.group(3)) < last:
                    info.total_row = last
                    info.last_data = int(m.group(3))
                    break
    return info


def table_warnings(info: TableInfo) -> list[str]:
    w = []
    if info.trailing_rows:
        w.append(f"{info.trailing_rows} non-empty row(s) below a blank row after row "
                 f"{info.last_data if info.total_row is None else info.total_row} were not treated "
                 "as part of the table")
    return w


# --------------------------------------------------------------------------
# Row snapshots (used by sort / dedupe)
# --------------------------------------------------------------------------
def snapshot_row(ws, r, c1, c2):
    snap = []
    for c in range(c1, c2 + 1):
        cell = ws._cells.get((r, c))
        if cell is None:
            snap.append(None)
            continue
        snap.append((cell._value, cell.data_type, copy(cell._style), cell.hyperlink, cell.comment))
    height = ws.row_dimensions[r].height if r in ws.row_dimensions else None
    return snap, height


def write_row(ws, r, snapshot, c1, src_row):
    snap, height = snapshot
    for i, item in enumerate(snap):
        c = c1 + i
        if item is None:
            ws._cells.pop((r, c), None)
            continue
        value, dtype, style, link, comment = item
        cell = ws.cell(row=r, column=c)
        if dtype == "f" and isinstance(value, str) and src_row != r:
            value = translate(value, c, src_row, c, r)
        cell._value = value
        cell.data_type = dtype
        cell._style = copy(style)
        cell.hyperlink = copy(link) if link is not None else None
        cell.comment = copy(comment) if comment is not None else None
    if height is not None:
        ws.row_dimensions[r].height = height
    elif r in ws.row_dimensions:
        ws.row_dimensions[r].height = None


def translate(formula: str, c0, r0, c1, r1) -> str:
    """Shift relative references the way Excel's copy/paste does."""
    try:
        return Translator(formula, origin=f"{get_column_letter(c0)}{r0}").translate_formula(
            f"{get_column_letter(c1)}{r1}")
    except Exception:
        return formula


def copy_style(src, dst, number_format=True):
    if src is None or not src.has_style:
        return
    if number_format:
        dst._style = copy(src._style)
    else:
        nf = dst.number_format
        dst._style = copy(src._style)
        dst.number_format = nf


def bolded(font):
    """A copy of a cell's font with bold switched on (other attributes kept)."""
    f = copy(font)
    f.b = True
    return f


def finish(book: Book, out_path, details: dict, warnings=None, changed=True):
    """Save, attach the formula check, and build the (changed, warnings, details) tuple."""
    details = dict(details)
    details["formula_check"] = book.save(out_path)
    details["path"] = str(book.saved_to)
    return changed, list(warnings or []) + book.warnings, details


# Common schema fragments (kept short: every schema must fit in 600 chars).
S_PATH = {"type": "string"}
S_SHEET = {"type": "string"}
S_OUT = {"type": "string"}
S_FORCE = {"type": "boolean"}
S_HDR = {"type": "integer"}
S_COLS = {"type": "array", "items": {"type": "string"}, "minItems": 1}


# --------------------------------------------------------------------------
# Value helpers shared by operations
# --------------------------------------------------------------------------
_ISO_DT = re.compile(r"(\d{4})-(\d{2})-(\d{2})(?:[T ](\d{2}):(\d{2})(?::(\d{2}))?)?")


def coerce_input(v, infer_dates=True):
    """Turn a JSON value from the agent into what goes into a cell.

    Returns (value, kind) where kind is 'formula', 'text', 'date', 'number', 'bool', 'blank'.
    """
    if v is None:
        return None, "blank"
    if isinstance(v, bool):
        return v, "bool"
    if isinstance(v, (int, float)):
        if isinstance(v, float) and not math.isfinite(v):
            raise UserError("NaN/Infinity cannot be written to Excel")
        return v, "number"
    if isinstance(v, str):
        if v.startswith("'"):
            return v[1:], "text"
        if v.startswith("=") and len(v) > 1:
            err = fx.check_syntax(v)
            if err:
                raise UserError(f"formula {v!r} has a syntax error: {err}", "formula_syntax")
            return fx.add_future_prefixes(v), "formula"
        if infer_dates:
            m = _ISO_DT.fullmatch(v.strip())
            if m:
                try:
                    if m.group(4) is None:
                        return _dt.date(int(m[1]), int(m[2]), int(m[3])), "date"
                    return _dt.datetime(int(m[1]), int(m[2]), int(m[3]), int(m[4]), int(m[5]),
                                        int(m[6] or 0)), "date"
                except ValueError:
                    pass
        return v, "text"
    raise UserError(f"unsupported value type {type(v).__name__}")


def put(cell, value, kind):
    """Write a coerced value; text that looks like a formula stays text."""
    from openpyxl.cell.cell import MergedCell
    if isinstance(cell, MergedCell):
        raise UserError(f"{cell.coordinate} is inside a merged range; write to its top-left cell",
                        "merged_cell")
    cell.value = value
    if kind == "text" and isinstance(value, str):
        cell.data_type = "s"
    if kind == "date" and cell.number_format == "General":
        cell.number_format = "yyyy-mm-dd" if not isinstance(value, _dt.datetime) else "yyyy-mm-dd hh:mm"


class ValueReader:
    """Computed cell values for operations (formulas evaluated, cache fallback)."""

    def __init__(self, book: Book):
        self.book = book
        self.ev = book.evaluator()
        self.unknown = 0

    def get(self, ws, r, c, raw=False):
        cell = ws._cells.get((r, c))
        if cell is None:
            return None
        if cell.data_type != "f":
            if cell.data_type == "e":
                return fx.XLError(str(cell.value)) if str(cell.value) in fx.ERROR_CODES else cell.value
            return cell.value if raw else self.ev.convert_constant(cell)
        try:
            return self.ev.cell_value(ws.title, r, c)
        except (fx.Unsupported, fx._DependsOnCircular, fx._Circular):
            self.unknown += 1
            return None
        except RecursionError:
            self.unknown += 1
            return None

    def warnings(self):
        if self.unknown:
            return [f"{self.unknown} formula value(s) could not be computed and were treated as blank"]
        return []


def out_value(v):
    """A computed value as something a cell can hold."""
    if isinstance(v, fx.XLError):
        return v.code
    return v


def norm_key(v, ignore_case=True, trim=True):
    """Normalise a value for grouping / matching (Excel-like: case-insensitive)."""
    if v is None:
        return ("blank",)
    if isinstance(v, fx.XLError):
        return ("err", v.code)
    if isinstance(v, bool):
        return ("bool", v)
    if isinstance(v, (_dt.datetime, _dt.date)):
        return ("num", float(to_number_like(v)))
    if isinstance(v, (int, float)):
        return ("num", float(v))
    s = str(v)
    if trim:
        s = s.strip()
    if s == "":
        return ("blank",)
    return ("text", s.casefold() if ignore_case else s)


def display_width(cell) -> float:
    """Estimated display width (characters) of a cell's content."""
    v = cell.value
    if v is None:
        return 0
    if cell.data_type == "f":
        return 10
    if isinstance(v, (_dt.datetime, _dt.date, _dt.time)) or (
            cell.is_date and isinstance(v, (int, float))):
        fmt = cell.number_format or "yyyy-mm-dd"
        n = len(re.sub(r'"[^"]*"|\\|\[[^\]]*\]', "", fmt))
        return max(n, 8)
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        try:
            return len(fx.format_text(v, cell.number_format or "General"))
        except Exception:
            return len(fx.general_text(v))
    text = str(v)
    return max((len(line) for line in text.split("\n")), default=0)


def autofit_columns(ws, cols, rows=None, min_width=6, max_width=80):
    """Set column widths from an estimate of the longest displayed value."""
    widths = {}
    for (r, c), cell in ws._cells.items():
        if c not in cols or (rows is not None and r not in rows) or cell.value is None:
            continue
        w = display_width(cell)
        size = (cell.font.sz or 11) if cell.has_style else 11
        bold = cell.has_style and cell.font.b
        w = w * (size / 11) * (1.1 if bold else 1.0)
        widths[c] = max(widths.get(c, 0), w)
    out = {}
    for c in cols:
        w = round(min(max(widths.get(c, 0) + 2, min_width), max_width), 1)
        ws.column_dimensions[get_column_letter(c)].width = w
        out[get_column_letter(c)] = w
    return out


def append_column(ws, info: TableInfo, header, values, number_format=None, col=None):
    """Write a new column right of the table (or at an empty ``col``).

    ``values`` are (value, kind) pairs, one per data row. Returns the column index.
    """
    from openpyxl.worksheet.table import TableColumn
    if any(h is not None and str(h).strip().casefold() == str(header).strip().casefold()
           for h in info.headers.values()):
        raise UserError(f"a column named {header!r} already exists", "column_exists")
    if col is None:
        col = info.last_col + 1
    rows = [info.header_row] + list(info.data_rows)
    if info.total_row:
        rows.append(info.total_row)
    busy = [f"{get_column_letter(col)}{r}" for r in rows
            if ws._cells.get((r, col)) is not None and ws._cells[(r, col)].value is not None]
    if busy:
        raise UserError(f"column {get_column_letter(col)} is not empty ({busy[0]}); "
                        "pick an empty column", "column_not_empty")
    left = col - 1 if col > 1 else None
    h = ws.cell(row=info.header_row, column=col)
    if left:
        copy_style(ws._cells.get((info.header_row, left)), h)
    put(h, str(header), "text")
    for r, (v, kind) in zip(info.data_rows, values):
        cell = ws.cell(row=r, column=col)
        if left:
            copy_style(ws._cells.get((r, left)), cell, number_format=False)
        put(cell, v, kind)
        if number_format:
            cell.number_format = number_format
    if left and get_column_letter(left) in ws.column_dimensions:
        width = ws.column_dimensions[get_column_letter(left)].width
    else:
        width = None
    est = max(len(str(header)) * 1.1 + 2, 8)
    ws.column_dimensions[get_column_letter(col)].width = max(width or 0, est) if width else est
    # Grow Excel tables and the autofilter that end at the column on the left.
    for t in ws.tables.values():
        c1, r1, c2, r2 = range_boundaries(t.ref)
        if r1 == info.header_row and c2 == col - 1:
            t.ref = f"{get_column_letter(c1)}{r1}:{get_column_letter(col)}{r2}"
            ids = [tc.id for tc in t.tableColumns] or [0]
            t.tableColumns.append(TableColumn(id=max(ids) + 1, name=str(header)))
            if t.autoFilter is not None and t.autoFilter.ref:
                t.autoFilter.ref = t.ref
    if ws.auto_filter.ref:
        c1, r1, c2, r2 = range_boundaries(ws.auto_filter.ref)
        if r1 == info.header_row and c2 == col - 1:
            ws.auto_filter.ref = f"{get_column_letter(c1)}{r1}:{get_column_letter(col)}{r2}"
    info.headers[col] = str(header)
    info.last_col = max(info.last_col, col)
    return col
