# SPDX-License-Identifier: Apache-2.0
"""Reading, cell writes, sheets, columns and formulas."""
from __future__ import annotations

import re

from openpyxl.styles import Font
from openpyxl.utils.cell import get_column_letter

from . import formula as fx
from ._core import (bolded, S_COLS, S_FORCE, S_HDR, S_OUT, S_PATH, S_SHEET, Book, UserError,
                    ValueReader, append_column, check_new_sheet_name, coerce_input, copy_style,
                    find_table, finish, operation, parse_cell, parse_range, put,
                    table_warnings, translate)

_MAX_WRITE = 100_000


@operation("read_summary", {
    "type": "object", "required": ["path"],
    "properties": {"path": S_PATH, "sheet": S_SHEET,
                   "sample_rows": {"type": "integer", "minimum": 0, "maximum": 50},
                   "header_row": S_HDR},
})
def read_summary(path, sheet=None, sample_rows=5, header_row=None):
    """Sheets, used ranges, headers, row counts, sample rows and formula counts."""
    book = Book(path, for_write=False)
    wb = book.wb
    reader = ValueReader(book)
    sheets = []
    targets = [book.ws(sheet)] if sheet else wb.worksheets
    for ws in targets:
        vals = [(r, c) for (r, c), cell in ws._cells.items() if cell.value is not None]
        entry = {"name": ws.title, "state": ws.sheet_state,
                 "formulas_count": sum(1 for cell in ws._cells.values() if cell.data_type == "f"),
                 "charts": len(ws._charts), "images": len(ws._images),
                 "tables": [{"name": t.displayName, "ref": t.ref} for t in ws.tables.values()],
                 "merged_ranges": [str(m) for m in ws.merged_cells.ranges][:20],
                 "freeze_panes": ws.freeze_panes,
                 "data_validations": len(ws.data_validations.dataValidation),
                 "conditional_formats": sum(len(cf.rules) for cf in ws.conditional_formatting)}
        if not vals:
            entry.update(used_range=None, header_row=None, headers=[], row_count=0, sample_rows=[])
            sheets.append(entry)
            continue
        r1, r2 = min(r for r, _ in vals), max(r for r, _ in vals)
        c1, c2 = min(c for _, c in vals), max(c for _, c in vals)
        entry["used_range"] = f"{get_column_letter(c1)}{r1}:{get_column_letter(c2)}{r2}"
        try:
            info = find_table(ws, header_row)
        except UserError:
            info = None
        if info is not None:
            entry["header_row"] = info.header_row
            entry["headers"] = info.header_list()
            entry["header_columns"] = f"{get_column_letter(info.first_col)}:{get_column_letter(info.last_col)}"
            entry["row_count"] = info.n_rows
            entry["total_row"] = info.total_row
            rows = []
            for r in list(info.data_rows)[:sample_rows]:
                rows.append([_sample(reader, ws, r, c) for c in range(info.first_col, info.last_col + 1)])
            entry["sample_rows"] = rows
            entry["number_formats"] = {
                info.name(c): ws._cells[(info.first_data, c)].number_format
                for c in range(info.first_col, info.last_col + 1)
                if (info.first_data, c) in ws._cells
                and ws._cells[(info.first_data, c)].number_format != "General"}
            if info.trailing_rows:
                entry["rows_below_table"] = info.trailing_rows
        sheets.append(entry)
    details = {
        "path": str(book.path),
        "sheets": sheets,
        "sheet_names": wb.sheetnames,
        "active_sheet": wb.active.title if wb.active is not None else None,
        "chartsheets": [cs.title for cs in wb.chartsheets],
        "defined_names": [{"name": n, "value": d.attr_text} for n, d in wb.defined_names.items()],
    }
    warnings = []
    if book.lossy:
        details["content_lost_on_write"] = book.lossy
        warnings.append("write operations on this file need force=true: " + ", ".join(book.lossy))
    if book.notices:
        details["not_preserved_on_write"] = book.notices
    return False, warnings + reader.warnings(), details


def _sample(reader, ws, r, c):
    cell = ws._cells.get((r, c))
    if cell is None:
        return None
    if cell.data_type == "f":
        v = reader.get(ws, r, c)
        return {"formula": cell.value if isinstance(cell.value, str) else "array formula",
                "value": fx.general_text(v) if isinstance(v, fx.XLError) else v}
    return cell.value


@operation("write_cells", {
    "type": "object", "required": ["path"],
    "properties": {"path": S_PATH, "sheet": S_SHEET,
                   "cells": {"type": "object", "description": "{\"B2\": value}; '=..' is a formula"},
                   "start": {"type": "string"},
                   "values": {"type": "array", "items": {"type": "array"}},
                   "out_path": S_OUT, "force": S_FORCE},
})
def write_cells(path, sheet=None, cells=None, start=None, values=None, out_path=None, force=False):
    """Write values/formulas to cells: a {ref: value} map and/or a 2-D block at start."""
    if not cells and values is None:
        raise UserError("give 'cells' ({\"A1\": value}) or 'start' + 'values' (2-D list)")
    if values is not None and not start:
        start = "A1"
    book = Book(path, force=force)
    ws = book.ws(sheet)
    plan = []
    for ref, v in (cells or {}).items():
        r, c = parse_cell(ref)
        plan.append((r, c, v))
    if values is not None:
        r0, c0 = parse_cell(start)
        for i, row in enumerate(values):
            if not isinstance(row, list):
                raise UserError("'values' must be a list of rows (lists)")
            for j, v in enumerate(row):
                plan.append((r0 + i, c0 + j, v))
    if len(plan) > _MAX_WRITE:
        raise UserError(f"too many cells in one call ({len(plan)} > {_MAX_WRITE})")
    coerced = [(r, c, *coerce_input(v)) for r, c, v in plan]
    previous = []
    for r, c, v, kind in coerced:
        old = ws._cells.get((r, c))
        if old is not None and old.value is not None and len(previous) < 20:
            previous.append({"cell": f"{get_column_letter(c)}{r}", "old": old.value})
        put(ws.cell(row=r, column=c), v, kind)
    details = {"sheet": ws.title, "cells_written": len(coerced),
               "formulas_written": sum(1 for x in coerced if x[3] == "formula"),
               "overwritten": previous}
    result = finish(book, out_path, details)
    vals = book.last_values
    formula_values = {f"{get_column_letter(c)}{r}": vals.get((ws.title, r, c))
                      for r, c, v, kind in coerced if kind == "formula"}
    if formula_values:
        result[2]["formula_values"] = dict(list(formula_values.items())[:50])
    return result


@operation("add_sheet", {
    "type": "object", "required": ["path", "name"],
    "properties": {"path": S_PATH, "name": {"type": "string", "minLength": 1, "maxLength": 31},
                   "index": {"type": "integer", "minimum": 0},
                   "headers": {"type": "array", "items": {"type": ["string", "number"]}},
                   "out_path": S_OUT, "force": S_FORCE},
})
def add_sheet(path, name, index=None, headers=None, out_path=None, force=False):
    """Add a worksheet (creates the workbook if the file does not exist)."""
    book = Book(path, create=True, force=force)
    wb = book.wb
    if book.created:
        ws = wb.active
        ws.title = check_new_sheet_name(None, name)
    else:
        check_new_sheet_name(wb, name)
        ws = wb.create_sheet(title=name, index=index)
    if headers:
        for j, h in enumerate(headers, 1):
            cell = ws.cell(row=1, column=j)
            put(cell, h, "text" if isinstance(h, str) else "number")
            cell.font = Font(bold=True)
            ws.column_dimensions[get_column_letter(j)].width = max(len(str(h)) * 1.1 + 2, 8)
        ws.freeze_panes = "A2"
    details = {"sheet": ws.title, "index": wb.sheetnames.index(ws.title),
               "sheets": wb.sheetnames, "created_workbook": book.created}
    return finish(book, out_path, details)


@operation("add_column", {
    "type": "object", "required": ["path", "header"],
    "properties": {"path": S_PATH, "sheet": S_SHEET, "header": {"type": "string", "minLength": 1},
                   "values": {"type": "array"}, "fill": {},
                   "column": {"type": "string", "pattern": "^[A-Za-z]{1,3}$"},
                   "number_format": {"type": "string"}, "header_row": S_HDR,
                   "out_path": S_OUT, "force": S_FORCE},
})
def add_column(path, header, sheet=None, values=None, fill=None, column=None, number_format=None,
               header_row=None, out_path=None, force=False):
    """Append a column after the table (values per data row, or one fill value)."""
    book = Book(path, force=force)
    ws = book.ws(sheet)
    info = find_table(ws, header_row)
    n = info.n_rows
    warnings = table_warnings(info)
    if values is not None and fill is not None:
        raise UserError("give either 'values' or 'fill', not both")
    from openpyxl.utils.cell import column_index_from_string
    col = column_index_from_string(column.upper()) if column else info.last_col + 1
    if info.first_col <= col <= info.last_col:
        raise UserError(f"column {column} is inside the table; inserting columns is not "
                        "supported because it would break references. Use an empty column.")
    if values is not None:
        if len(values) > n:
            raise UserError(f"'values' has {len(values)} items but the table has {n} data rows")
        if len(values) < n:
            warnings.append(f"'values' has {len(values)} items for {n} rows; the rest are blank")
        items = [coerce_input(v) for v in values] + [(None, "blank")] * (n - len(values))
    else:
        v, kind = coerce_input(fill)
        if kind == "formula":
            items = [(translate(v, col, info.first_data, col, r), kind) for r in info.data_rows]
        else:
            items = [(v, kind)] * n
    col = append_column(ws, info, header, items, number_format, col)
    details = {"sheet": ws.title, "column": get_column_letter(col), "header": header, "rows": n}
    return finish(book, out_path, details, warnings)


_PLACEHOLDER = re.compile(r"\[@(?:\[([^\]]+)\]|([^\[\]]+))\]")


@operation("calculated_column", {
    "type": "object", "required": ["path", "header", "formula"],
    "properties": {"path": S_PATH, "sheet": S_SHEET, "header": {"type": "string", "minLength": 1},
                   "formula": {"type": "string", "description":
                               "for the first data row, e.g. =B2*C2 or =[@Qty]*[@Price]"},
                   "number_format": {"type": "string"}, "header_row": S_HDR,
                   "out_path": S_OUT, "force": S_FORCE},
})
def calculated_column(path, header, formula, sheet=None, number_format=None, header_row=None,
                      out_path=None, force=False):
    """Add (or refill) a column with one formula per data row, relative refs shifted."""
    book = Book(path, force=force)
    ws = book.ws(sheet)
    info = find_table(ws, header_row)
    warnings = table_warnings(info)
    if info.n_rows == 0:
        raise UserError("the table has no data rows")
    if not formula.startswith("="):
        formula = "=" + formula
    first = info.first_data

    def resolve(m):
        name = m.group(1) or m.group(2)
        return f"{get_column_letter(info.col(name))}{{row}}"

    uses_names = bool(_PLACEHOLDER.search(formula))
    template = _PLACEHOLDER.sub(resolve, formula) if uses_names else None
    probe = template.replace("{row}", str(first)) if uses_names else formula
    err = fx.check_syntax(probe)
    if err:
        raise UserError(f"formula {probe!r} has a syntax error: {err}", "formula_syntax")
    existing = None
    for c, h in info.headers.items():
        if h is not None and str(h).strip().casefold() == header.strip().casefold():
            existing = c
    target = existing or info.last_col + 1

    def row_formula(r):
        if uses_names:
            f = template.replace("{row}", str(r))
        else:
            f = translate(formula, target, first, target, r)
        return fx.add_future_prefixes(f)

    items = [(row_formula(r), "formula") for r in info.data_rows]
    if existing:
        warnings.append(f"column {header!r} already existed; its {info.n_rows} data cells were "
                        "replaced with the formula")
        for r, (f, kind) in zip(info.data_rows, items):
            cell = ws.cell(row=r, column=existing)
            put(cell, f, kind)
            if number_format:
                cell.number_format = number_format
        col = existing
    else:
        col = append_column(ws, info, header, items, number_format)
    changed, warns, details = finish(book, out_path, {
        "sheet": ws.title, "column": get_column_letter(col), "header": header,
        "first_formula": items[0][0], "rows": info.n_rows}, warnings)
    vals = [book.last_values.get((ws.title, r, col)) for r in list(info.data_rows)]
    details["sample_values"] = vals[:5]
    errs = [i for i, v in enumerate(vals) if isinstance(v, fx.XLError)]
    if errs:
        warns.append(f"{len(errs)} row(s) evaluate to an error, first at "
                     f"{get_column_letter(col)}{first + errs[0]} ({vals[errs[0]]})")
    return changed, warns, details


@operation("set_formula", {
    "type": "object", "required": ["path", "cell", "formula"],
    "properties": {"path": S_PATH, "sheet": S_SHEET,
                   "cell": {"type": "string", "description": "B2, or B2:B20 to fill down/right"},
                   "formula": {"type": "string"}, "number_format": {"type": "string"},
                   "out_path": S_OUT, "force": S_FORCE},
})
def set_formula(path, cell, formula, sheet=None, number_format=None, out_path=None, force=False):
    """Set a formula in a cell, or fill a range from its top-left cell with relative refs."""
    if not formula.startswith("="):
        formula = "=" + formula
    err = fx.check_syntax(formula)
    if err:
        raise UserError(f"formula {formula!r} has a syntax error: {err}", "formula_syntax")
    formula = fx.add_future_prefixes(formula)
    c1, r1, c2, r2 = parse_range(cell)
    if (r2 - r1 + 1) * (c2 - c1 + 1) > _MAX_WRITE:
        raise UserError("range too large")
    book = Book(path, force=force)
    ws = book.ws(sheet)
    written = []
    for r in range(r1, r2 + 1):
        for c in range(c1, c2 + 1):
            f = formula if (r, c) == (r1, c1) else translate(formula, c1, r1, c, r)
            target = ws.cell(row=r, column=c)
            put(target, f, "formula")
            if number_format:
                target.number_format = number_format
            written.append((r, c, f))
    changed, warns, details = finish(book, out_path, {"sheet": ws.title, "cells": len(written)})
    details["results"] = [{"cell": f"{get_column_letter(c)}{r}", "formula": f,
                           "value": book.last_values.get((ws.title, r, c))}
                          for r, c, f in written[:20]]
    unevaluated = [x for x in details["results"] if (ws.title, *_rc(x["cell"])) not in book.last_values]
    if unevaluated:
        warns.append("value not computed by xlsxops for: "
                     + ", ".join(x["cell"] for x in unevaluated[:10]))
    return changed, warns, details


def _rc(ref):
    r, c = parse_cell(ref)
    return r, c


_TOTAL_FUNCS = ("SUM", "AVERAGE", "COUNT", "COUNTA", "MIN", "MAX")


@operation("total_row", {
    "type": "object", "required": ["path"],
    "properties": {"path": S_PATH, "sheet": S_SHEET, "columns": S_COLS,
                   "func": {"enum": list(_TOTAL_FUNCS)},
                   "funcs": {"type": "object", "additionalProperties": {"enum": list(_TOTAL_FUNCS)}},
                   "label": {"type": "string"}, "header_row": S_HDR,
                   "out_path": S_OUT, "force": S_FORCE},
})
def total_row(path, sheet=None, columns=None, func="SUM", funcs=None, label="Total",
              header_row=None, out_path=None, force=False):
    """Write SUM/AVERAGE/COUNT (etc.) formulas in the row under the table.

    Re-running replaces an existing total row instead of stacking a second one.
    """
    book = Book(path, force=force)
    ws = book.ws(sheet)
    info = find_table(ws, header_row)
    warnings = table_warnings(info)
    if info.n_rows == 0:
        raise UserError("the table has no data rows")
    reader = ValueReader(book)
    funcs = {k: v for k, v in (funcs or {}).items()}
    if columns:
        cols = [info.col(c) for c in columns]
    elif funcs:
        cols = []
    else:
        cols = [c for c in range(info.first_col, info.last_col + 1)
                if _numeric_share(reader, ws, info, c) >= 0.8]
        if not cols:
            raise UserError("no numeric columns found; pass 'columns'")
    for k in funcs:
        c = info.col(k)
        if c not in cols:
            cols.append(c)
    col_func = {c: func for c in cols}
    for k, v in funcs.items():
        col_func[info.col(k)] = v
    row = info.total_row or info.last_data + 1
    replaced = info.total_row is not None
    first, last = info.first_data, info.last_data
    written = {}
    for c in range(info.first_col, info.last_col + 1):
        cell = ws._cells.get((row, c))
        if replaced and cell is not None and c not in col_func:
            cell.value = None
    for c, f in sorted(col_func.items()):
        L = get_column_letter(c)
        cell = ws.cell(row=row, column=c)
        copy_style(ws._cells.get((last, c)), cell)
        put(cell, f"={f}({L}{first}:{L}{last})", "formula")
        cell.font = bolded(cell.font)
        if f in ("COUNT", "COUNTA"):
            cell.number_format = "General"
        written[info.name(c)] = cell.value
    label_col = info.first_col
    if label and label_col not in col_func:
        cell = ws.cell(row=row, column=label_col)
        copy_style(ws._cells.get((last, label_col)), cell, number_format=False)
        put(cell, label, "text")
        cell.font = bolded(cell.font)
    elif label:
        warnings.append("first column is totalled, so the label was not written")
    changed, warns, details = finish(book, out_path, {
        "sheet": ws.title, "row": row, "replaced_existing": replaced, "formulas": written},
        warnings + reader.warnings())
    details["values"] = {info.name(c): book.last_values.get((ws.title, row, c)) for c in col_func}
    return changed, warns, details


def _numeric_share(reader, ws, info, c):
    n = num = 0
    for r in info.data_rows:
        v = reader.get(ws, r, c)
        if v is None or v == "":
            continue
        n += 1
        if isinstance(v, (int, float)) and not isinstance(v, bool):
            cell = ws._cells.get((r, c))
            if cell is not None and cell.is_date:
                continue
            num += 1
    return num / n if n else 0


@operation("verify_formulas", {
    "type": "object", "required": ["path"],
    "properties": {"path": S_PATH, "write_cache": {"type": "boolean"},
                   "limit": {"type": "integer", "minimum": 1, "maximum": 5000}, "force": S_FORCE},
})
def verify_formulas(path, write_cache=False, limit=200, force=False):
    """Evaluate every formula; report errors (#REF!, #VALUE!...), circular refs, unsupported.

    With write_cache=true the computed values are stored as cached values in the file.
    """
    book = Book(path, force=force, for_write=write_cache)
    report, values, _ = fx.verify(book.wb, limit=limit)
    report["supported_functions"] = fx.SUPPORTED_FUNCTIONS
    warnings = []
    if report["error_count"]:
        warnings.append(f"{report['error_count']} formula(s) evaluate to errors")
    if report["circular_count"]:
        warnings.append(f"{report['circular_count']} cell(s) in circular references")
    if report["unsupported_count"]:
        warnings.append(f"{report['unsupported_count']} formula(s) could not be evaluated")
    report["ok_formulas"] = report["evaluated"] - report["error_count"]
    if write_cache:
        changed, warns, details = finish(book, None, report, warnings)
        return changed, warns, details
    report["path"] = str(book.path)
    return False, warnings, report
