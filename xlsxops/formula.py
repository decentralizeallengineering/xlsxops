# SPDX-License-Identifier: Apache-2.0
"""A small pure-Python Excel formula engine.

It tokenizes, parses and evaluates the subset of Excel formulas that agents
typically write: arithmetic, comparison, concatenation, cell/range references
(including other sheets and defined names) and a fixed list of functions.
Anything outside that subset raises :class:`Unsupported` so callers can list
the formula as "not evaluated" instead of guessing a value.
"""
from __future__ import annotations

import datetime as _dt
import math
import re
import sys
import time
from dataclasses import dataclass
from decimal import ROUND_DOWN, ROUND_HALF_UP, ROUND_UP, Decimal, InvalidOperation

from openpyxl.utils.cell import column_index_from_string, get_column_letter
from openpyxl.utils.datetime import from_excel, to_excel

MAX_ROW = 1_048_576
MAX_COL = 16_384


# --------------------------------------------------------------------------
# Values and exceptions
# --------------------------------------------------------------------------
class XLError:
    """An Excel error value such as ``#REF!``."""

    __slots__ = ("code",)

    def __init__(self, code: str):
        self.code = code

    def __eq__(self, other):
        return isinstance(other, XLError) and other.code == self.code

    def __hash__(self):
        return hash(self.code)

    def __repr__(self):
        return self.code

    __str__ = __repr__


ERROR_CODES = ("#NULL!", "#DIV/0!", "#VALUE!", "#REF!", "#NAME?", "#NUM!", "#N/A",
               "#GETTING_DATA", "#SPILL!", "#CALC!")
NULL, DIV0, VALUE, REF, NAME, NUM, NA = (XLError(c) for c in ERROR_CODES[:7])

ERROR_KINDS = {"#REF!": "ref", "#VALUE!": "type", "#DIV/0!": "div0", "#N/A": "na",
               "#NAME?": "name", "#NUM!": "num", "#NULL!": "null"}


class Unsupported(Exception):
    """The formula uses something this engine does not evaluate."""


class FormulaSyntaxError(Exception):
    """The formula text cannot be parsed."""


class _Raise(Exception):
    """Internal: carries an Excel error value up the evaluation stack."""

    def __init__(self, err: XLError):
        self.err = err


class _Circular(Exception):
    def __init__(self, cells):
        self.cells = cells


class _DependsOnCircular(Exception):
    def __init__(self, cell):
        self.cell = cell


# --------------------------------------------------------------------------
# Tokenizer
# --------------------------------------------------------------------------
@dataclass
class RefTok:
    """A parsed reference token. Rows/cols are 1-based; None for open ends."""

    prefix: str          # original sheet prefix text including '!', or ''
    sheet: str | None    # unquoted sheet name, None = current sheet
    external: bool
    kind: str            # 'cell', 'area', 'cols', 'rows', 'err'
    r1: int | None = None
    c1: int | None = None
    r2: int | None = None
    c2: int | None = None
    abs_r1: bool = False
    abs_c1: bool = False
    abs_r2: bool = False
    abs_c2: bool = False

    def render(self) -> str:
        def col(c, a):
            return ("$" if a else "") + get_column_letter(c)

        def row(r, a):
            return ("$" if a else "") + str(r)

        if self.kind == "err":
            body = "#REF!"
        elif self.kind == "cell":
            body = col(self.c1, self.abs_c1) + row(self.r1, self.abs_r1)
        elif self.kind == "area":
            body = (col(self.c1, self.abs_c1) + row(self.r1, self.abs_r1) + ":"
                    + col(self.c2, self.abs_c2) + row(self.r2, self.abs_r2))
        elif self.kind == "cols":
            body = col(self.c1, self.abs_c1) + ":" + col(self.c2, self.abs_c2)
        else:
            body = row(self.r1, self.abs_r1) + ":" + row(self.r2, self.abs_r2)
        return self.prefix + body


@dataclass
class Token:
    kind: str   # NUM STR BOOL ERR REF FUNC NAME OP LP RP SEP SEMI LB RB STRUCT
    value: object
    start: int
    end: int


_SHEET = r"(?:(?P<q>'(?:[^']|'')+')|(?P<u>(?:\[[^\]]+\])?[A-Za-z0-9_.¡-￿]+))!"
_CELL = r"(\$?)([A-Za-z]{1,3})(\$?)([0-9]+)"
_RE_REF = re.compile(
    rf"(?P<prefix>{_SHEET})?(?:"
    rf"(?P<err>#REF!)"
    rf"|(?P<a>{_CELL})(?::(?P<b>{_CELL}))?"
    rf"|(?P<cols>(\$?)([A-Za-z]{{1,3}}):(\$?)([A-Za-z]{{1,3}}))"
    rf"|(?P<rows>(\$?)([0-9]+):(\$?)([0-9]+))"
    r")(?![\w.(\[!])"
)
_RE_CELL = re.compile(_CELL)
_RE_NUM = re.compile(r"(?:[0-9]+\.?[0-9]*|\.[0-9]+)(?:[eE][+-]?[0-9]+)?")
_RE_IDENT = re.compile(r"[A-Za-z_\\¡-￿][\w.¡-￿]*")
_OPS2 = ("<>", "<=", ">=")
_OPS1 = "+-*/^&=<>%"


def _parse_ref(m: re.Match) -> RefTok:
    prefix = m.group("prefix") or ""
    sheet, external = None, False
    if prefix:
        if m.group("q"):
            sheet = m.group("q")[1:-1].replace("''", "'")
        else:
            sheet = m.group("u")
        if sheet.startswith("["):
            external = True
    if m.group("err"):
        return RefTok(prefix, sheet, external, "err")
    if m.group("a"):
        ca = _RE_CELL.fullmatch(m.group("a"))
        r1, c1 = int(ca.group(4)), column_index_from_string(ca.group(2).upper())
        t = RefTok(prefix, sheet, external, "cell", r1, c1, r1, c1,
                   bool(ca.group(3)), bool(ca.group(1)), bool(ca.group(3)), bool(ca.group(1)))
        if m.group("b"):
            cb = _RE_CELL.fullmatch(m.group("b"))
            t.kind = "area"
            t.r2, t.c2 = int(cb.group(4)), column_index_from_string(cb.group(2).upper())
            t.abs_r2, t.abs_c2 = bool(cb.group(3)), bool(cb.group(1))
        return t
    if m.group("cols"):
        whole = m.group("cols")
        mm = re.fullmatch(r"(\$?)([A-Za-z]{1,3}):(\$?)([A-Za-z]{1,3})", whole)
        return RefTok(prefix, sheet, external, "cols", None,
                      column_index_from_string(mm.group(2).upper()), None,
                      column_index_from_string(mm.group(4).upper()),
                      False, bool(mm.group(1)), False, bool(mm.group(3)))
    whole = m.group("rows")
    mm = re.fullmatch(r"(\$?)([0-9]+):(\$?)([0-9]+)", whole)
    return RefTok(prefix, sheet, external, "rows", int(mm.group(2)), None,
                  int(mm.group(4)), None, bool(mm.group(1)), False, bool(mm.group(3)), False)


def tokenize(text: str) -> list[Token]:
    """Tokenize a formula (with or without the leading '=')."""
    s = text[1:] if text.startswith("=") else text
    off = len(text) - len(s)
    out: list[Token] = []
    i, n = 0, len(s)
    while i < n:
        ch = s[i]
        if ch in " \t\r\n":
            i += 1
            continue
        if ch == "@":  # implicit-intersection marker; a no-op for single values
            i += 1
            continue
        if ch == '"':
            j, buf = i + 1, []
            while True:
                if j >= n:
                    raise FormulaSyntaxError("unterminated string")
                if s[j] == '"':
                    if j + 1 < n and s[j + 1] == '"':
                        buf.append('"')
                        j += 2
                        continue
                    break
                buf.append(s[j])
                j += 1
            out.append(Token("STR", "".join(buf), off + i, off + j + 1))
            i = j + 1
            continue
        if ch == "#":
            for code in ERROR_CODES:
                if s.startswith(code, i):
                    out.append(Token("ERR", XLError(code), off + i, off + i + len(code)))
                    i += len(code)
                    break
            else:
                raise FormulaSyntaxError(f"unknown error literal at {i}")
            continue
        m = _RE_REF.match(s, i)
        if m and (ch.isdigit() is False or m.group("rows")):
            # A bare identifier like LOG10 is a function, handled by the lookahead.
            out.append(Token("REF", _parse_ref(m), off + i, off + m.end()))
            i = m.end()
            continue
        if ch.isdigit() or (ch == "." and i + 1 < n and s[i + 1].isdigit()):
            m = _RE_NUM.match(s, i)
            txt = m.group(0)
            val = float(txt)
            if val.is_integer() and "." not in txt and "e" not in txt.lower():
                val = int(txt)
            out.append(Token("NUM", val, off + i, off + m.end()))
            i = m.end()
            continue
        if ch == "[":
            raise Unsupported("external or structured reference")
        if ch == "'":
            raise FormulaSyntaxError("bad quoted sheet reference")
        m = _RE_IDENT.match(s, i)
        if m:
            word = m.group(0)
            j = m.end()
            if j < n and s[j] == "(":
                out.append(Token("FUNC", word, off + i, off + j + 1))
                i = j + 1
                continue
            if j < n and s[j] == "[":
                raise Unsupported(f"structured table reference {word}[...]")
            if word.upper() in ("TRUE", "FALSE"):
                out.append(Token("BOOL", word.upper() == "TRUE", off + i, off + j))
            else:
                out.append(Token("NAME", word, off + i, off + j))
            i = j
            continue
        two = s[i:i + 2]
        if two in _OPS2:
            out.append(Token("OP", two, off + i, off + i + 2))
            i += 2
            continue
        if ch in _OPS1:
            out.append(Token("OP", ch, off + i, off + i + 1))
            i += 1
            continue
        simple = {"(": "LP", ")": "RP", ",": "SEP", ";": "SEMI", "{": "LB", "}": "RB"}
        if ch in simple:
            out.append(Token(simple[ch], ch, off + i, off + i + 1))
            i += 1
            continue
        if ch == ":":
            raise Unsupported("range operator with a non-reference operand")
        raise FormulaSyntaxError(f"unexpected character {ch!r} at {i}")
    return out


# --------------------------------------------------------------------------
# Parser (Pratt) -> tuples
# --------------------------------------------------------------------------
_INFIX_BP = {"=": 10, "<>": 10, "<": 10, ">": 10, "<=": 10, ">=": 10,
             "&": 20, "+": 30, "-": 30, "*": 40, "/": 40, "^": 50}
_PCT_BP = 60
_NEG_BP = 70


class _Parser:
    def __init__(self, toks):
        self.toks = toks
        self.i = 0

    def peek(self):
        return self.toks[self.i] if self.i < len(self.toks) else None

    def next(self):
        t = self.peek()
        if t is None:
            raise FormulaSyntaxError("unexpected end of formula")
        self.i += 1
        return t

    def expr(self, rbp=0):
        node = self.prefix(self.next())
        while True:
            t = self.peek()
            if t is None or t.kind != "OP":
                break
            if t.value == "%":
                if _PCT_BP <= rbp:
                    break
                self.i += 1
                node = ("pct", node)
                continue
            bp = _INFIX_BP.get(t.value)
            if bp is None or bp <= rbp:
                break
            self.i += 1
            node = ("bin", t.value, node, self.expr(bp))
        return node

    def prefix(self, t):
        k = t.kind
        if k == "NUM":
            return ("num", t.value)
        if k == "STR":
            return ("str", t.value)
        if k == "BOOL":
            return ("bool", t.value)
        if k == "ERR":
            return ("err", t.value)
        if k == "REF":
            return ("ref", t.value)
        if k == "NAME":
            return ("name", t.value)
        if k == "OP" and t.value in "+-":
            return ("neg" if t.value == "-" else "pos", self.expr(_NEG_BP))
        if k == "LP":
            node = self.expr()
            if self.next().kind != "RP":
                raise FormulaSyntaxError("expected )")
            return node
        if k == "FUNC":
            args = []
            if self.peek() is not None and self.peek().kind == "RP":
                self.i += 1
                return ("func", t.value, args)
            while True:
                p = self.peek()
                if p is not None and p.kind in ("SEP", "RP"):
                    args.append(("missing",))
                else:
                    args.append(self.expr())
                nt = self.next()
                if nt.kind == "RP":
                    break
                if nt.kind != "SEP":
                    raise FormulaSyntaxError("expected , or ) in argument list")
            return ("func", t.value, args)
        if k == "LB":
            rows, row = [], []
            while True:
                tok = self.next()
                neg = False
                if tok.kind == "OP" and tok.value == "-":
                    neg, tok = True, self.next()
                if tok.kind not in ("NUM", "STR", "BOOL", "ERR"):
                    raise FormulaSyntaxError("bad array constant")
                v = -tok.value if neg else tok.value
                row.append(v)
                sep = self.next()
                if sep.kind == "SEP":
                    continue
                rows.append(row)
                row = []
                if sep.kind == "SEMI":
                    continue
                if sep.kind == "RB":
                    break
                raise FormulaSyntaxError("bad array constant")
            return ("array", rows)
        raise FormulaSyntaxError(f"unexpected token {t.value!r}")


def parse(text: str):
    toks = tokenize(text)
    if not toks:
        raise FormulaSyntaxError("empty formula")
    p = _Parser(toks)
    node = p.expr()
    if p.peek() is not None:
        raise FormulaSyntaxError(f"unexpected {p.peek().value!r}")
    return node


# --------------------------------------------------------------------------
# Formula text utilities used by the operations
# --------------------------------------------------------------------------
FUTURE_FUNCS = {
    "XLOOKUP", "XMATCH", "IFS", "IFNA", "MAXIFS", "MINIFS", "SWITCH", "CONCAT", "TEXTJOIN",
    "UNIQUE", "SEQUENCE", "RANDARRAY", "SORTBY", "LET", "LAMBDA", "STDEV.S", "STDEV.P",
    "VAR.S", "VAR.P", "AGGREGATE", "DAYS", "ISOWEEKNUM", "PERCENTILE.INC", "PERCENTILE.EXC",
    "QUARTILE.INC", "QUARTILE.EXC", "MODE.SNGL", "RANK.EQ", "RANK.AVG", "CEILING.MATH",
    "FLOOR.MATH", "TEXTBEFORE", "TEXTAFTER", "TEXTSPLIT", "VSTACK", "HSTACK", "TAKE", "DROP",
    "CHOOSECOLS", "CHOOSEROWS", "TOCOL", "TOROW", "FORECAST.LINEAR",
}
_XLWS_FUNCS = {"FILTER", "SORT"}


def add_future_prefixes(formula: str) -> str:
    """Add the ``_xlfn.`` storage prefix Excel needs for post-2007 functions.

    Without it, Excel shows #NAME? for e.g. XLOOKUP written by openpyxl.
    """
    try:
        toks = tokenize(formula)
    except (FormulaSyntaxError, Unsupported):
        return formula
    out, last = [], 0
    for t in toks:
        if t.kind != "FUNC":
            continue
        name = t.value.upper()
        if name in FUTURE_FUNCS:
            out.append(formula[last:t.start] + "_xlfn.")
            last = t.start
        elif name in _XLWS_FUNCS:
            out.append(formula[last:t.start] + "_xlfn._xlws.")
            last = t.start
    out.append(formula[last:])
    return "".join(out)


def rewrite_refs(formula: str, fn) -> str:
    """Rewrite reference tokens: ``fn(RefTok) -> RefTok | str | None``."""
    toks = tokenize(formula)
    out, last = [], 0
    for t in toks:
        if t.kind != "REF":
            continue
        new = fn(t.value)
        if new is None:
            continue
        out.append(formula[last:t.start])
        out.append(new if isinstance(new, str) else new.render())
        last = t.end
    out.append(formula[last:])
    return "".join(out)


def references(formula: str) -> list[RefTok]:
    return [t.value for t in tokenize(formula) if t.kind == "REF"]


def check_syntax(formula: str) -> str | None:
    """Return an error message if the formula cannot be parsed, else None."""
    try:
        parse(formula)
    except FormulaSyntaxError as e:
        return str(e)
    except Unsupported:
        return None
    return None


# --------------------------------------------------------------------------
# Coercion helpers
# --------------------------------------------------------------------------
_NUM_TEXT = re.compile(r"\s*[+-]?(?:\d+(?:,\d{3})*(?:\.\d*)?|\.\d+)(?:[eE][+-]?\d+)?\s*%?\s*")
_ISO_DATE = re.compile(r"\s*(\d{4})-(\d{1,2})-(\d{1,2})(?:[ T](\d{1,2}):(\d{2})(?::(\d{2}))?)?\s*")


def _text_to_number(s: str):
    if _NUM_TEXT.fullmatch(s):
        t = s.strip().replace(",", "")
        pct = t.endswith("%")
        if pct:
            t = t[:-1].strip()
        try:
            v = float(t)
        except ValueError:
            return None
        return v / 100 if pct else v
    m = _ISO_DATE.fullmatch(s)
    if m:
        try:
            d = _dt.datetime(int(m[1]), int(m[2]), int(m[3]), int(m[4] or 0), int(m[5] or 0),
                             int(m[6] or 0))
        except ValueError:
            return None
        return to_excel(d)
    return None


def general_text(v) -> str:
    """Render a value the way Excel's General format would in text contexts."""
    if v is None:
        return ""
    if isinstance(v, bool):
        return "TRUE" if v else "FALSE"
    if isinstance(v, int):
        return str(v)
    if isinstance(v, float):
        if v.is_integer() and abs(v) < 1e15:
            return str(int(v))
        s = f"{v:.15g}"
        if "e" in s:
            mant, exp = s.split("e")
            s = f"{mant}E{'+' if int(exp) >= 0 else '-'}{abs(int(exp)):02d}"
        return s
    if isinstance(v, XLError):
        return v.code
    return str(v)


def clean_number(x):
    """Normalise a numeric result: 15 significant digits, ints where exact."""
    if isinstance(x, bool):
        return x
    if isinstance(x, int):
        return x
    if math.isnan(x) or math.isinf(x):
        raise _Raise(NUM)
    x = float(f"{x:.15g}")
    if x.is_integer() and abs(x) < 2 ** 53:
        return int(x)
    return x


def _wild_regex(pattern: str):
    out, i = [], 0
    while i < len(pattern):
        c = pattern[i]
        if c == "~" and i + 1 < len(pattern) and pattern[i + 1] in "*?~":
            out.append(re.escape(pattern[i + 1]))
            i += 2
            continue
        out.append(".*" if c == "*" else "." if c == "?" else re.escape(c))
        i += 1
    return re.compile("".join(out), re.IGNORECASE | re.DOTALL)


def _has_wild(s: str) -> bool:
    return any(c in s for c in "*?~")


def _type_rank(v):
    if isinstance(v, bool):
        return 2
    if isinstance(v, (int, float)):
        return 0
    return 1


def compare(a, b) -> int:
    """Excel ordering: numbers < text < logicals; text is case-insensitive."""
    if a is None:
        a = "" if isinstance(b, str) else False if isinstance(b, bool) else 0
    if b is None:
        b = "" if isinstance(a, str) else False if isinstance(a, bool) else 0
    ra, rb = _type_rank(a), _type_rank(b)
    if ra != rb:
        return -1 if ra < rb else 1
    if ra == 1:
        a, b = a.casefold(), b.casefold()
    return (a > b) - (a < b)


# --------------------------------------------------------------------------
# Ranges
# --------------------------------------------------------------------------
class Range:
    """A rectangular block of cells, evaluated lazily."""

    __slots__ = ("ev", "sheet", "r1", "c1", "r2", "c2", "_vals")

    def __init__(self, ev, sheet, r1, c1, r2, c2):
        self.ev, self.sheet = ev, sheet
        self.r1, self.c1, self.r2, self.c2 = r1, c1, r2, c2
        self._vals = None

    @property
    def rows(self):
        return self.r2 - self.r1 + 1

    @property
    def cols(self):
        return self.c2 - self.c1 + 1

    @property
    def size(self):
        return self.rows * self.cols

    def cell(self, i, j):
        return self.ev.cell_value(self.sheet, self.r1 + i, self.c1 + j)

    def values(self):
        if self._vals is None:
            key = (self.sheet, self.r1, self.c1, self.r2, self.c2)
            cached = self.ev._range_cache.get(key)
            if cached is None:
                cached = [[self.cell(i, j) for j in range(self.cols)] for i in range(self.rows)]
                self.ev._range_cache[key] = cached
            self._vals = cached
        return self._vals

    def flat(self):
        for row in self.values():
            yield from row

    def vector(self):
        """1-D list for a single row or column, else raise."""
        if self.rows == 1:
            return self.values()[0]
        if self.cols == 1:
            return [r[0] for r in self.values()]
        raise _Raise(VALUE)

    def sub(self, i, j, rows, cols):
        if i < 0 or j < 0 or i + rows > self.rows or j + cols > self.cols:
            raise _Raise(REF)
        return Range(self.ev, self.sheet, self.r1 + i, self.c1 + j,
                     self.r1 + i + rows - 1, self.c1 + j + cols - 1)


class ArrayValue:
    """An in-formula array constant like {1,2;3,4}."""

    __slots__ = ("rows_",)

    def __init__(self, rows):
        self.rows_ = rows

    @property
    def rows(self):
        return len(self.rows_)

    @property
    def cols(self):
        return len(self.rows_[0]) if self.rows_ else 0

    @property
    def size(self):
        return self.rows * self.cols

    def values(self):
        return self.rows_

    def flat(self):
        for r in self.rows_:
            yield from r

    def vector(self):
        if self.rows == 1:
            return self.rows_[0]
        if self.cols == 1:
            return [r[0] for r in self.rows_]
        raise _Raise(VALUE)

    def cell(self, i, j):
        return self.rows_[i][j]


_Block = (Range, ArrayValue)


def scalar(v):
    if isinstance(v, _Block):
        if v.size == 1:
            return v.cell(0, 0)
        raise Unsupported("a range is used where one value is expected (array formula)")
    return v


def to_num(v):
    v = scalar(v)
    if isinstance(v, XLError):
        raise _Raise(v)
    if isinstance(v, bool):
        return int(v)
    if isinstance(v, (int, float)):
        return v
    if v is None:
        return 0
    if isinstance(v, str):
        x = _text_to_number(v)
        if x is None:
            raise _Raise(VALUE)
        return x
    raise _Raise(VALUE)


def to_text(v):
    v = scalar(v)
    if isinstance(v, XLError):
        raise _Raise(v)
    return general_text(v)


def to_bool(v):
    v = scalar(v)
    if isinstance(v, XLError):
        raise _Raise(v)
    if isinstance(v, bool):
        return v
    if v is None:
        return False
    if isinstance(v, (int, float)):
        return v != 0
    if isinstance(v, str):
        if v.upper() == "TRUE":
            return True
        if v.upper() == "FALSE":
            return False
    raise _Raise(VALUE)


def to_int(v):
    return int(math.floor(to_num(v)))


def _check(v):
    if isinstance(v, XLError):
        raise _Raise(v)
    return v


def _numbers(args, count_text_direct=True):
    """Numbers for SUM/AVERAGE/MIN/MAX: ranges skip text/logicals, literals coerce."""
    out = []
    for a in args:
        if isinstance(a, _Block):
            for v in a.flat():
                if isinstance(v, XLError):
                    raise _Raise(v)
                if isinstance(v, (int, float)) and not isinstance(v, bool):
                    out.append(v)
        elif a is None:
            continue
        else:
            out.append(to_num(a))
    return out


# --------------------------------------------------------------------------
# Criteria (COUNTIF, SUMIF, ...)
# --------------------------------------------------------------------------
def make_criterion(crit):
    crit = scalar(crit)
    if isinstance(crit, XLError):
        raise _Raise(crit)
    if crit is None:
        raise Unsupported("criterion refers to a blank cell")
    if isinstance(crit, bool):
        return lambda v: isinstance(v, bool) and v == crit
    if isinstance(crit, (int, float)):
        return lambda v: _crit_num_eq(v, crit)
    s = str(crit)
    op = "="
    for cand in ("<>", ">=", "<=", "=", ">", "<"):
        if s.startswith(cand):
            op, s = cand, s[len(cand):]
            break
    if s.upper() in ("TRUE", "FALSE"):
        b = s.upper() == "TRUE"
        eq = lambda v: isinstance(v, bool) and v == b  # noqa: E731
        return eq if op == "=" else (lambda v: not eq(v)) if op == "<>" else (lambda v: False)
    num = _text_to_number(s) if s.strip() else None
    if num is not None:
        if op == "=":
            return lambda v: _crit_num_eq(v, num)
        if op == "<>":
            return lambda v: not _crit_num_eq(v, num)
        cmp = {">": lambda a: a > num, "<": lambda a: a < num,
               ">=": lambda a: a >= num, "<=": lambda a: a <= num}[op]
        return lambda v: isinstance(v, (int, float)) and not isinstance(v, bool) and cmp(v)
    if op in ("=", "<>"):
        if s == "":
            blank = lambda v: v is None or v == ""  # noqa: E731
            if op == "<>":
                return lambda v: not blank(v)
            # "=" matches truly empty cells only; a bare "" also matches empty text.
            return (lambda v: v is None) if crit.startswith("=") else blank
        rx = _wild_regex(s)
        eq = lambda v: isinstance(v, str) and rx.fullmatch(v) is not None  # noqa: E731
        return eq if op == "=" else (lambda v: not eq(v))
    t = s.casefold()
    cmp = {">": lambda a: a > t, "<": lambda a: a < t, ">=": lambda a: a >= t,
           "<=": lambda a: a <= t}[op]
    return lambda v: isinstance(v, str) and cmp(v.casefold())


def _crit_num_eq(v, num):
    if isinstance(v, bool) or v is None:
        return False
    if isinstance(v, (int, float)):
        return v == num
    if isinstance(v, str):
        x = _text_to_number(v)
        return x is not None and x == num
    return False


def _as_block(v):
    if isinstance(v, _Block):
        return v
    raise _Raise(VALUE)


# --------------------------------------------------------------------------
# Lookup helpers
# --------------------------------------------------------------------------
def _lookup_eq(a, b, wildcard):
    """Exact-match equality as used by MATCH/VLOOKUP/XLOOKUP."""
    if isinstance(a, XLError) or isinstance(b, XLError) or b is None:
        return False
    if isinstance(a, bool) or isinstance(b, bool):
        return isinstance(a, bool) and isinstance(b, bool) and a == b
    if isinstance(a, (int, float)):
        return isinstance(b, (int, float)) and a == b
    if isinstance(a, str):
        if not isinstance(b, str):
            return False
        if wildcard and _has_wild(a):
            return _wild_regex(a).fullmatch(b) is not None
        return a.casefold() == b.casefold()
    return False


def _approx_index(vec, x, mode):
    """mode 1: largest <= x (sorted asc); -1: smallest >= x (sorted desc)."""
    best = None
    for i, v in enumerate(vec):
        if v is None or isinstance(v, XLError) or _type_rank(v) != _type_rank(x):
            continue
        c = compare(v, x)
        if mode == 1:
            if c <= 0:
                best = i
            else:
                break
        else:
            if c >= 0:
                best = i
            else:
                break
    return best


# --------------------------------------------------------------------------
# TEXT() number formatting
# --------------------------------------------------------------------------
def _split_sections(fmt):
    out, cur, q, i = [], [], False, 0
    while i < len(fmt):
        c = fmt[i]
        if c == '"':
            q = not q
        elif c == "\\" and not q and i + 1 < len(fmt):
            cur.append(fmt[i:i + 2])
            i += 2
            continue
        if c == ";" and not q:
            out.append("".join(cur))
            cur = []
        else:
            cur.append(c)
        i += 1
    out.append("".join(cur))
    return out


_DATE_TOKEN = re.compile(r'"[^"]*"|\\.|yyyy|yy|mmmmm|mmmm|mmm|mm|m|dddd|ddd|dd|d|hh|h|ss|s|'
                         r'am/pm|a/p|\.0+|\[[^\]]*\]|.', re.IGNORECASE)


def _is_date_format(sec):
    stripped = re.sub(r'"[^"]*"|\\.|\[[^\]]*\]', "", sec)
    return bool(re.search(r"[ymdhs]", stripped, re.IGNORECASE)) and not re.search(r"[0#?]", stripped.replace(".0", ""))


def _format_date(serial, sec, epoch):
    if isinstance(serial, bool) or not isinstance(serial, (int, float)):
        raise _Raise(VALUE)
    if serial < 0:
        raise _Raise(NUM)
    d = from_excel(serial, epoch=epoch)
    if isinstance(d, _dt.time):
        d = _dt.datetime.combine(_dt.date(1899, 12, 31), d)
    elif isinstance(d, _dt.timedelta):
        d = _dt.datetime(1899, 12, 31) + d
    elif not isinstance(d, _dt.datetime):
        d = _dt.datetime.combine(d, _dt.time())
    toks = _DATE_TOKEN.findall(sec)
    ampm = any(t.lower() in ("am/pm", "a/p") for t in toks)
    out = []
    for idx, t in enumerate(toks):
        lo = t.lower()
        prev = next((x.lower() for x in reversed(toks[:idx]) if re.fullmatch(r"[a-z]+", x, re.I)), "")
        nxt = next((x.lower() for x in toks[idx + 1:] if re.fullmatch(r"[a-z]+", x, re.I)), "")
        if t.startswith('"'):
            out.append(t[1:-1])
        elif t.startswith("\\"):
            out.append(t[1:])
        elif t.startswith("["):
            raise Unsupported(f"TEXT format {sec!r}")
        elif lo == "yyyy":
            out.append(f"{d.year:04d}")
        elif lo == "yy":
            out.append(f"{d.year % 100:02d}")
        elif lo in ("mm", "m") and (prev.startswith("h") or nxt.startswith("s")):
            out.append(f"{d.minute:02d}" if lo == "mm" else str(d.minute))
        elif lo == "mmmmm":
            out.append(_MONTHS[d.month - 1][0])
        elif lo == "mmmm":
            out.append(_MONTHS[d.month - 1])
        elif lo == "mmm":
            out.append(_MONTHS[d.month - 1][:3])
        elif lo == "mm":
            out.append(f"{d.month:02d}")
        elif lo == "m":
            out.append(str(d.month))
        elif lo == "dddd":
            out.append(_DAYS[d.weekday()])
        elif lo == "ddd":
            out.append(_DAYS[d.weekday()][:3])
        elif lo == "dd":
            out.append(f"{d.day:02d}")
        elif lo == "d":
            out.append(str(d.day))
        elif lo in ("hh", "h"):
            h = (d.hour % 12 or 12) if ampm else d.hour
            out.append(f"{h:02d}" if lo == "hh" else str(h))
        elif lo in ("ss", "s"):
            out.append(f"{d.second:02d}" if lo == "ss" else str(d.second))
        elif lo == "am/pm":
            out.append("AM" if d.hour < 12 else "PM")
        elif lo == "a/p":
            out.append("A" if d.hour < 12 else "P")
        elif lo.startswith(".0"):
            frac = d.microsecond / 1e6
            out.append(f"{frac:.{len(lo) - 1}f}"[1:])
        else:
            out.append(t)
    return "".join(out)


_MONTHS = ["January", "February", "March", "April", "May", "June", "July", "August",
           "September", "October", "November", "December"]
_DAYS = ["Monday", "Tuesday", "Wednesday", "Thursday", "Friday", "Saturday", "Sunday"]


def _round_half_up(x, places):
    try:
        d = Decimal(repr(float(x)))
        return d.quantize(Decimal(1).scaleb(-places), rounding=ROUND_HALF_UP)
    except InvalidOperation:
        raise _Raise(NUM)


def _format_number(x, sec):
    # Separate literal text from the digit-placeholder core.
    parts, i = [], 0  # list of (is_literal, text)
    while i < len(sec):
        c = sec[i]
        if c == '"':
            j = sec.index('"', i + 1) if '"' in sec[i + 1:] else len(sec)
            parts.append((True, sec[i + 1:j]))
            i = j + 1
        elif c == "\\" and i + 1 < len(sec):
            parts.append((True, sec[i + 1]))
            i += 2
        elif c == "_" and i + 1 < len(sec):
            parts.append((True, " "))
            i += 2
        elif c == "*" and i + 1 < len(sec):
            i += 2
        elif c in "0#?.,":
            parts.append((False, c))
            i += 1
        elif c == "%":
            parts.append((True, "%"))
            x *= 100
            i += 1
        elif c in "eE" or c == "/":
            raise Unsupported(f"TEXT format {sec!r}")
        elif c == "[":
            j = sec.find("]", i)
            if j < 0 or not re.fullmatch(r"\[\$[^\]]*\]", sec[i:j + 1]):
                raise Unsupported(f"TEXT format {sec!r}")
            sym = sec[i + 2:j].split("-")[0]
            parts.append((True, sym))
            i = j + 1
        else:
            parts.append((True, c))
            i += 1
    ph_idx = [k for k, (lit, _) in enumerate(parts) if not lit]
    if not ph_idx:
        return "".join(t for _, t in parts)
    first, last = ph_idx[0], ph_idx[-1]
    core = "".join(t for lit, t in parts[first:last + 1] if not lit)
    # Trailing commas scale by 1000.
    while core.endswith(","):
        core = core[:-1]
        x /= 1000
    int_part, has_dot, dec_part = core.partition(".")
    int_part = int_part.replace(".", "")
    decimals = sum(1 for c in dec_part if c in "0#?")
    min_dec = sum(1 for c in dec_part.rstrip("#,") if c in "0?")
    min_int = sum(1 for c in int_part if c in "0?")
    grouping = "," in int_part
    q = _round_half_up(x, decimals)
    ip, _, dp = f"{q:f}".partition(".")
    dp = (dp + "0" * decimals)[:decimals]
    while len(dp) > min_dec and dp.endswith("0"):
        dp = dp[:-1]
    ip = ip.lstrip("0").rjust(min_int, "0")
    if grouping and ip:
        ip = f"{int(ip):,}"
    num = ip + ("." + dp if has_dot else "")
    prefix = "".join(t for _, t in parts[:first])
    middle = "".join(t for lit, t in parts[first:last + 1] if lit)
    suffix = "".join(t for _, t in parts[last + 1:])
    return prefix + num + middle + suffix


def format_text(value, fmt: str, epoch=None):
    """Python implementation of TEXT(value, format) for common formats."""
    sections = _split_sections(fmt)
    if isinstance(value, str):
        x = _text_to_number(value)
        if x is None:
            if len(sections) >= 4:
                return sections[3].replace("@", value).replace('"', "")
            return value
        value = x
    if value is None:
        value = 0
    if isinstance(value, bool):
        return general_text(value)
    if fmt.strip().lower() == "general" or fmt == "":
        return general_text(value)
    if fmt.strip() == "@":
        return general_text(value)
    if len(sections) == 1:
        sec, neg_sign = sections[0], value < 0
    elif value > 0 or (value == 0 and len(sections) < 3):
        sec, neg_sign = sections[0], False
    elif value < 0:
        sec, neg_sign = sections[1], False
    else:
        sec, neg_sign = sections[2], False
    if _is_date_format(sec):
        return _format_date(value, sec, epoch)
    out = _format_number(abs(value), sec)
    if neg_sign and re.search(r"[1-9]", out):
        out = "-" + out
    return out


# --------------------------------------------------------------------------
# Functions
# --------------------------------------------------------------------------
FUNCS: dict[str, tuple] = {}


def _fn(name, lazy=False, minargs=0, maxargs=None):
    def deco(f):
        FUNCS[name] = (f, lazy, minargs, maxargs)
        return f
    return deco


@_fn("SUM", minargs=1)
def _sum(ev, a):
    return sum(_numbers(a))


@_fn("AVERAGE", minargs=1)
def _average(ev, a):
    nums = _numbers(a)
    if not nums:
        raise _Raise(DIV0)
    return sum(nums) / len(nums)


@_fn("MIN", minargs=1)
def _min(ev, a):
    nums = _numbers(a)
    return min(nums) if nums else 0


@_fn("MAX", minargs=1)
def _max(ev, a):
    nums = _numbers(a)
    return max(nums) if nums else 0


@_fn("COUNT", minargs=1)
def _count(ev, a):
    n = 0
    for x in a:
        if isinstance(x, _Block):
            n += sum(1 for v in x.flat() if isinstance(v, (int, float)) and not isinstance(v, bool))
        elif isinstance(x, (int, float)):
            n += 1
        elif isinstance(x, str) and _text_to_number(x) is not None:
            n += 1
    return n


@_fn("COUNTA", minargs=1)
def _counta(ev, a):
    n = 0
    for x in a:
        if isinstance(x, _Block):
            n += sum(1 for v in x.flat() if v is not None)
        elif x is not None:
            n += 1
    return n


@_fn("COUNTBLANK", minargs=1, maxargs=1)
def _countblank(ev, a):
    return sum(1 for v in _as_block(a[0]).flat() if v is None or v == "")


def _crit_mask(pairs):
    """pairs: [(range, criterion)] -> (rows, cols, mask list)."""
    shape = None
    masks = []
    for rng, crit in pairs:
        rng = _as_block(rng)
        if shape is None:
            shape = (rng.rows, rng.cols)
        elif (rng.rows, rng.cols) != shape:
            raise _Raise(VALUE)
        pred = make_criterion(crit)
        masks.append([pred(v) for v in rng.flat()])
    mask = [all(m) for m in zip(*masks)]
    return shape, mask


def _sized_like(rng, shape):
    rng = _as_block(rng)
    if isinstance(rng, ArrayValue):
        return rng
    return Range(rng.ev, rng.sheet, rng.r1, rng.c1, rng.r1 + shape[0] - 1, rng.c1 + shape[1] - 1)


def _masked_numbers(rng, mask):
    out = []
    for v, m in zip(rng.flat(), mask):
        if not m:
            continue
        if isinstance(v, XLError):
            raise _Raise(v)
        if isinstance(v, (int, float)) and not isinstance(v, bool):
            out.append(v)
    return out


@_fn("COUNTIF", minargs=2, maxargs=2)
def _countif(ev, a):
    _, mask = _crit_mask([(a[0], a[1])])
    return sum(mask)


@_fn("COUNTIFS", minargs=2)
def _countifs(ev, a):
    if len(a) % 2:
        raise _Raise(VALUE)
    _, mask = _crit_mask(list(zip(a[::2], a[1::2])))
    return sum(mask)


@_fn("SUMIF", minargs=2, maxargs=3)
def _sumif(ev, a):
    shape, mask = _crit_mask([(a[0], a[1])])
    target = _sized_like(a[2], shape) if len(a) > 2 and a[2] is not None else _as_block(a[0])
    return sum(_masked_numbers(target, mask))


@_fn("SUMIFS", minargs=3)
def _sumifs(ev, a):
    if len(a) % 2 == 0:
        raise _Raise(VALUE)
    shape, mask = _crit_mask(list(zip(a[1::2], a[2::2])))
    tgt = _as_block(a[0])
    if (tgt.rows, tgt.cols) != shape:
        raise _Raise(VALUE)
    return sum(_masked_numbers(tgt, mask))


@_fn("AVERAGEIF", minargs=2, maxargs=3)
def _averageif(ev, a):
    shape, mask = _crit_mask([(a[0], a[1])])
    target = _sized_like(a[2], shape) if len(a) > 2 and a[2] is not None else _as_block(a[0])
    nums = _masked_numbers(target, mask)
    if not nums:
        raise _Raise(DIV0)
    return sum(nums) / len(nums)


def _ifs_agg(a, agg):
    if len(a) % 2 == 0:
        raise _Raise(VALUE)
    shape, mask = _crit_mask(list(zip(a[1::2], a[2::2])))
    tgt = _as_block(a[0])
    if (tgt.rows, tgt.cols) != shape:
        raise _Raise(VALUE)
    return agg(_masked_numbers(tgt, mask))


@_fn("AVERAGEIFS", minargs=3)
def _averageifs(ev, a):
    def avg(n):
        if not n:
            raise _Raise(DIV0)
        return sum(n) / len(n)
    return _ifs_agg(a, avg)


@_fn("MAXIFS", minargs=3)
def _maxifs(ev, a):
    return _ifs_agg(a, lambda n: max(n) if n else 0)


@_fn("MINIFS", minargs=3)
def _minifs(ev, a):
    return _ifs_agg(a, lambda n: min(n) if n else 0)


@_fn("IF", lazy=True, minargs=2, maxargs=3)
def _if(ev, a):
    cond = to_bool(ev.eval(a[0]))
    if cond:
        return ev.eval(a[1]) if a[1][0] != "missing" else 0
    if len(a) > 2:
        return ev.eval(a[2]) if a[2][0] != "missing" else 0
    return False


@_fn("IFERROR", lazy=True, minargs=2, maxargs=2)
def _iferror(ev, a):
    try:
        v = scalar(ev.eval(a[0]))
    except _Raise:
        return ev.eval(a[1])
    if isinstance(v, XLError):
        return ev.eval(a[1])
    return v


@_fn("IFNA", lazy=True, minargs=2, maxargs=2)
def _ifna(ev, a):
    try:
        v = scalar(ev.eval(a[0]))
    except _Raise as e:
        if e.err == NA:
            return ev.eval(a[1])
        raise
    if v == NA:
        return ev.eval(a[1])
    return v


@_fn("AND", minargs=1)
def _and(ev, a):
    vals = _logicals(a)
    if not vals:
        raise _Raise(VALUE)
    return all(vals)


@_fn("OR", minargs=1)
def _or(ev, a):
    vals = _logicals(a)
    if not vals:
        raise _Raise(VALUE)
    return any(vals)


def _logicals(a):
    out = []
    for x in a:
        if isinstance(x, _Block):
            for v in x.flat():
                _check(v)
                if isinstance(v, bool):
                    out.append(v)
                elif isinstance(v, (int, float)):
                    out.append(v != 0)
        else:
            out.append(to_bool(x))
    return out


@_fn("NOT", minargs=1, maxargs=1)
def _not(ev, a):
    return not to_bool(a[0])


def _round_impl(x, n, mode):
    x, n = to_num(x), to_int(n)
    try:
        d = Decimal(repr(float(x)))
        r = d.quantize(Decimal(1).scaleb(-n), rounding=mode)
    except InvalidOperation:
        raise _Raise(NUM)
    return float(r)


@_fn("ROUND", minargs=2, maxargs=2)
def _round(ev, a):
    return _round_impl(a[0], a[1], ROUND_HALF_UP)


@_fn("ROUNDUP", minargs=2, maxargs=2)
def _roundup(ev, a):
    return _round_impl(a[0], a[1], ROUND_UP)


@_fn("ROUNDDOWN", minargs=2, maxargs=2)
def _rounddown(ev, a):
    return _round_impl(a[0], a[1], ROUND_DOWN)


@_fn("ABS", minargs=1, maxargs=1)
def _abs(ev, a):
    return abs(to_num(a[0]))


@_fn("INT", minargs=1, maxargs=1)
def _int(ev, a):
    return math.floor(to_num(a[0]))


@_fn("MOD", minargs=2, maxargs=2)
def _mod(ev, a):
    x, y = to_num(a[0]), to_num(a[1])
    if y == 0:
        raise _Raise(DIV0)
    return x - y * math.floor(x / y)


@_fn("SQRT", minargs=1, maxargs=1)
def _sqrt(ev, a):
    x = to_num(a[0])
    if x < 0:
        raise _Raise(NUM)
    return math.sqrt(x)


@_fn("SUMPRODUCT", minargs=1)
def _sumproduct(ev, a):
    blocks = [_as_block(x) for x in a]
    shape = (blocks[0].rows, blocks[0].cols)
    if any((b.rows, b.cols) != shape for b in blocks):
        raise _Raise(VALUE)
    total = 0
    cols = [list(b.flat()) for b in blocks]
    for vals in zip(*cols):
        p = 1
        for v in vals:
            _check(v)
            p *= v if isinstance(v, (int, float)) and not isinstance(v, bool) else 0
        total += p
    return total


@_fn("SUBTOTAL", minargs=2)
def _subtotal(ev, a):
    code = to_int(a[0]) % 100
    impl = {1: _average, 2: _count, 3: _counta, 4: _max, 5: _min, 9: _sum}.get(code)
    if impl is None:
        raise Unsupported(f"SUBTOTAL function code {code}")
    return impl(ev, a[1:])


@_fn("VLOOKUP", minargs=3, maxargs=4)
def _vlookup(ev, a):
    x = _check(scalar(a[0]))
    table = _as_block(a[1])
    col = to_int(a[2])
    approx = to_bool(a[3]) if len(a) > 3 and a[3] is not None else True
    if col < 1:
        raise _Raise(VALUE)
    if col > table.cols:
        raise _Raise(REF)
    keys = [r[0] for r in table.values()]
    if approx:
        idx = _approx_index(keys, x, 1)
    else:
        idx = next((i for i, k in enumerate(keys) if _lookup_eq(x, k, True)), None)
    if idx is None:
        raise _Raise(NA)
    return table.cell(idx, col - 1)


@_fn("HLOOKUP", minargs=3, maxargs=4)
def _hlookup(ev, a):
    x = _check(scalar(a[0]))
    table = _as_block(a[1])
    row = to_int(a[2])
    approx = to_bool(a[3]) if len(a) > 3 and a[3] is not None else True
    if row < 1:
        raise _Raise(VALUE)
    if row > table.rows:
        raise _Raise(REF)
    keys = table.values()[0]
    idx = _approx_index(keys, x, 1) if approx else next(
        (i for i, k in enumerate(keys) if _lookup_eq(x, k, True)), None)
    if idx is None:
        raise _Raise(NA)
    return table.cell(row - 1, idx)


@_fn("MATCH", minargs=2, maxargs=3)
def _match(ev, a):
    x = _check(scalar(a[0]))
    vec = _as_block(a[1]).vector()
    mode = to_int(a[2]) if len(a) > 2 and a[2] is not None else 1
    if mode == 0:
        idx = next((i for i, k in enumerate(vec) if _lookup_eq(x, k, True)), None)
    else:
        idx = _approx_index(vec, x, 1 if mode > 0 else -1)
    if idx is None:
        raise _Raise(NA)
    return idx + 1


@_fn("XLOOKUP", minargs=3, maxargs=6)
def _xlookup(ev, a):
    x = _check(scalar(a[0]))
    look = _as_block(a[1])
    ret = _as_block(a[2])
    not_found = a[3] if len(a) > 3 and a[3] is not None else None
    match_mode = to_int(a[4]) if len(a) > 4 and a[4] is not None else 0
    search_mode = to_int(a[5]) if len(a) > 5 and a[5] is not None else 1
    vec = look.vector()
    vertical = look.cols == 1
    if vertical and ret.rows != len(vec) or not vertical and ret.cols != len(vec):
        raise _Raise(VALUE)
    order = range(len(vec)) if search_mode > 0 else range(len(vec) - 1, -1, -1)
    idx = None
    if match_mode in (0, 2):
        idx = next((i for i in order if _lookup_eq(x, vec[i], match_mode == 2)), None)
    elif match_mode in (-1, 1):
        idx = next((i for i in order if _lookup_eq(x, vec[i], False)), None)
        if idx is None:
            best = None
            for i in order:
                v = vec[i]
                if v is None or isinstance(v, XLError) or _type_rank(v) != _type_rank(x):
                    continue
                c = compare(v, x)
                if (match_mode == -1 and c < 0 and (best is None or compare(v, vec[best]) > 0)) or \
                   (match_mode == 1 and c > 0 and (best is None or compare(v, vec[best]) < 0)):
                    best = i
            idx = best
    else:
        raise _Raise(VALUE)
    if idx is None:
        if not_found is not None:
            return not_found
        raise _Raise(NA)
    if isinstance(ret, ArrayValue):
        return ret.cell(idx, 0) if vertical else ret.cell(0, idx)
    return ret.sub(idx, 0, 1, ret.cols) if vertical else ret.sub(0, idx, ret.rows, 1)


@_fn("INDEX", minargs=2, maxargs=3)
def _index(ev, a):
    arr = _as_block(a[0])
    r = to_int(a[1]) if a[1] is not None else 0
    c = to_int(a[2]) if len(a) > 2 and a[2] is not None else None
    if c is None:
        if arr.rows == 1 and arr.cols > 1:
            r, c = 1, r
        elif arr.cols == 1:
            c = 1
        else:
            c = 0
    if r < 0 or c < 0 or r > arr.rows or c > arr.cols:
        raise _Raise(REF)
    if isinstance(arr, ArrayValue):
        if r == 0 or c == 0:
            raise Unsupported("INDEX returning a whole row/column of an array constant")
        return arr.cell(r - 1, c - 1)
    if r == 0 and c == 0:
        return arr
    if r == 0:
        return arr.sub(0, c - 1, arr.rows, 1)
    if c == 0:
        return arr.sub(r - 1, 0, 1, arr.cols)
    return arr.sub(r - 1, c - 1, 1, 1)


@_fn("TEXT", minargs=2, maxargs=2)
def _text(ev, a):
    v = _check(scalar(a[0]))
    return format_text(v, to_text(a[1]), ev.epoch)


@_fn("CONCATENATE", minargs=1)
def _concatenate(ev, a):
    return "".join(to_text(x) for x in a)


@_fn("CONCAT", minargs=1)
def _concat(ev, a):
    out = []
    for x in a:
        if isinstance(x, _Block):
            out.extend(general_text(_check(v)) for v in x.flat())
        else:
            out.append(to_text(x))
    return "".join(out)


@_fn("LEN", minargs=1, maxargs=1)
def _len(ev, a):
    return len(to_text(a[0]))


@_fn("LEFT", minargs=1, maxargs=2)
def _left(ev, a):
    n = to_int(a[1]) if len(a) > 1 else 1
    if n < 0:
        raise _Raise(VALUE)
    return to_text(a[0])[:n]


@_fn("RIGHT", minargs=1, maxargs=2)
def _right(ev, a):
    n = to_int(a[1]) if len(a) > 1 else 1
    if n < 0:
        raise _Raise(VALUE)
    s = to_text(a[0])
    return s[len(s) - n:] if n else ""


@_fn("MID", minargs=3, maxargs=3)
def _mid(ev, a):
    s, start, n = to_text(a[0]), to_int(a[1]), to_int(a[2])
    if start < 1 or n < 0:
        raise _Raise(VALUE)
    return s[start - 1:start - 1 + n]


@_fn("UPPER", minargs=1, maxargs=1)
def _upper(ev, a):
    return to_text(a[0]).upper()


@_fn("LOWER", minargs=1, maxargs=1)
def _lower(ev, a):
    return to_text(a[0]).lower()


@_fn("TRIM", minargs=1, maxargs=1)
def _trim(ev, a):
    return re.sub(" +", " ", to_text(a[0])).strip(" ")


@_fn("ISBLANK", minargs=1, maxargs=1)
def _isblank(ev, a):
    return scalar(a[0]) is None


@_fn("ISNUMBER", minargs=1, maxargs=1)
def _isnumber(ev, a):
    v = scalar(a[0])
    return isinstance(v, (int, float)) and not isinstance(v, bool)


@_fn("ISTEXT", minargs=1, maxargs=1)
def _istext(ev, a):
    return isinstance(scalar(a[0]), str)


@_fn("ISERROR", lazy=True, minargs=1, maxargs=1)
def _iserror(ev, a):
    try:
        return isinstance(scalar(ev.eval(a[0])), XLError)
    except _Raise:
        return True


@_fn("ISNA", lazy=True, minargs=1, maxargs=1)
def _isna(ev, a):
    try:
        return scalar(ev.eval(a[0])) == NA
    except _Raise as e:
        return e.err == NA


@_fn("DATE", minargs=3, maxargs=3)
def _date(ev, a):
    y, m, d = to_int(a[0]), to_int(a[1]), to_int(a[2])
    if 0 <= y < 1900:
        y += 1900
    total = y * 12 + (m - 1)
    yy, mm = divmod(total, 12)
    try:
        base = _dt.datetime(yy, mm + 1, 1) + _dt.timedelta(days=d - 1)
    except (ValueError, OverflowError):
        raise _Raise(NUM)
    return to_excel(base, epoch=ev.epoch)


def _serial_to_date(ev, v):
    x = to_num(v)
    if x < 0:
        raise _Raise(NUM)
    d = from_excel(x, epoch=ev.epoch)
    if isinstance(d, _dt.time):
        return _dt.datetime(1900, 1, 1) if ev.epoch.year == 1899 else _dt.datetime(1904, 1, 1)
    return d


@_fn("YEAR", minargs=1, maxargs=1)
def _year(ev, a):
    return _serial_to_date(ev, a[0]).year


@_fn("MONTH", minargs=1, maxargs=1)
def _month(ev, a):
    return _serial_to_date(ev, a[0]).month


@_fn("DAY", minargs=1, maxargs=1)
def _day(ev, a):
    return _serial_to_date(ev, a[0]).day


SUPPORTED_FUNCTIONS = sorted(FUNCS)


# --------------------------------------------------------------------------
# Evaluator
# --------------------------------------------------------------------------
_UNSUP = object()
_CIRC = object()


class Evaluator:
    """Evaluates formulas of an openpyxl workbook loaded with formulas."""

    def __init__(self, wb, fallback=None):
        self.wb = wb
        self.epoch = getattr(wb, "epoch", None) or _dt.datetime(1899, 12, 30)
        self.sheets = {ws.title.casefold(): ws for ws in wb.worksheets}
        self.memo: dict = {}
        self.unsupported: dict = {}
        self.circular: set = set()
        self.origin: dict = {}
        self._stack: list = []
        self._in_progress: set = set()
        self._dep_flags: list = []
        self._range_cache: dict = {}
        self._dim_cache: dict = {}
        self._names_active: set = set()
        self.fallback = fallback or {}

    # -- cell access ------------------------------------------------------
    def _ws(self, sheet):
        ws = self.sheets.get(sheet.casefold())
        if ws is None:
            raise _Raise(REF)
        return ws

    def cell_value(self, sheet, r, c):
        """Computed value of a cell (errors returned as XLError values)."""
        key = (sheet, r, c)
        if key in self.memo:
            v = self.memo[key]
            if v is _UNSUP:
                raise Unsupported(f"depends on {self.unsupported[key]['cell']} "
                                  f"({self.unsupported[key]['reason']})")
            if v is _CIRC:
                raise _DependsOnCircular(key)
            if isinstance(v, XLError) and self._dep_flags:
                self._dep_flags[-1] = True
            return v
        ws = self._ws(sheet)
        cell = ws._cells.get((r, c))
        if cell is None:
            return None
        if cell.data_type == "f":
            v = self._eval_cell(key, ws, cell)
        else:
            v = self.convert_constant(cell)
        if isinstance(v, XLError) and self._dep_flags:
            self._dep_flags[-1] = True
        return v

    def convert_constant(self, cell):
        raw = cell.value
        if raw is None:
            return None
        if cell.data_type == "e":
            return XLError(str(raw)) if str(raw) in ERROR_CODES else VALUE
        if isinstance(raw, bool):
            return raw
        if isinstance(raw, (_dt.datetime, _dt.date, _dt.time)):
            return to_excel(raw, epoch=self.epoch)
        if isinstance(raw, _dt.timedelta):
            return raw.total_seconds() / 86400
        if isinstance(raw, (int, float, str)):
            return raw
        return str(raw)

    def _eval_cell(self, key, ws, cell):
        if key in self._in_progress:
            i = self._stack.index(key)
            raise _Circular(self._stack[i:])
        raw = cell.value
        if not isinstance(raw, str):
            # ArrayFormula / DataTableFormula objects.
            reason = f"{type(raw).__name__} is not evaluated"
            self.memo[key] = _UNSUP
            self.unsupported[key] = {"cell": _addr(key), "reason": reason}
            raise Unsupported(reason)
        self._in_progress.add(key)
        self._stack.append(key)
        self._dep_flags.append(False)
        saved_sheet = getattr(self, "_sheet", None)
        try:
            try:
                v = self.eval_formula(raw, key[0])
                origin = not self._dep_flags[-1]
            except Unsupported as e:
                fb = self.fallback.get(key)
                self.memo[key] = _UNSUP
                self.unsupported[key] = {"cell": _addr(key), "reason": str(e)}
                if fb is not None:
                    self.memo[key] = fb
                    return fb
                raise
            except FormulaSyntaxError as e:
                v, origin = NAME, True
                self.unsupported[key] = {"cell": _addr(key), "reason": f"syntax: {e}"}
                self.memo[key] = _UNSUP
                raise Unsupported(f"syntax error: {e}")
            except RecursionError:
                self.memo[key] = _UNSUP
                self.unsupported[key] = {"cell": _addr(key), "reason": "dependency chain too deep"}
                raise Unsupported("dependency chain too deep")
        finally:
            self._sheet = saved_sheet
            self._in_progress.discard(key)
            self._stack.pop()
            self._dep_flags.pop()
        self.memo[key] = v
        if isinstance(v, XLError):
            self.origin[key] = origin
        return v

    def eval_formula(self, text, sheet):
        try:
            node = parse(text)
            v = self.eval(node, sheet)
            v = scalar(v)
        except _Raise as e:
            return e.err
        if v is None:
            return 0
        if isinstance(v, (int, float)) and not isinstance(v, bool):
            try:
                return clean_number(v)
            except _Raise as e:
                return e.err
        return v

    # -- AST evaluation ---------------------------------------------------
    def eval(self, node, sheet=None):
        if sheet is not None:
            self._sheet = sheet
        k = node[0]
        if k == "num" or k == "str" or k == "bool":
            return node[1]
        if k == "err":
            raise _Raise(node[1])
        if k == "ref":
            return self._ref(node[1])
        if k == "func":
            return self._call(node[1], node[2])
        if k == "bin":
            return self._binop(node[1], node[2], node[3])
        if k == "neg":
            return -to_num(self.eval(node[1]))
        if k == "pos":
            return self.eval(node[1])
        if k == "pct":
            return to_num(self.eval(node[1])) / 100
        if k == "name":
            return self._name(node[1])
        if k == "array":
            return ArrayValue(node[1])
        if k == "missing":
            return None
        raise Unsupported(f"node {k}")

    def _ref(self, t: RefTok):
        if t.external:
            raise Unsupported("reference to another workbook")
        sheet = t.sheet if t.sheet is not None else self._sheet
        if t.kind == "err":
            raise _Raise(REF)
        ws = self.sheets.get(sheet.casefold())
        if ws is None:
            raise _Raise(REF)
        sheet = ws.title
        if t.kind in ("cell", "area"):
            r1, c1, r2, c2 = t.r1, t.c1, t.r2, t.c2
        else:
            max_r, max_c = self._dims(ws)
            if t.kind == "cols":
                r1, r2, c1, c2 = 1, max_r, t.c1, t.c2
            else:
                r1, r2, c1, c2 = t.r1, t.r2, 1, max_c
        r1, r2 = min(r1, r2), max(r1, r2)
        c1, c2 = min(c1, c2), max(c1, c2)
        if r2 > MAX_ROW or c2 > MAX_COL or r1 < 1 or c1 < 1:
            raise _Raise(REF)
        if t.kind in ("cols", "rows"):
            max_r, max_c = self._dims(ws)
            r2, c2 = min(r2, max_r), min(c2, max_c)
        return Range(self, sheet, r1, c1, r2, c2)

    def _dims(self, ws):
        d = self._dim_cache.get(ws.title)
        if d is None:
            d = self._dim_cache[ws.title] = (max(ws.max_row, 1), max(ws.max_column, 1))
        return d

    def _name(self, name):
        key = name.casefold()
        if key in self._names_active:
            raise _Raise(REF)
        ws = self.sheets.get(self._sheet.casefold())
        dn = None
        if ws is not None and getattr(ws, "defined_names", None):
            for k, v in ws.defined_names.items():
                if k.casefold() == key:
                    dn = v
        if dn is None:
            for k, v in self.wb.defined_names.items():
                if k.casefold() == key:
                    dn = v
        if dn is None:
            raise _Raise(NAME)
        text = dn.attr_text or ""
        self._names_active.add(key)
        saved = self._sheet
        try:
            return self.eval(parse("=" + text), saved)
        finally:
            self._sheet = saved
            self._names_active.discard(key)

    def _call(self, name, args):
        up = name.upper()
        for p in ("_XLFN._XLWS.", "_XLFN.", "_XLWS."):
            if up.startswith(p):
                up = up[len(p):]
        spec = FUNCS.get(up)
        if spec is None:
            raise Unsupported(f"function {up}")
        impl, lazy, mn, mx = spec
        if len(args) < mn or (mx is not None and len(args) > mx):
            raise _Raise(VALUE)
        sheet = self._sheet
        if lazy:
            return impl(_LazyCtx(self, sheet), args)
        vals = []
        for a in args:
            vals.append(self.eval(a, sheet))
            self._sheet = sheet
        result = impl(self, vals)
        self._sheet = sheet
        return result

    def _binop(self, op, ln, rn):
        sheet = self._sheet
        a = self.eval(ln, sheet)
        b = self.eval(rn, sheet)
        self._sheet = sheet
        if op == "&":
            return to_text(a) + to_text(b)
        if op in ("=", "<>", "<", ">", "<=", ">="):
            a, b = _check(scalar(a)), _check(scalar(b))
            c = compare(a, b)
            return {"=": c == 0, "<>": c != 0, "<": c < 0, ">": c > 0,
                    "<=": c <= 0, ">=": c >= 0}[op]
        x, y = to_num(a), to_num(b)
        if op == "+":
            return x + y
        if op == "-":
            return x - y
        if op == "*":
            return x * y
        if op == "/":
            if y == 0:
                raise _Raise(DIV0)
            return x / y
        if op == "^":
            if x == 0 and y < 0:
                raise _Raise(DIV0)
            try:
                r = x ** y
            except OverflowError:
                raise _Raise(NUM)
            if isinstance(r, complex):
                raise _Raise(NUM)
            return r
        raise Unsupported(f"operator {op}")

    # -- whole-workbook pass ---------------------------------------------
    def formula_cells(self):
        for ws in self.wb.worksheets:
            for (r, c) in sorted(ws._cells):
                cell = ws._cells[(r, c)]
                if cell.data_type == "f":
                    yield ws, r, c, cell

    def evaluate_all(self, time_budget: float = 60.0):
        """Evaluate every formula. Returns the verification report dict."""
        deadline = time.monotonic() + time_budget
        old_limit = sys.getrecursionlimit()
        sys.setrecursionlimit(max(old_limit, 20000))
        results = {}
        timed_out = []
        try:
            for ws, r, c, cell in self.formula_cells():
                key = (ws.title, r, c)
                if time.monotonic() > deadline:
                    timed_out.append(key)
                    continue
                try:
                    results[key] = ("ok", self.cell_value(*key))
                except _Circular as e:
                    for k in e.cells:
                        self.memo[k] = _CIRC
                        self.circular.add(k)
                    if key not in self.circular:
                        results[key] = ("circular_dep", None)
                except _DependsOnCircular:
                    results[key] = ("circular_dep", None)
                except Unsupported as e:
                    if key not in self.unsupported:
                        self.unsupported[key] = {"cell": _addr(key), "reason": str(e)}
                    results[key] = ("unsupported", self.unsupported[key]["reason"])
                except Exception as e:  # an evaluator gap must not sink the operation
                    self.memo[key] = _UNSUP
                    self.unsupported[key] = {"cell": _addr(key),
                                             "reason": f"evaluator error: {type(e).__name__}"}
                    results[key] = ("unsupported", self.unsupported[key]["reason"])
            for k in self.circular:
                results[k] = ("circular", None)
        finally:
            sys.setrecursionlimit(old_limit)
        for k in timed_out:
            results[k] = ("unsupported", "time budget exceeded")
        return results


class _LazyCtx:
    """Gives lazy functions an eval() bound to the calling sheet."""

    __slots__ = ("ev", "sheet", "epoch")

    def __init__(self, ev, sheet):
        self.ev, self.sheet, self.epoch = ev, sheet, ev.epoch

    def eval(self, node):
        try:
            return self.ev.eval(node, self.sheet)
        finally:
            self.ev._sheet = self.sheet


def _addr(key):
    sheet, r, c = key
    return f"{quote_sheet(sheet)}!{get_column_letter(c)}{r}"


def quote_sheet(name: str) -> str:
    if re.fullmatch(r"[A-Za-z_][A-Za-z0-9_.]*", name) and not re.fullmatch(r"[A-Za-z]{1,3}[0-9]+", name):
        return name
    return "'" + name.replace("'", "''") + "'"


def verify(wb, time_budget: float = 60.0, limit: int = 200):
    """Evaluate all formulas in ``wb`` and build a verification report."""
    ev = Evaluator(wb)
    results = ev.evaluate_all(time_budget)
    errors, circular, unsupported = [], [], []
    values, unsupported_keys = {}, []
    for key, (status, v) in results.items():
        cell = wb[key[0]]._cells[(key[1], key[2])]
        formula = cell.value if isinstance(cell.value, str) else str(type(cell.value).__name__)
        addr = _addr(key)
        if status == "ok":
            values[key] = v
            if isinstance(v, XLError):
                errors.append({"cell": addr, "formula": formula, "error": v.code,
                               "kind": ERROR_KINDS.get(v.code, "error"),
                               "origin": ev.origin.get(key, True)})
        elif status == "circular":
            circular.append({"cell": addr, "formula": formula})
        elif status == "circular_dep":
            errors.append({"cell": addr, "formula": formula, "error": "circular",
                           "kind": "circular_dependency", "origin": False})
        else:
            unsupported.append({"cell": addr, "formula": formula, "reason": v})
            unsupported_keys.append(key)
    report = {
        "formulas": len(results),
        "evaluated": len(values),
        "error_count": len(errors),
        "circular_count": len(circular),
        "unsupported_count": len(unsupported),
        "errors": errors[:limit],
        "circular": circular[:limit],
        "unsupported": unsupported[:limit],
    }
    return report, values, unsupported_keys
