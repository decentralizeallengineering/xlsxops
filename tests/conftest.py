# SPDX-License-Identifier: Apache-2.0
"""Fixtures: every test builds its own workbooks in a temp folder."""
from __future__ import annotations

import datetime as dt

import openpyxl
import pytest
from openpyxl.chart import BarChart, Reference
from openpyxl.formatting.rule import DataBarRule
from openpyxl.styles import Font, PatternFill
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation

SALES = [
    # Region, Rep, Product, Qty, Price, Date
    ("West", "Ana", "Widget", 10, 2.5, dt.date(2024, 1, 5)),
    ("East", "Bo", "Gadget", 3, 10.0, dt.date(2024, 1, 7)),
    ("West", "Cy", "Gizmo", 7, 4.25, dt.date(2024, 2, 1)),
    ("North", "Di", "Widget", 1, 2.5, dt.date(2024, 2, 3)),
    ("East", "Bo", "Gadget", 3, 10.0, dt.date(2024, 1, 7)),   # duplicate of row 3
    ("South", "Ed", "Doohickey", 12, 1.75, dt.date(2024, 3, 9)),
    ("West", "Ana", "Gadget", 2, 10.0, dt.date(2024, 3, 11)),
    ("north", "Di", "widget", 1, 2.5, dt.date(2024, 2, 3)),   # duplicate (case) of row 5
    ("East", "Fay", "Gizmo", 40, 4.25, dt.date(2024, 4, 2)),
    ("South", "Ed", "Widget", 6, 2.5, dt.date(2024, 4, 20)),
]
PRODUCTS = [("Widget", "Hardware", 1.1), ("Gadget", "Electronics", 6.0),
            ("Gizmo", "Electronics", 2.0), ("Sprocket", "Hardware", 0.5)]

HEADER_FILL = "FF1F4E78"


def build_sales(path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sales"
    heads = ["Region", "Rep", "Product", "Qty", "Price", "Date", "Total"]
    ws.append(heads)
    for i, row in enumerate(SALES, start=2):
        ws.append(list(row) + [f"=D{i}*E{i}"])
        ws[f"E{i}"].number_format = '"$"#,##0.00'
        ws[f"F{i}"].number_format = "yyyy-mm-dd"
        ws[f"G{i}"].number_format = '"$"#,##0.00'
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFFFF")
        c.fill = PatternFill("solid", start_color=HEADER_FILL, end_color=HEADER_FILL)
    for col, w in zip("ABCDEFG", (12, 9, 13, 7, 10, 12, 11)):
        ws.column_dimensions[col].width = w
    last = len(SALES) + 1
    dv = DataValidation(type="list", formula1='"West,East,North,South"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"A2:A{last}")
    ws.conditional_formatting.add(f"D2:D{last}", DataBarRule(start_type="min", end_type="max",
                                                             color="638EC6"))
    ch = BarChart()
    ch.title = "Qty by rep"
    ch.add_data(Reference(ws, min_col=4, min_row=1, max_row=last), titles_from_data=True)
    ch.set_categories(Reference(ws, min_col=2, min_row=2, max_row=last))
    ws.add_chart(ch, "J2")

    p = wb.create_sheet("Products")
    p.append(["Product", "Category", "Cost"])
    for row in PRODUCTS:
        p.append(list(row))
    p["A1"].font = Font(italic=True)
    p.column_dimensions["B"].width = 18

    n = wb.create_sheet("Notes")
    n["A1"] = "Tax rate"
    n["B1"] = 0.0825
    n["B1"].number_format = "0.00%"
    n["A2"] = "Total qty"
    n["B2"] = "=SUM(Sales!D2:D11)"
    n["A3"] = "Merged note"
    n.merge_cells("A3:C3")
    n["A3"].fill = PatternFill("solid", start_color="FFFFFF00", end_color="FFFFFF00")
    n["A4"] = "Revenue with tax"
    n["B4"] = "=SUM(Sales!G2:G11)*(1+TaxRate)"
    n.column_dimensions["A"].width = 22
    wb.defined_names["TaxRate"] = DefinedName("TaxRate", attr_text="Notes!$B$1")
    wb.save(path)
    return path


@pytest.fixture
def sales(tmp_path):
    return str(build_sales(tmp_path / "sales.xlsx"))


def snapshot(path):
    """Things every operation must leave alone (for the round-trip test)."""
    wb = openpyxl.load_workbook(path)
    out = {"sheets": wb.sheetnames, "names": {k: v.attr_text for k, v in wb.defined_names.items()}}
    n = wb["Notes"]
    out["notes"] = {c.coordinate: (None if c.data_type == "f" else c.value, c.number_format,
                                   bool(c.font.b), c.fill.fgColor.rgb if c.fill and c.fill.fill_type else None)
                    for row in n.iter_rows() for c in row}
    out["notes_formulas"] = {c.coordinate: c.value for row in n.iter_rows() for c in row
                             if c.data_type == "f"}
    out["notes_merged"] = sorted(str(m) for m in n.merged_cells.ranges)
    out["notes_width"] = n.column_dimensions["A"].width
    p = wb["Products"]
    out["products"] = [[c.value for c in row] for row in p.iter_rows()]
    out["products_style"] = (p["A1"].font.i, p.column_dimensions["B"].width)
    s = wb["Sales"]
    out["sales_header"] = [(c.value, bool(c.font.b), c.fill.fgColor.rgb) for c in s[1]][:7]
    out["sales_widths"] = {L: s.column_dimensions[L].width for L in "ABCDEFG"}
    out["sales_formats"] = (s["E2"].number_format, s["F2"].number_format)
    out["charts"] = {ws.title: len(ws._charts) for ws in wb.worksheets}
    out["chart_titles"] = [c.title is not None for c in s._charts]
    out["dv"] = len(s.data_validations.dataValidation)
    out["cf"] = sum(len(cf.rules) for cf in s.conditional_formatting)
    return out
